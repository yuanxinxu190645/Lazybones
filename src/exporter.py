# F:\MOFCNM\src\exporter.py

import json
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from pathlib import Path


def flatten_enum_field(value) -> str:
    """把枚举标签列表展开为可读字符串"""
    if isinstance(value, list):
        parts = []
        for item in value:
            if isinstance(item, dict):
                v = item.get("value", "")
                src = item.get("source", "")
                conf = item.get("confidence", "")
                flag = "★" if src == "ai_generated" else ""
                parts.append(f"{flag}{v}[{conf}]")
            else:
                parts.append(str(item))
        return "; ".join(parts)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value) if value is not None else ""


def _build_all_columns(schema: dict, lang: str) -> list:
    """
    构建完整列定义：用户字段在前，身份字段补缺在后。
    与 schema_loader.build_prompt 的合并逻辑保持一致。
    返回 [(key, label), ...]
    """
    from schema_loader import IDENTITY_FIELDS, IDENTITY_KEYS

    label_key = "label_" + lang
    user_fields = schema.get("fields", [])
    user_keys = {f["key"] for f in user_fields}

    columns = [(f["key"], f.get(label_key, f["key"])) for f in user_fields]

    for f in IDENTITY_FIELDS:
        if f["key"] not in user_keys:
            columns.append((f["key"], f.get(label_key, f["key"])))

    return columns


def export_excel(papers: list, schema: dict, output_path: str, lang: str = "zh"):
    """
    导出 Excel，列顺序：用户字段在前，身份字段补缺在后。
    """
    columns = _build_all_columns(schema, lang)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "文献数据" if lang == "zh" else "Papers"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["序号", "文献ID"] + [label for _, label in columns]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    ai_fill = PatternFill("solid", fgColor="FFF2CC")
    na_fill = PatternFill("solid", fgColor="F2F2F2")

    for row_idx, paper in enumerate(papers, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=paper.get("_paper_id", ""))

        for col_idx, (key, label) in enumerate(columns, 3):
            value = paper.get(key, "N/A")
            cell = ws.cell(row=row_idx, column=col_idx)

            has_ai_tag = False
            if isinstance(value, list):
                for item in value:
                    if isinstance(item, dict) and item.get("source") == "ai_generated":
                        has_ai_tag = True
                        break

            cell.value = flatten_enum_field(value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            if has_ai_tag:
                cell.fill = ai_fill
            elif cell.value in ("N/A", "", None):
                cell.fill = na_fill

    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def export_markdown(papers: list, schema: dict, output_path: str,
                    lang: str = "zh", seq_map: dict = None) -> str:
    """
    导出 Markdown 文件。
    每篇文献一个二级标题，身份字段放在折叠块里，
    schema 字段逐条展开，enum_multi 展开为带置信度的列表。

    seq_map: {base_id: seq_num}，来自 ui._paper_seq_map，用于显示序号。
             传 None 时自动按列表顺序编号。
    """
    from schema_loader import IDENTITY_FIELDS, IDENTITY_KEYS

    NL = chr(10)
    label_key = "label_" + lang
    user_fields = schema.get("fields", [])
    user_keys = {f["key"] for f in user_fields}

    # 身份字段标签映射
    identity_label_map = {
        f["key"]: f.get(label_key, f["key"]) for f in IDENTITY_FIELDS
    }

    lines = []
    # 文档标题
    doc_title = "文献数据库导出" if lang == "zh" else "Literature Database Export"
    lines.append("# " + doc_title)
    lines.append("")
    lines.append("> 导出字段：用户 Schema 字段 + 身份字段  |  语言：" + lang)
    lines.append("")
    lines.append("---")
    lines.append("")

    for i, paper in enumerate(papers, 1):
        base_id = paper.get("_paper_id", "").rsplit("__", 1)[0]
        seq = seq_map.get(base_id, i) if seq_map else i
        paper_id_full = paper.get("_paper_id", "")

        # ── 二级标题 ──
        lines.append("## " + str(seq) + ". " + paper_id_full)
        lines.append("")

        # ── 元信息行 ──
        meta_parts = []
        schema_ver = paper.get("_schema_ver", "N/A")
        model_used = paper.get("_model_used", paper.get("_model", "N/A"))
        tokens = paper.get("_tokens_used", "N/A")
        added_at = str(paper.get("_added_at", "N/A"))[:19]
        meta_parts.append("**Schema:​** `" + schema_ver + "`")
        meta_parts.append("**模型:​** `" + str(model_used) + "`")
        meta_parts.append("**Tokens:​** " + str(tokens))
        meta_parts.append("**入库:​** " + added_at)
        lines.append("  ".join(meta_parts))
        lines.append("")

        # ── 身份字段（折叠块）──
        lines.append("<details>")
        lines.append("<summary><strong>📄 文献身份信息</strong></summary>")
        lines.append("")
        lines.append("| 字段 | 值 |")
        lines.append("|---|---|")
        for f in IDENTITY_FIELDS:
            key = f["key"]
            label = identity_label_map.get(key, key)
            val = paper.get(key, "N/A")
            if isinstance(val, list):
                parts = []
                for it in val:
                    if isinstance(it, dict):
                        parts.append(str(it.get("value", "")))
                    else:
                        parts.append(str(it))
                val = "; ".join(p for p in parts if p) or "N/A"
            else:
                val = str(val) if val not in (None, "") else "N/A"
            # 表格内转义管道符
            val_escaped = val.replace("|", "\\|")
            if val in ("N/A", "None", ""):
                lines.append("| " + label + " | *N/A* |")
            else:
                lines.append("| " + label + " | " + val_escaped + " |")
        lines.append("")
        lines.append("</details>")
        lines.append("")

        # ── 用户 Schema 字段 ──
        lines.append("### 抽取字段" if lang == "zh" else "### Extracted Fields")
        lines.append("")

        for field in user_fields:
            key = field["key"]
            if key in IDENTITY_KEYS:
                continue
            label = field.get(label_key, key)
            value = paper.get(key, "N/A")
            ftype = field.get("type", "text")

            lines.append("**​" + label + "​** (`" + key + "`)")
            lines.append("")

            if value in ("N/A", None, ""):
                lines.append("*N/A*")
            elif ftype == "enum_multi" and isinstance(value, list):
                for item in value:
                    if isinstance(item, dict):
                        v = str(item.get("value", ""))
                        src = item.get("source", "")
                        conf = item.get("confidence", "")
                        ev = item.get("evidence", "")
                        ai_mark = " ★AI自创" if src == "ai_generated" else ""
                        conf_mark = " `[" + conf + "]`" if conf else ""
                        lines.append("- " + v + ai_mark + conf_mark)
                        if ev:
                            lines.append("  > " + ev)
                    else:
                        lines.append("- " + str(item))
            elif isinstance(value, list):
                for item in value:
                    if isinstance(item, dict):
                        v = str(item.get("value", ""))
                        conf = item.get("confidence", "")
                        ev = item.get("evidence", "")
                        conf_mark = " `[" + conf + "]`" if conf else ""
                        lines.append("- " + v + conf_mark)
                        if ev:
                            lines.append("  > " + ev)
                    else:
                        lines.append("- " + str(item))
            else:
                lines.append(str(value))

            lines.append("")

        lines.append("---")
        lines.append("")

    content = NL.join(lines)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    Path(output_path).write_text(content, encoding="utf-8")
    return output_path

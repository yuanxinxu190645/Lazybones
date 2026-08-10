import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext, simpledialog
import threading
import time
import json
import os
import gc
import uuid
import shutil
import sqlite3
import re
import sys
import webbrowser
import tkinterdnd2 as TkinterDnD
from pathlib import Path
from datetime import datetime

from schema_loader import (load_settings, save_settings, list_prompt_versions,
                           load_schema, build_prompt, get_field_labels)
from extractor import (match_si_files, generate_paper_id,
                       process_to_text_cache, deduplicate_pairs,
                       record_document_identity, document_requires_ocr,
                       SI_PATTERN)
from performance import (MODE_IDS_BY_LABEL, MODE_LABELS,
                         PerformanceManager)
from ai_client import PROVIDERS, test_connection, extract_paper, quick_ask
from database import (init_db, get_db_path, get_db_path_by_name,
                      list_databases, create_database,
                      insert_paper, delete_paper, reject_paper,
                      get_all_papers, get_paper, update_paper_field,
                      get_stats, export_zip, import_zip,
                      inspect_backup,
                      find_duplicate_groups, merge_duplicate_group,
                      init_chat_db, list_chat_sessions, save_chat_session,
                      save_chat_message, load_chat_messages,
                      delete_chat_session)
from exporter import export_excel
from search_index import (
    build_paper_search_document as _build_paper_search_document,
    build_cached_search_documents,
    normalize_search_text as _normalize_search_text,
)
from citation_manager import (
    AVAILABLE_FIELDS, DEFAULT_INLINE_ITEM_TEMPLATE,
    DEFAULT_INLINE_TEMPLATE, DEFAULT_SCHEME, DEFAULT_TEMPLATE,
    delete_citation_scheme, format_citation_list,
    format_citation, format_inline_citation,
    load_citation_schemes, normalize_citation_scheme,
    save_citation_schemes,
    select_citation_record, upsert_citation_scheme,
)
from citation_session import CitationSessionStore
from office_bridge import OfficeBridge, OfficeBridgeError
from workflow import (can_auto_continue_pairs, next_item_index,
                      normalize_pending_files)
from theme import (DENSITY_IDS_BY_LABEL, DENSITY_LABELS,
                   THEME_IDS_BY_LABEL, THEME_LABELS, ThemeManager,
                   normalize_appearance)
from pagination import (DEFAULT_PAGE_SIZE, PAGE_SIZE_OPTIONS,
                        normalize_page_size, paginate_rows)
from update_manager import (UpdateError, check_latest_release,
                            download_update, launch_portable_updater,
                            updates_dir)
from version import APP_VERSION, GITHUB_RELEASES_URL
NL = chr(10)
NL2 = chr(10) + chr(10)



class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()
        self.title("文献智能抽取工具")
        self.geometry("1280x820")

        self.settings = load_settings()
        self.theme_manager = ThemeManager(self, self.settings)
        self.performance_manager = PerformanceManager(self.settings)
        self._ensure_project_dirs()
        self._init_db()

        self._queue_files = []
        self._review_items = []
        self._stop_flag = False
        self._all_papers_cache = []
        self._manage_papers_by_id = {}
        self._manage_search_index = {}
        self._manage_base_search_index = {}
        self._manage_filtered_rows = []
        self._manage_selected_ids = set()
        self._manage_page = 1
        self._manage_page_size = normalize_page_size(
            self.settings.get("manage_page_size", DEFAULT_PAGE_SIZE))
        self._manage_selection_sync = False
        self._filter_after_id = None
        self._chat_messages = []
        self._find_replace_history = []
        self._pending_theme_roots = {}
        self._theme_flush_after_id = None

        self._build_ui()
        self.theme_manager.apply(self.settings)
        self.bind_all("<Map>", self._on_theme_widget_map, add="+")
        self.after(2500, self._poll_system_theme)
        self._restore_pending_queue()
        self.after(5000, self._auto_check_for_updates)

    def _ensure_project_dirs(self):
        base = Path(self.settings.get("project_dir", ""))
        if not base or not str(base).strip():
            return
        for d in ["inbox", "raw", "text_cache", "extract_cache",
                  "database", "exports"]:
            (base / d).mkdir(parents=True, exist_ok=True)

    def _init_db(self):
        project_dir = self.settings.get("project_dir", "")
        db_name = self.settings.get("current_db", "main")
        self.db_path = get_db_path_by_name(project_dir, db_name)
        init_db(self.db_path)

    # ── AI 按钮工厂 ────────────────────────────────────
    def _ai_button(self, parent, text, command, **kwargs):
        """创建带茶绿色标识的 AI 功能按钮"""
        btn = tk.Button(
            parent,
            text=text,
            command=command,
            bg="#A8D5B5",
            fg="#1A3A2A",
            activebackground="#7EC8A0",
            activeforeground="#1A3A2A",
            relief=tk.FLAT,
            padx=6,
            pady=2,
            font=("微软雅黑", 9),
            cursor="hand2",
            **kwargs
        )
        btn._theme_role = "accent"
        return btn

    def _on_theme_widget_map(self, event):
        widget = getattr(event, "widget", None)
        if widget is None:
            return
        if self.theme_manager.widget_is_current(widget):
            return
        try:
            target = widget.winfo_toplevel()
            key = str(target)
        except tk.TclError:
            return
        # A page can map dozens of children in one event cycle.  Queue its
        # top-level once instead of repainting every child one after another.
        self._pending_theme_roots[key] = target
        if self._theme_flush_after_id is None:
            self._theme_flush_after_id = self.after_idle(
                self._flush_mapped_theme_widgets)

    def _flush_mapped_theme_widgets(self):
        self._theme_flush_after_id = None
        pending = list(self._pending_theme_roots.values())
        self._pending_theme_roots.clear()
        for target in pending:
            try:
                if target.winfo_exists():
                    self.theme_manager.apply_tree(target)
            except tk.TclError:
                continue

    def _poll_system_theme(self):
        try:
            self.theme_manager.poll_system_change(self.settings)
        finally:
            self.after(2500, self._poll_system_theme)

    def _build_ui(self):
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        self.tab_extract = ttk.Frame(self.notebook)
        self.tab_review = ttk.Frame(self.notebook)
        self.tab_manage = ttk.Frame(self.notebook)
        self.tab_settings = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_extract,  text="  📥 抽取  ")
        self.notebook.add(self.tab_review,   text="  🔍 审核  ")
        self.notebook.add(self.tab_manage,   text="  🗄 数据管理  ")
        self.notebook.add(self.tab_settings, text="  ⚙ 设置  ")

        self._build_extract_tab()
        self._build_review_tab()
        self._build_manage_tab()
        self._build_settings_tab()

        self.status_var = tk.StringVar(value="就绪")
        status_bar = tk.Label(self, textvariable=self.status_var,
                              bd=1, relief=tk.SUNKEN, anchor=tk.W,
                              bg="#E0E0E0", padx=8)
        status_bar._theme_role = "status"
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)

    # ════════════════════════════════════════════════════
    # 抽取页
    # ════════════════════════════════════════════════════

    def _build_extract_tab(self):
        frame = self.tab_extract

        left = ttk.LabelFrame(frame, text="文件队列", padding=6)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=False, padx=8, pady=8)

        self.tip_label = tk.Label(
            left,
            text="拖放文件到此窗口或点击下方按钮选择",
            fg="#0070C0", font=("微软雅黑", 9, "bold"),
            justify=tk.LEFT)
        self.tip_label._theme_role = "accent"
        self.tip_label.pack(pady=4)

        btn_frame = ttk.Frame(left)
        btn_frame.pack(fill=tk.X)
        ttk.Button(btn_frame, text="选择文件",
                   command=self._select_files).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="扫描inbox",
                   command=self._scan_inbox).pack(side=tk.LEFT, padx=2)

        ttk.Separator(left, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=6)
        ttk.Label(left, text="待处理队列:").pack(anchor=tk.W)

        self.queue_listbox = tk.Listbox(left, width=38, height=22,
                                        selectmode=tk.EXTENDED,
                                        font=("微软雅黑", 9))
        self.queue_listbox.pack(fill=tk.BOTH, expand=True, pady=4)

        ttk.Button(left, text="移除选中",
                   command=self._remove_from_queue).pack(pady=2)
        ttk.Button(left, text="清空队列",
                   command=self._clear_queue).pack(pady=2)

        right = ttk.Frame(frame)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=4, pady=8)

        ctrl = ttk.LabelFrame(right, text="抽取控制", padding=8)
        ctrl.pack(fill=tk.X, padx=4, pady=4)

        row1 = ttk.Frame(ctrl)
        row1.pack(fill=tk.X, pady=2)
        ttk.Label(row1, text="提示词版本:").pack(side=tk.LEFT)
        self.schema_var = tk.StringVar()
        self.schema_combo = ttk.Combobox(row1, textvariable=self.schema_var,
                                         width=28, state="readonly")
        self.schema_combo.pack(side=tk.LEFT, padx=4)
        self.schema_combo.bind("<<ComboboxSelected>>", self._on_schema_change)
        ttk.Button(row1, text="🔄", command=self._refresh_schema_list,
                   width=3).pack(side=tk.LEFT)
        self._refresh_schema_list()

        row2 = ttk.Frame(ctrl)
        row2.pack(fill=tk.X, pady=4)
        self.lang_zh_var = tk.BooleanVar(
            value=self.settings.get("lang_zh", True))
        self.lang_en_var = tk.BooleanVar(
            value=self.settings.get("lang_en", True))
        ttk.Checkbutton(row2, text="抽取中文版",
                        variable=self.lang_zh_var,
                        command=self._persist_workflow_options
                        ).pack(side=tk.LEFT)
        ttk.Checkbutton(row2, text="抽取英文版",
                        variable=self.lang_en_var,
                        command=self._persist_workflow_options
                        ).pack(side=tk.LEFT, padx=20)

        row3 = ttk.Frame(ctrl)
        row3.pack(fill=tk.X, pady=3)
        self.auto_confirm_pairs_var = tk.BooleanVar(value=self.settings.get(
            "auto_confirm_clean_pairs", True))
        self.auto_open_review_var = tk.BooleanVar(value=self.settings.get(
            "auto_open_review", True))
        ttk.Checkbutton(
            row3, text="配对无异常时自动继续",
            variable=self.auto_confirm_pairs_var,
            command=self._persist_workflow_options).pack(side=tk.LEFT)
        ttk.Checkbutton(
            row3, text="完成后自动进入审核",
            variable=self.auto_open_review_var,
            command=self._persist_workflow_options).pack(
                side=tk.LEFT, padx=20)
        ttk.Label(row3, text="（孤立 SI 仍会要求确认）",
                  foreground="#777").pack(side=tk.LEFT)

        btn_row = ttk.Frame(ctrl)
        btn_row.pack(pady=6)
        self.start_btn = self._ai_button(
            btn_row, text="▶ 开始抽取",
            command=self._start_extraction)
        self.start_btn.pack(side=tk.LEFT, padx=4)
        self.stop_btn = ttk.Button(btn_row, text="■ 停止",
                                   command=self._stop_extraction,
                                   state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=4)

        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(right, variable=self.progress_var,
                                            maximum=100)
        self.progress_bar.pack(fill=tk.X, padx=4, pady=4)

        log_frame = ttk.LabelFrame(right, text="运行日志")
        log_frame.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        self.log_text = scrolledtext.ScrolledText(
            log_frame, height=22, state=tk.DISABLED,
            font=("Consolas", 9), bg="#1E1E1E", fg="#D4D4D4")
        self.log_text._theme_role = "log"
        self.log_text.pack(fill=tk.BOTH, expand=True)

        self.log_text.tag_config("time",   foreground="#6A9955")
        self.log_text.tag_config("ok",     foreground="#4EC9B0")
        self.log_text.tag_config("error",  foreground="#F44747")
        self.log_text.tag_config("warn",   foreground="#CE9178")
        self.log_text.tag_config("info",   foreground="#9CDCFE")
        self.log_text.tag_config("ai",     foreground="#C586C0")
        self.log_text.tag_config("title",  foreground="#DCDCAA")
        self.log_text.tag_config("detail", foreground="#858585")

        try:
            from tkinterdnd2 import DND_FILES
            for widget in [frame, left, self.queue_listbox,
                           right, self.log_text]:
                widget.drop_target_register(DND_FILES)
                widget.dnd_bind("<<Drop>>", self._on_drop_files)
        except Exception:
            self.tip_label.configure(
                text="将文件扔进 inbox/ 文件夹后点击「扫描inbox」或直接选择文件",
                fg="#666")

    def _on_schema_change(self, event=None):
        self.settings["schema_version"] = self.schema_var.get()
        save_settings(self.settings)

    def _refresh_schema_list(self):
        versions = list_prompt_versions()
        self.schema_combo["values"] = versions
        saved = self.settings.get("schema_version", "")
        if saved in versions:
            self.schema_var.set(saved)
        elif versions:
            self.schema_var.set(versions[-1])
        if hasattr(self, "review_schema_combo"):
            self.review_schema_combo["values"] = ["全部"] + versions

    def _log(self, msg: str):
        """✅ 修复:所有日志末尾补回换行符"""
        self.log_text.configure(state=tk.NORMAL)
        ts = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, "[" + ts + "] ", "time")
        if any(k in msg for k in ("成功", "✓", "命中", "完成", "收录")):
            tag = "ok"
        elif any(k in msg for k in ("✗", "错误", "失败", "Error")):
            tag = "error"
        elif any(k in msg for k in ("⚠", "警告", "跳过", "截断")):
            tag = "warn"
        elif any(k in msg for k in ("调用AI", "tokens=")):
            tag = "ai"
        elif "=========" in msg:
            tag = "title"
        elif msg.startswith("   "):
            tag = "detail"
        else:
            tag = "info"
        self.log_text.insert(tk.END, msg + "\n", tag)
        self.log_text.see(tk.END)
        self.log_text.configure(state=tk.DISABLED)
        self.update_idletasks()

    def _select_files(self):
        initial_dir = self.settings.get("last_import_dir", "")
        if not Path(initial_dir).is_dir():
            initial_dir = self.settings.get("project_dir", "")
        files = filedialog.askopenfilenames(
            title="选择论文文件",
            initialdir=initial_dir or None,
            filetypes=[("文档", "*.pdf *.docx *.doc"), ("所有文件", "*.*")])
        if files:
            self.settings["last_import_dir"] = str(Path(files[0]).parent)
            save_settings(self.settings)
            self._add_to_queue(list(files))

    def _persist_workflow_options(self):
        self.settings.update({
            "lang_zh": bool(self.lang_zh_var.get()),
            "lang_en": bool(self.lang_en_var.get()),
            "auto_confirm_clean_pairs": bool(
                self.auto_confirm_pairs_var.get()),
            "auto_open_review": bool(self.auto_open_review_var.get()),
        })
        save_settings(self.settings)

    def _persist_pending_queue(self):
        self.settings["pending_files"] = list(self._queue_files)
        save_settings(self.settings)

    def _restore_pending_queue(self):
        pending = self.settings.get("pending_files", [])
        if not isinstance(pending, list):
            pending = []
        restored = self._add_to_queue(pending, persist=False)
        # Remove stale paths from the persisted state at startup.
        self._persist_pending_queue()
        if restored:
            self._log("已恢复上次未完成队列 " + str(restored) + " 个文件")

    def _scan_inbox(self):
        inbox = Path(self.settings.get("project_dir", "")) / "inbox"
        if not inbox.exists():
            messagebox.showwarning("提示", "inbox不存在:" + str(inbox))
            return
        files = (list(inbox.glob("*.pdf")) +
                 list(inbox.glob("*.docx")) +
                 list(inbox.glob("*.doc")))
        if not files:
            messagebox.showinfo("提示", "inbox文件夹为空")
            return
        self._add_to_queue([str(f) for f in files])
        self._log("从inbox扫描到 " + str(len(files)) + " 个文件")

    def _add_to_queue(self, files: list, persist: bool = True):
        additions = normalize_pending_files(files, self._queue_files)
        for path in additions:
            self._queue_files.append(path)
            self.queue_listbox.insert(tk.END, Path(path).name)
        if persist:
            self._persist_pending_queue()
        self.status_var.set("队列中 " + str(len(self._queue_files)) + " 个文件")
        return len(additions)

    def _remove_from_queue(self):
        for idx in reversed(self.queue_listbox.curselection()):
            self.queue_listbox.delete(idx)
            self._queue_files.pop(idx)
        self._persist_pending_queue()
        self.status_var.set("队列中 " + str(len(self._queue_files)) + " 个文件")

    def _clear_queue(self):
        self._queue_files.clear()
        self.queue_listbox.delete(0, tk.END)
        self._persist_pending_queue()
        self.status_var.set("就绪")

    def _stop_extraction(self):
        self._stop_flag = True
        self._log("⚠ 已请求停止")

    def _on_drop_files(self, event):
        files = self._parse_dnd_paths(event.data)
        valid = []
        for f in files:
            p = Path(f)
            if p.is_file() and p.suffix.lower() in (".pdf", ".docx", ".doc"):
                valid.append(str(p))
            elif p.is_dir():
                for ext in ("*.pdf", "*.docx", "*.doc"):
                    valid.extend(str(x) for x in p.rglob(ext))
        if valid:
            self.settings["last_import_dir"] = str(Path(valid[0]).parent)
            save_settings(self.settings)
            self._add_to_queue(valid)
            self._log("拖入 " + str(len(valid)) + " 个文件")

    @staticmethod
    def _parse_dnd_paths(raw: str) -> list:
        paths = []
        i = 0
        while i < len(raw):
            if raw[i] == "{":
                end = raw.find("}", i)
                if end == -1:
                    break
                paths.append(raw[i + 1:end])
                i = end + 1
            elif raw[i] == " ":
                i += 1
            else:
                end = raw.find(" ", i)
                if end == -1:
                    paths.append(raw[i:])
                    break
                paths.append(raw[i:end])
                i = end + 1
        return paths

    def _remove_done_file(self, file_path: str):
        """按文件名找 listbox 索引,避免异步删除时索引错位"""
        target_name = Path(file_path).name

        def _do_delete():
            if file_path not in self._queue_files:
                return
            self._queue_files.remove(file_path)
            for i in range(self.queue_listbox.size()):
                if self.queue_listbox.get(i) == target_name:
                    self.queue_listbox.delete(i)
                    break
            self.status_var.set(
                "队列中 " + str(len(self._queue_files)) + " 个文件")
            self._persist_pending_queue()

        self.after(0, _do_delete)

    def _start_extraction(self):
        if not self._queue_files:
            messagebox.showwarning("提示", "队列为空")
            return
        if not self.settings.get("api_key"):
            messagebox.showerror("错误", "请先到【设置】页填写API Key")
            return
        if not self.lang_zh_var.get() and not self.lang_en_var.get():
            messagebox.showwarning("提示", "请至少选择一种语言")
            return
        self.start_btn.configure(state=tk.DISABLED)
        self.stop_btn.configure(state=tk.NORMAL)
        self._stop_flag = False
        threading.Thread(target=self._run_extraction, daemon=True).start()

    def _validate_text_cache(self, paper_id: str, current_si_paths: list,
                             text_cache_dir: str, extract_cache_dir: str,
                             schema_ver: str) -> bool:
        """
        校验文本缓存是否与当前 SI 配对一致。
        返回 True 表示缓存有效可复用,False 表示已被清理需要重建。

        关键逻辑:对比 meta.json 里的 si_files 和当前配对的 SI 文件名列表。
        不一致时:删除 text_cache/<paper_id>/ 整个目录,
                 同时清掉 extract_cache/<paper_id>/<schema>_*.json,
                 强制下游重新提取文本和重新调用 AI。
        """
        cache_dir = Path(text_cache_dir) / paper_id
        meta_file = cache_dir / "meta.json"
        if not meta_file.exists():
            return False

        try:
            meta = json.loads(meta_file.read_text(encoding="utf-8"))
        except Exception:
            # meta.json 损坏,直接当失效
            self._safe_log("   ⚠ meta.json 损坏,缓存判定失效:" + paper_id)
            shutil.rmtree(cache_dir, ignore_errors=True)
            return False

        cached_si = sorted(meta.get("si_files", []) or [])
        current_si = sorted([Path(p).name for p in current_si_paths])

        if cached_si == current_si:
            # Older releases cached the "OCR unavailable/failed" message as
            # if it were paper text.  Now that OCR is bundled, rebuild those
            # caches once instead of sending the error marker to the AI.
            cached_text_files = [cache_dir / "main.txt"]
            cached_text_files.extend(sorted(cache_dir.glob("SI_*.txt")))
            for text_file in cached_text_files:
                try:
                    sample = text_file.read_text(
                        encoding="utf-8", errors="ignore")[:500]
                except OSError:
                    continue
                if ("[OCR 不可用" in sample or
                        "[OCR 处理失败" in sample):
                    self._safe_log(
                        "   ⚠ 发现旧版 OCR 失败缓存，自动重建:" + paper_id)
                    shutil.rmtree(cache_dir, ignore_errors=True)
                    return False
            return True

        # SI 列表不一致 → 缓存失效,清理 text_cache 和对应的 extract_cache
        self._safe_log(
            "   ⚠ 检测到 SI 配对变化,旧缓存失效,自动重建:" + paper_id)
        self._safe_log(
            "      旧 SI: " + (str(cached_si) if cached_si else "[]"))
        self._safe_log(
            "      新 SI: " + (str(current_si) if current_si else "[]"))

        shutil.rmtree(cache_dir, ignore_errors=True)

        extract_dir = Path(extract_cache_dir) / paper_id
        if extract_dir.exists():
            for cf in extract_dir.glob(schema_ver + "_*.json"):
                try:
                    cf.unlink()
                    self._safe_log("      已删除旧 AI 缓存: " + cf.name)
                except Exception:
                    pass

        return False

    def _run_extraction(self):
        files = self._queue_files.copy()
        
        self._safe_log("========= 开始分析文件类型 =========")
        raw_pairs, orphan_sis = match_si_files(files, log_fn=self._safe_log)

        if can_auto_continue_pairs(
                raw_pairs, orphan_sis,
                self.settings.get("auto_confirm_clean_pairs", True)):
            pairs = raw_pairs
            self._safe_log(
                "✓ 正文/SI 配对无异常，已按设置自动继续（共 " +
                str(len(pairs)) + " 篇）")
        else:
            result_holder = {"value": None}
            confirm_event = threading.Event()
            self.after(0, lambda: self._confirm_pairs(
                raw_pairs, orphan_sis, result_holder, confirm_event))
            confirm_event.wait()

            if result_holder["value"] is False:
                self._safe_log("已取消")
                self.after(0, lambda: (
                    self.start_btn.configure(state=tk.NORMAL),
                    self.stop_btn.configure(state=tk.DISABLED)))
                return
            pairs = result_holder["value"]

        project_dir = self.settings.get("project_dir", "")
        pairs, duplicates, identities = deduplicate_pairs(
            pairs, project_dir, log_fn=self._safe_log,
            current_db_path=self.db_path,
            current_db_name=(self.db_name_var.get()
                             if hasattr(self, "db_name_var")
                             else self.settings.get("current_db", "main")))

        for duplicate in duplicates:
            for path in ([duplicate["main_path"]] +
                         list(duplicate["si_paths"])):
                self._remove_done_file(path)

        si_updates = sum(
            1 for identity in identities.values()
            if identity.get("_dedup_state") == "si_changed")
        if pairs and (duplicates or si_updates):
            library_hits = sum(
                1 for duplicate in duplicates
                if duplicate.get("current_library"))
            history_hits = len(duplicates) - library_hits
            summary = (
                "已在 AI 调用前截停 " + str(len(duplicates)) + " 篇重复文献。"
                + chr(10) +
                "当前文献库命中：" + str(library_hits) + " 篇"
                + chr(10) +
                "历史处理记录命中：" + str(history_hits) + " 篇"
                + chr(10) +
                "正文相同但 SI 更新：" + str(si_updates) +
                " 篇（将重新抽取并更新稳定文献 ID）"
                + chr(10) + chr(10) +
                "详细命中条目和判重依据已写入抽取日志。")
            self.after(0, lambda text=summary: messagebox.showinfo(
                "导入去重", text))

        if not pairs:
            self._safe_log("所有配对均为已处理文献，本次未调用 AI。")
            self.after(0, lambda: (
                self.start_btn.configure(state=tk.NORMAL),
                self.stop_btn.configure(state=tk.DISABLED)))
            self.after(0, lambda: messagebox.showinfo(
                "去重完成",
                "本次文献均已处理，已在正文/SI 匹配阶段剔除。\n"
                "没有调用 AI，也不会产生重复费用。"))
            return

        total = len(pairs)
        # Refresh free memory at the start of each batch so balanced/high mode
        # automatically scales down when the computer is already busy.
        budget = self.performance_manager.refresh(
            self.settings, redetect=True)
        concurrent = budget.ai_workers
        self._safe_log("[性能] " + self.performance_manager.budget_summary())
        self._safe_log("========= 开始处理 " + str(total) +
                       " 篇文献(并发 " + str(concurrent) + ") =========")

        schema_ver = self.schema_var.get()
        schema = load_schema(schema_ver)
        text_cache_dir = str(Path(project_dir) / "text_cache")
        extract_cache_dir = str(Path(project_dir) / "extract_cache")

        langs = []
        if self.lang_zh_var.get():
            langs.append("zh")
        if self.lang_en_var.get():
            langs.append("en")

        # Step 1: CPU 密集型转换使用独立进程；OCR 使用单独的单进程池，
        # 避免多个模型副本同时占用显存。
        import concurrent.futures as cf

        prepared = []
        conversion_work = []
        performance_options = budget.to_dict()
        for i, (main_path, si_paths) in enumerate(pairs.items()):
            if self._stop_flag:
                break
            identity = identities.get(main_path) or {}
            paper_id = generate_paper_id(
                main_path, doi=identity.get("doi"), identity=identity)
            self._safe_log("[文本] [" + str(i + 1) + "/" + str(total) +
                           "] " + paper_id)
            try:
                cache_valid = self._validate_text_cache(
                    paper_id, si_paths,
                    text_cache_dir, extract_cache_dir, schema_ver)
                if cache_valid:
                    self._safe_log("   ✓ 文本缓存命中")
                    prepared.append(
                        (paper_id, main_path, list(si_paths), identity))
                else:
                    needs_ocr = document_requires_ocr(
                        [main_path] + list(si_paths))
                    conversion_work.append((
                        paper_id, main_path, list(si_paths), identity,
                        needs_ocr))
            except Exception as exc:
                self._safe_log(
                    "   ✗ 文本预处理失败:" + type(exc).__name__ +
                    ": " + str(exc)[:200])

        if conversion_work and not self._stop_flag:
            has_ocr = any(item[4] for item in conversion_work)
            normal_workers = min(
                budget.cpu_workers,
                max(1, budget.active_documents - (1 if has_ocr else 0)))
            self._safe_log(
                "[性能] 文本转换进程 " + str(normal_workers) +
                ("，OCR 专用进程 1" if has_ocr else ""))
            normal_pool = cf.ProcessPoolExecutor(max_workers=normal_workers)
            ocr_pool = (cf.ProcessPoolExecutor(max_workers=1)
                        if has_ocr else None)
            futures = {}
            try:
                for item in conversion_work:
                    paper_id, main_path, si_paths, identity, needs_ocr = item
                    pool = ocr_pool if needs_ocr else normal_pool
                    future = pool.submit(
                        process_to_text_cache,
                        main_path, si_paths, text_cache_dir, paper_id,
                        self.settings.get("max_chars", 80000), identity,
                        performance_options)
                    futures[future] = item

                for future in cf.as_completed(futures):
                    item = futures[future]
                    paper_id, main_path, si_paths, identity, needs_ocr = item
                    if self._stop_flag:
                        for pending in futures:
                            pending.cancel()
                        break
                    try:
                        future.result()
                        prepared.append(
                            (paper_id, main_path, si_paths, identity))
                        provider = "OCR" if needs_ocr else "CPU"
                        self._safe_log(
                            "   ✓ 文本提取完成 [" + provider + "] " +
                            paper_id)
                    except Exception as exc:
                        self._safe_log(
                            "   ✗ 文本提取失败(" + paper_id + "): " +
                            type(exc).__name__ + ": " + str(exc)[:200])
            finally:
                normal_pool.shutdown(wait=True, cancel_futures=True)
                if ocr_pool is not None:
                    ocr_pool.shutdown(wait=True, cancel_futures=True)

        # Step 2: 网络型 AI 调用使用线程池，并发数由性能模式统一控制

        success_counter = {"n": 0}
        failed_counter = {"n": 0}
        done_counter = {"n": 0}
        lock = threading.Lock()
        start_ts = time.time()

        def process_one(item):
            paper_id, main_path, si_paths, identity = item
            if self._stop_flag:
                return ("stopped", paper_id)

            try:
                merged_path = (Path(text_cache_dir) /
                               paper_id / "merged.txt")
                tables_path = (Path(text_cache_dir) /
                               paper_id / "tables.json")
                merged_text = merged_path.read_text(encoding="utf-8")
                tables = json.loads(
                    tables_path.read_text(encoding="utf-8"))

                for lang in langs:
                    if self._stop_flag:
                        return ("stopped", paper_id)
                    cache_file = (Path(extract_cache_dir) / paper_id /
                                  (schema_ver + "_" + lang + ".json"))
                    if cache_file.exists():
                        self._safe_log(
                            "   ✓ AI缓存命中 (" + paper_id +
                            " / " + lang + ")")
                        continue
                    self._safe_log(
                        "   → 调用AI (" + paper_id +
                        " / " + lang + ")...")
                    prompt = build_prompt(schema, lang)
                    result = extract_paper(
                        merged_text, tables, prompt,
                        self.settings["base_url"],
                        self.settings["api_key"],
                        self.settings["model"], lang,
                        timeout=float(self.settings.get(
                            "ai_timeout", 600)),
                        max_retries=int(self.settings.get(
                            "ai_max_retries", 3)),
                        max_input_tokens=int(self.settings.get(
                            "max_input_tokens", 100000)),
                        log_fn=self._safe_log,
                    )
                    result["_paper_id"] = paper_id
                    result["_schema_ver"] = schema_ver
                    result["_lang"] = lang
                    result["_model"] = self.settings["model"]
                    cache_file.parent.mkdir(parents=True, exist_ok=True)
                    cache_file.write_text(
                        json.dumps(result, ensure_ascii=False, indent=2),
                        encoding="utf-8")
                    self._safe_log(
                        "   ✓ AI完成 (" + paper_id + " / " + lang +
                        "),tokens=" +
                        str(result.get("_tokens_used", 0)))

                # 仅在所有目标语言均已成功或命中 AI 缓存后登记去重身份。
                # 失败的任务不登记，用户下次拖入时仍可正常重试。
                record_document_identity(
                    project_dir, paper_id, main_path, identity)
                for sp in [main_path] + list(si_paths):
                    self._remove_done_file(sp)
                return ("ok", paper_id)

            except Exception as e:
                self._safe_log(
                    "   ✗ 错误(" + paper_id + "):" +
                    type(e).__name__ + ": " + str(e)[:200])
                return ("fail", paper_id)

        with cf.ThreadPoolExecutor(max_workers=concurrent) as pool:
            futures = [pool.submit(process_one, p) for p in prepared]
            for fut in cf.as_completed(futures):
                status, _ = fut.result()
                with lock:
                    done_counter["n"] += 1
                    if status == "ok":
                        success_counter["n"] += 1
                    elif status == "fail":
                        failed_counter["n"] += 1
                    done = done_counter["n"]

                pct = done / max(1, len(prepared)) * 100
                elapsed = time.time() - start_ts
                avg = elapsed / done if done else 0
                remaining = avg * (len(prepared) - done)
                eta_str = self._format_seconds(remaining)
                self.after(0, lambda v=pct: self.progress_var.set(v))
                self.after(0, lambda d=done, t=len(prepared),
                           e=eta_str: self.status_var.set(
                               f"已完成 {d}/{t} · 预计剩余 {e}"))

        self._safe_log(
            "========= 完成:成功 " + str(success_counter["n"]) +
            ",失败 " + str(failed_counter["n"]) + " =========")
        success_count = success_counter["n"]
        failed_count = failed_counter["n"]

        def finish_ui():
            self.start_btn.configure(state=tk.NORMAL)
            self.stop_btn.configure(state=tk.DISABLED)
            if langs and self.review_lang_var.get() not in langs:
                # Setting the variable triggers a persisted review refresh.
                self.review_lang_var.set(langs[0])
            else:
                self._refresh_review_tab()
            if (self.settings.get("auto_open_review", True)
                    and success_count > 0):
                self.notebook.select(self.tab_review)
                self.review_status_label._theme_role = "success"
                self.theme_manager.apply_widget(self.review_status_label)
                self.review_status_var.set(
                    "✓ 抽取完成：成功 " + str(success_count) +
                    "，失败 " + str(failed_count) +
                    "；已自动打开待审核首篇")
                self.after(5000, lambda: self.review_status_var.set(""))
            else:
                messagebox.showinfo(
                    "完成",
                    "抽取结束" + chr(10) +
                    "成功 " + str(success_count) +
                    ",失败 " + str(failed_count) +
                    chr(10) + "请前往【审核】页面")

        self.after(0, finish_ui)

    def _do_rerun_one(self, paper_id: str, schema_ver: str, lang: str):
        """
        通用单篇重跑：基于 text_cache 重新调用 AI，
        覆盖 extract_cache，如果该记录已入库则同时覆盖数据库。
        审核页和管理页都复用这个方法。

        线程安全：本方法跑在后台线程，所有 UI 更新走 self.after / self._safe_log。
        """
        project_dir = self.settings.get("project_dir", "")
        text_cache_dir = Path(project_dir) / "text_cache"
        extract_cache_dir = Path(project_dir) / "extract_cache"

        merged_path = text_cache_dir / paper_id / "merged.txt"
        tables_path = text_cache_dir / paper_id / "tables.json"

        if not merged_path.exists():
            self._safe_log("✗ 文本缓存丢失，无法重跑：" + paper_id)
            self.after(0, lambda: messagebox.showerror(
                "重跑失败",
                "找不到文本缓存：" + paper_id + chr(10) +
                "请回抽取页重新拖入原始 PDF / Word 文件"))
            return False

        try:
            merged_text = merged_path.read_text(encoding="utf-8")
            tables = json.loads(tables_path.read_text(encoding="utf-8"))
            schema = load_schema(schema_ver)
            prompt = build_prompt(schema, lang)

            self._safe_log("→ 重跑 AI: " + paper_id + " / " + lang)

            # 删旧 AI 缓存，确保结果是新的，不是旧缓存命中
            cache_file = (extract_cache_dir / paper_id /
                          (schema_ver + "_" + lang + ".json"))
            if cache_file.exists():
                cache_file.unlink()

            result = extract_paper(
                merged_text, tables, prompt,
                self.settings["base_url"],
                self.settings["api_key"],
                self.settings["model"], lang,
                timeout=float(self.settings.get("ai_timeout", 600)),
                max_retries=int(self.settings.get("ai_max_retries", 3)),
                max_input_tokens=int(self.settings.get(
                    "max_input_tokens", 100000)),
                log_fn=self._safe_log,
            )
            result["_paper_id"] = paper_id
            result["_schema_ver"] = schema_ver
            result["_lang"] = lang
            result["_model"] = self.settings["model"]

            cache_file.parent.mkdir(parents=True, exist_ok=True)
            cache_file.write_text(
                json.dumps(result, ensure_ascii=False, indent=2),
                encoding="utf-8")
            self._safe_log("✓ 重跑完成: " + paper_id +
                           "，tokens=" +
                           str(result.get("_tokens_used", 0)))

            # 如果已入库，自动覆盖数据库记录
            full_id = paper_id + "__" + lang
            if get_paper(self.db_path, full_id):
                insert_paper(self.db_path, full_id, result,
                             schema_ver, self.settings["model"], lang)
                self._safe_log("✓ 数据库已用新结果覆盖: " + full_id)

            return True

        except Exception as e:
            self._safe_log("✗ 重跑失败 (" + paper_id + "): " +
                           type(e).__name__ + ": " + str(e)[:200])
            return False

    def _safe_log(self, msg: str):
        """线程安全地写日志(任何线程都可调)"""
        self.after(0, lambda m=msg: self._log(m))

    @staticmethod
    def _format_seconds(sec: float) -> str:
        if sec <= 0 or sec != sec:
            return "--"
        sec = int(sec)
        if sec < 60:
            return f"{sec}s"
        if sec < 3600:
            return f"{sec // 60}m{sec % 60:02d}s"
        return f"{sec // 3600}h{(sec % 3600) // 60:02d}m"

    def _confirm_pairs(self, raw_pairs: dict, orphan_sis: list,
                       result_holder: dict, confirm_event: threading.Event):
        win = tk.Toplevel(self)
        win.title("确认正文与SI配对")
        win.geometry("900x720")
        win.grab_set()

        # 工作状态：把所有数据集中到这里，UI 渲染一律读这两个字典
        # pairs:  {main_path: [si_path, ...]}
        # orphans: [si_path, ...] 未配对的 SI 列表
        state = {
            "pairs": {m: list(sis) for m, sis in raw_pairs.items()},
            "orphans": list(orphan_sis),
        }

        ttk.Label(win,
                  text="请检查正文/SI配对关系，可在下方互相调整后再开始抽取",
                  font=("微软雅黑", 10),
                  foreground="#1F4E79").pack(pady=8)

        # ── 主配对表格 ──────────────────────────────────────
        main_frame = ttk.LabelFrame(win, text="已识别的正文 + SI 配对",
                                    padding=4)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        cols = ("main", "si", "status")
        tree = ttk.Treeview(main_frame, columns=cols,
                            show="headings", height=10)
        tree.heading("main",   text="正文文件")
        tree.heading("si",     text="SI文件（多个用 ; 分隔）")
        tree.heading("status", text="状态")
        tree.column("main",   width=260)
        tree.column("si",     width=400)
        tree.column("status", width=100, anchor=tk.CENTER)
        tree.pack(fill=tk.BOTH, expand=True, side=tk.TOP)
        tree.tag_configure("ok",    background="#E8F5E9")
        tree.tag_configure("no_si", background="#FFF9C4")

        main_btn_row = ttk.Frame(main_frame)
        main_btn_row.pack(fill=tk.X, pady=4)

        # ── 孤立 SI 区域 ─────────────────────────────────────
        orphan_frame = ttk.LabelFrame(
            win, text="⚠ 未配对的 SI 文件（可指派给上方某个正文 / 改判为正文 / 放弃）",
            padding=4)
        orphan_frame.pack(fill=tk.BOTH, expand=False, padx=8, pady=4)

        orphan_tree = ttk.Treeview(
            orphan_frame, columns=("si_file",),
            show="headings", height=4)
        orphan_tree.heading("si_file", text="孤立 SI 文件")
        orphan_tree.column("si_file", width=600)
        orphan_tree.pack(fill=tk.X, side=tk.TOP)

        orphan_btn_row = ttk.Frame(orphan_frame)
        orphan_btn_row.pack(fill=tk.X, pady=4)

        # ── 渲染函数 ─────────────────────────────────────────
        def _refresh_main_tree():
            for item in tree.get_children():
                tree.delete(item)
            for main, si_list in state["pairs"].items():
                si_names = ("; ".join(Path(s).name for s in si_list)
                            if si_list else "(无SI)")
                tag = "ok" if si_list else "no_si"
                status = "✓ 已配对" if si_list else "无SI"
                tree.insert("", tk.END,
                            values=(Path(main).name, si_names, status),
                            tags=(tag,))

        def _refresh_orphan_tree():
            for item in orphan_tree.get_children():
                orphan_tree.delete(item)
            for s in state["orphans"]:
                orphan_tree.insert("", tk.END,
                                   values=(Path(s).name,))

        _refresh_main_tree()
        _refresh_orphan_tree()

        # ── 主表格操作 ───────────────────────────────────────
        def _selected_main():
            """返回当前选中的 main_path，或 None"""
            sel = tree.selection()
            if not sel:
                return None
            name = tree.item(sel[0], "values")[0]
            for m in state["pairs"]:
                if Path(m).name == name:
                    return m
            return None

        def _demote_main_to_orphan_si():
            """把选中行的正文整个降级为孤立 SI（误判修正）"""
            main = _selected_main()
            if not main:
                messagebox.showwarning("提示", "请先选中一行", parent=win)
                return
            if not messagebox.askyesno(
                    "确认",
                    "把「" + Path(main).name + "」改判为 SI？\n"
                    "原本配在它下面的 SI 文件会全部变成孤立 SI",
                    parent=win):
                return
            si_list = state["pairs"].pop(main)
            state["orphans"].append(main)
            state["orphans"].extend(si_list)
            _refresh_main_tree()
            _refresh_orphan_tree()

        def _release_si_from_main():
            """把选中行的 SI 全部释放回孤立 SI 区"""
            main = _selected_main()
            if not main:
                messagebox.showwarning("提示", "请先选中一行", parent=win)
                return
            si_list = state["pairs"][main]
            if not si_list:
                messagebox.showinfo("提示", "该行没有 SI 可释放", parent=win)
                return
            state["pairs"][main] = []
            state["orphans"].extend(si_list)
            _refresh_main_tree()
            _refresh_orphan_tree()

        def _remove_main_row():
            """从队列里完全移除这一行（正文+SI 都不参与抽取）"""
            main = _selected_main()
            if not main:
                messagebox.showwarning("提示", "请先选中一行", parent=win)
                return
            if not messagebox.askyesno(
                    "确认",
                    "移除「" + Path(main).name + "」及其 SI？\n该正文将不参与本次抽取",
                    parent=win):
                return
            del state["pairs"][main]
            _refresh_main_tree()

        ttk.Button(main_btn_row, text="↓ 改判为SI",
                   command=_demote_main_to_orphan_si).pack(side=tk.LEFT, padx=4)
        ttk.Button(main_btn_row, text="↓ 释放此行SI",
                   command=_release_si_from_main).pack(side=tk.LEFT, padx=4)
        ttk.Button(main_btn_row, text="✗ 移除此正文",
                   command=_remove_main_row).pack(side=tk.LEFT, padx=4)
        ttk.Label(main_btn_row,
                  text="（误判为正文的可点「改判为SI」，已配错的可点「释放此行SI」）",
                  foreground="#888").pack(side=tk.LEFT, padx=8)

        # ── 孤立 SI 操作 ─────────────────────────────────────
        def _selected_orphan_paths():
            sel = orphan_tree.selection()
            if not sel:
                return []
            names = [orphan_tree.item(s, "values")[0] for s in sel]
            return [s for s in state["orphans"] if Path(s).name in names]

        def _assign_to_existing_main():
            """把选中的孤立 SI 指派给主表格里已有的某个正文"""
            orphans_picked = _selected_orphan_paths()
            if not orphans_picked:
                messagebox.showwarning("提示", "请先选中孤立 SI", parent=win)
                return
            mains_available = list(state["pairs"].keys())
            if not mains_available:
                messagebox.showwarning(
                    "提示",
                    "主表格里没有正文可挂靠，请先用「↑ 改判为正文」创建一个，"
                    "或用「指定外部正文」从磁盘选择",
                    parent=win)
                return

            picker = tk.Toplevel(win)
            picker.title("选择要挂靠的正文")
            picker.geometry("600x320")
            picker.grab_set()
            ttk.Label(picker, text="选择正文：",
                      font=("微软雅黑", 9, "bold")).pack(anchor=tk.W,
                                                       padx=8, pady=4)
            lb = tk.Listbox(picker, font=("微软雅黑", 9),
                            selectmode=tk.SINGLE)
            lb.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
            for m in mains_available:
                lb.insert(tk.END, Path(m).name)

            def _do_pick():
                idx = lb.curselection()
                if not idx:
                    messagebox.showwarning("提示", "请选一个正文",
                                           parent=picker)
                    return
                target_main = mains_available[idx[0]]
                for s in orphans_picked:
                    state["pairs"][target_main].append(s)
                    state["orphans"].remove(s)
                _refresh_main_tree()
                _refresh_orphan_tree()
                picker.destroy()

            btn_row = ttk.Frame(picker)
            btn_row.pack(pady=6)
            ttk.Button(btn_row, text="✓ 确定",
                       command=_do_pick).pack(side=tk.LEFT, padx=8)
            ttk.Button(btn_row, text="✗ 取消",
                       command=picker.destroy).pack(side=tk.LEFT, padx=8)

        def _assign_to_external_main():
            """从磁盘选一个文件作为正文，挂靠选中的孤立 SI"""
            orphans_picked = _selected_orphan_paths()
            if not orphans_picked:
                messagebox.showwarning("提示", "请先选中孤立 SI", parent=win)
                return
            main_path = filedialog.askopenfilename(
                title="选择对应的正文文件",
                filetypes=[("文档", "*.pdf *.docx *.doc")],
                parent=win)
            if not main_path:
                return
            if main_path not in state["pairs"]:
                state["pairs"][main_path] = []
            for s in orphans_picked:
                state["pairs"][main_path].append(s)
                state["orphans"].remove(s)
            _refresh_main_tree()
            _refresh_orphan_tree()

        def _promote_orphan_to_main():
            """把孤立 SI 直接当成正文加进主表格（误判修正反向操作）"""
            orphans_picked = _selected_orphan_paths()
            if not orphans_picked:
                messagebox.showwarning("提示", "请先选中孤立 SI", parent=win)
                return
            for s in orphans_picked:
                if s not in state["pairs"]:
                    state["pairs"][s] = []
                state["orphans"].remove(s)
            _refresh_main_tree()
            _refresh_orphan_tree()

        def _drop_orphan():
            """放弃选中的孤立 SI（不参与抽取）"""
            orphans_picked = _selected_orphan_paths()
            if not orphans_picked:
                messagebox.showwarning("提示", "请先选中孤立 SI", parent=win)
                return
            for s in orphans_picked:
                state["orphans"].remove(s)
            _refresh_orphan_tree()

        ttk.Button(orphan_btn_row, text="📎 挂靠到上方正文",
                   command=_assign_to_existing_main).pack(side=tk.LEFT, padx=4)
        ttk.Button(orphan_btn_row, text="📁 指定外部正文",
                   command=_assign_to_external_main).pack(side=tk.LEFT, padx=4)
        ttk.Button(orphan_btn_row, text="↑ 改判为正文",
                   command=_promote_orphan_to_main).pack(side=tk.LEFT, padx=4)
        ttk.Button(orphan_btn_row, text="✗ 放弃",
                   command=_drop_orphan).pack(side=tk.LEFT, padx=4)
        ttk.Label(orphan_btn_row,
                  text="（多选用 Ctrl/Shift 点击）",
                  foreground="#888").pack(side=tk.LEFT, padx=8)

        # ── 确认/取消 ─────────────────────────────────────────
        bottom = ttk.Frame(win)
        bottom.pack(pady=8)

        def on_confirm():
            # 检查是否还有孤立 SI 没处理
            if state["orphans"]:
                if not messagebox.askyesno(
                        "提示",
                        "还有 " + str(len(state["orphans"])) +
                        " 个孤立 SI 未处理，它们将不参与抽取，确认继续？",
                        parent=win):
                    return
            result_holder["value"] = dict(state["pairs"])
            confirm_event.set()
            win.destroy()

        def on_cancel():
            result_holder["value"] = False
            confirm_event.set()
            win.destroy()

        ttk.Button(bottom, text="✓ 确认，开始抽取",
                   command=on_confirm,
                   style="Accent.TButton").pack(side=tk.LEFT, padx=8)
        ttk.Button(bottom, text="✗ 取消",
                   command=on_cancel).pack(side=tk.LEFT, padx=8)
        win.protocol("WM_DELETE_WINDOW", on_cancel)

    # ════════════════════════════════════════════════════
    # 审核页
    # ════════════════════════════════════════════════════

    def _build_review_tab(self):
        frame = self.tab_review

        top = ttk.Frame(frame)
        top.pack(fill=tk.X, padx=8, pady=4)
        ttk.Button(top, text="🔄 刷新",
                   command=self._refresh_review_tab).pack(side=tk.LEFT, padx=4)
        ttk.Label(top, text="语言:").pack(side=tk.LEFT, padx=(12, 2))
        saved_review_lang = self.settings.get("review_lang", "zh")
        if saved_review_lang not in ("zh", "en"):
            saved_review_lang = "zh"
        self.review_lang_var = tk.StringVar(value=saved_review_lang)
        ttk.Combobox(top, textvariable=self.review_lang_var,
                     values=["zh", "en"], width=6,
                     state="readonly").pack(side=tk.LEFT)
        ttk.Label(top, text="版本:").pack(side=tk.LEFT, padx=(12, 2))
        self.review_schema_var = tk.StringVar(value=self.settings.get(
            "review_schema_filter", "全部"))
        self.review_schema_combo = ttk.Combobox(
            top, textvariable=self.review_schema_var,
            width=22, state="readonly")
        self.review_schema_combo.pack(side=tk.LEFT)
        self.review_schema_combo.bind(
            "<<ComboboxSelected>>",
            self._on_review_filter_change)
        self.review_lang_var.trace_add(
            "write", lambda *args: self._on_review_filter_change())

        # ── 勾选操作快捷按钮 ──
        ttk.Separator(top, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=8)
        ttk.Button(top, text="☑ 全选",
                   command=self._review_check_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="☐ 全不选",
                   command=self._review_uncheck_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="⇄ 反选",
                   command=self._review_invert_check).pack(side=tk.LEFT, padx=2)
        self.review_check_count_var = tk.StringVar(value="已勾 0 条")
        tk.Label(top, textvariable=self.review_check_count_var,
                 fg="#0070C0",
                 font=("微软雅黑", 9)).pack(side=tk.LEFT, padx=8)

        paned = ttk.PanedWindow(frame, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
        self.review_paned = paned
        self._review_split_initialized = False
        paned.bind("<Map>", self._initialize_review_split, add="+")

        # ── 左侧:Treeview 替代 Listbox,首列做勾选框 ──
        left_frame = ttk.LabelFrame(paned, text="待审核(点击首列勾选)")
        paned.add(left_frame, weight=2)

        cols = ("check", "paper_id", "ver_lang")
        self.review_tree = ttk.Treeview(
            left_frame, columns=cols, show="headings",
            selectmode="browse")
        self.review_tree.heading("check", text="✓")
        self.review_tree.heading("paper_id", text="文献ID")
        self.review_tree.heading("ver_lang", text="版本/语言")
        self.review_tree.column("check", width=40, minwidth=40,
                                anchor=tk.CENTER,
                                stretch=False)
        # ID 吸收宽度变化；版本/语言保持完整可见，避免默认打开时被裁切。
        self.review_tree.column("paper_id", width=205, minwidth=120,
                                stretch=True)
        self.review_tree.column("ver_lang", width=132, minwidth=118,
                                anchor=tk.CENTER, stretch=False)
        self.review_tree.tag_configure("checked", background="#E8F5E9")

        rv_sb = ttk.Scrollbar(left_frame, orient=tk.VERTICAL,
                              command=self.review_tree.yview)
        self.review_tree.configure(yscrollcommand=rv_sb.set)
        self.review_tree.pack(side=tk.LEFT, fill=tk.BOTH,
                              expand=True, padx=4, pady=4)
        rv_sb.pack(side=tk.RIGHT, fill=tk.Y)

        # 单击首列切换勾选;点击其它列只选中(右侧自动渲染)
        self.review_tree.bind("<Button-1>", self._on_review_tree_click)
        self.review_tree.bind("<<TreeviewSelect>>",
                              self._on_review_tree_select)
        # 记录勾选状态:{tree_item_id: True/False}
        self._review_checked = {}

        right_frame = ttk.Frame(paned)
        paned.add(right_frame, weight=5)
        self.review_detail = scrolledtext.ScrolledText(
            right_frame, font=("微软雅黑", 9), wrap=tk.WORD)
        self.review_detail.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        self.review_detail.tag_config("title", foreground="#1F4E79",
                                      font=("微软雅黑", 11, "bold"))
        self.review_detail.tag_config("identity_block",
                                      background="#F0F7FF",
                                      foreground="#1F4E79",
                                      font=("微软雅黑", 9))
        self.review_detail.tag_config("identity_label",
                                      foreground="#1F4E79",
                                      font=("微软雅黑", 9, "bold"))
        self.review_detail.tag_config("section",
                                      foreground="#107C10",
                                      font=("微软雅黑", 10, "bold"))
        self.review_detail.tag_config("field", foreground="#0070C0",
                                      font=("微软雅黑", 9, "bold"))
        self.review_detail.tag_config("ai_tag", background="#FFF2CC")
        self.review_detail.tag_config("evidence", foreground="#666",
                                      font=("微软雅黑", 8, "italic"))
        self.review_detail.tag_config("na", foreground="#999")
        self.review_detail.tag_config(
            "doi_link", foreground="#0066CC",
            underline=True, font=("微软雅黑", 9, "bold"))
        self.review_detail.tag_bind(
            "doi_link", "<Enter>",
            lambda event: self.review_detail.configure(cursor="hand2"))
        self.review_detail.tag_bind(
            "doi_link", "<Leave>",
            lambda event: self.review_detail.configure(cursor="xterm"))
        self.review_detail.bind("<Key>", self._readonly_text_key)
        self._install_readonly_copy_support(self.review_detail)
        self.review_detail.bind(
            "<Control-Return>", self._review_accept_shortcut)
        self.review_tree.bind(
            "<Control-Return>", self._review_accept_shortcut)
        self._review_current_doi = ""

        btn_frame = ttk.Frame(right_frame)
        btn_frame.pack(fill=tk.X, padx=4, pady=4)
        ttk.Button(btn_frame, text="✓ 收录入库(当前)",
                   command=self._accept_paper).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="✓ 收录并下一篇",
                   command=self._accept_and_next,
                   style="Accent.TButton").pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="✗ 丢弃(当前)",
                   command=self._reject_paper).pack(side=tk.LEFT, padx=4)
        ttk.Separator(btn_frame, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=6)
        # 批量操作:基于勾选
        ttk.Button(btn_frame, text="✓✓ 批量收录(已勾)",
                   command=self._batch_accept_checked
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_frame, text="✗✗ 批量丢弃(已勾)",
                   command=self._batch_reject_checked
                   ).pack(side=tk.LEFT, padx=4)
        self._ai_button(btn_frame, text="↺ 重跑",
                        command=self._rerun_paper).pack(side=tk.LEFT, padx=4)

        # 重跑状态标签:常态隐藏,重跑时显示进度,完成后自动隐藏
        self.review_status_var = tk.StringVar(value="")
        self.review_status_label = tk.Label(
            btn_frame, textvariable=self.review_status_var,
            fg="#0070C0", font=("微软雅黑", 9, "bold"))
        self.review_status_label._theme_role = "accent"
        self.review_status_label.pack(side=tk.LEFT, padx=12)

        ttk.Button(btn_frame, text="📂 打开缓存目录",
                   command=self._open_cache_dir).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text="📋 复制 DOI",
                   command=self._copy_review_doi).pack(side=tk.RIGHT, padx=4)
        ttk.Button(btn_frame, text="下一篇 ▶",
                   command=lambda: self._move_review_selection(1)
                   ).pack(side=tk.RIGHT, padx=2)
        ttk.Button(btn_frame, text="◀ 上一篇",
                   command=lambda: self._move_review_selection(-1)
                   ).pack(side=tk.RIGHT, padx=2)

        self.bind("<Control-Return>", self._review_accept_shortcut, add="+")
        self.bind("<Alt-Left>",
                  lambda event: self._review_navigation_shortcut(-1), add="+")
        self.bind("<Alt-Right>",
                  lambda event: self._review_navigation_shortcut(1), add="+")

    def _initialize_review_split(self, event=None):
        """Give the review list a useful width the first time it is shown."""
        if self._review_split_initialized:
            return
        self.after_idle(self._apply_initial_review_split)

    def _apply_initial_review_split(self):
        if self._review_split_initialized:
            return
        try:
            width = self.review_paned.winfo_width()
            if width <= 100:
                self.after(50, self._apply_initial_review_split)
                return
            # Roughly one third of the page, with enough room for all 3 columns.
            target = max(390, round(width * 0.32))
            target = min(target, max(300, width - 560))
            self.review_paned.sashpos(0, target)
            self._review_split_initialized = True
        except tk.TclError:
            pass

    def _on_review_filter_change(self, event=None):
        self.settings["review_lang"] = self.review_lang_var.get()
        self.settings["review_schema_filter"] = self.review_schema_var.get()
        save_settings(self.settings)
        self._refresh_review_tab()

    def _refresh_review_tab(self):
        project_dir = self.settings.get("project_dir", "")
        extract_cache_dir = Path(project_dir) / "extract_cache"

        versions = ["全部"] + list_prompt_versions()
        self.review_schema_combo["values"] = versions
        if self.review_schema_var.get() not in versions:
            self.review_schema_var.set("全部")

        # 清空 tree 和勾选状态
        for item in self.review_tree.get_children():
            self.review_tree.delete(item)
        self._review_items = []
        self._review_checked = {}

        if not extract_cache_dir.exists():
            self._update_review_check_count()
            return

        filter_lang = self.review_lang_var.get()
        filter_ver = self.review_schema_var.get()
        current_ver = self.schema_var.get() if hasattr(self, "schema_var") else ""
        effective_filter = filter_ver if filter_ver != "全部" else current_ver

        for paper_dir in sorted(extract_cache_dir.iterdir()):
            if not paper_dir.is_dir():
                continue
            for cache_file in sorted(paper_dir.glob("*.json")):
                stem = cache_file.stem
                if not stem.endswith("_" + filter_lang):
                    continue
                schema_ver = stem[:-len("_" + filter_lang)]
                if effective_filter and schema_ver != effective_filter:
                    continue

                paper_id = paper_dir.name
                in_db = get_paper(self.db_path,
                                  paper_id + "__" + filter_lang)
                if in_db:
                    continue

                ver_lang = schema_ver + "/" + filter_lang
                tree_id = self.review_tree.insert(
                    "", tk.END,
                    values=("☐", paper_id, ver_lang))
                self._review_checked[tree_id] = False
                self._review_items.append({
                    "paper_id": paper_id,
                    "schema_ver": schema_ver,
                    "lang": filter_lang,
                    "cache_file": str(cache_file),
                    "tree_id": tree_id,
                })

        self.status_var.set("待审核 " + str(len(self._review_items)) + " 条")
        self._update_review_check_count()
        children = self.review_tree.get_children()
        if children:
            first = children[0]
            self.review_tree.selection_set(first)
            self.review_tree.focus(first)
            self.review_tree.see(first)
            self._on_review_tree_select()
        elif hasattr(self, "review_detail"):
            self.review_detail.configure(state=tk.NORMAL)
            self.review_detail.delete("1.0", tk.END)
            self.review_detail.insert(
                tk.END, "当前筛选条件下没有待审核文献。")
            self.review_detail.configure(state=tk.DISABLED)

    def _on_review_tree_click(self, event):
        """单击 tree:点首列时切换勾选,点其它列只走 select 逻辑"""
        region = self.review_tree.identify("region", event.x, event.y)
        if region != "cell":
            return
        col = self.review_tree.identify_column(event.x)
        item = self.review_tree.identify_row(event.y)
        if not item:
            return
        if col == "#1":   # 首列 check
            self._toggle_review_check(item)
            return "break"   # 阻止默认 select,避免和勾选冲突
        # 其它列让默认行为继续(会触发 <<TreeviewSelect>>)

    def _toggle_review_check(self, tree_id: str):
        cur = self._review_checked.get(tree_id, False)
        new_state = not cur
        self._review_checked[tree_id] = new_state
        vals = list(self.review_tree.item(tree_id, "values"))
        vals[0] = "☑" if new_state else "☐"
        self.review_tree.item(tree_id, values=vals,
                              tags=("checked",) if new_state else ())
        self._update_review_check_count()

    def _review_check_all(self):
        for item in self.review_tree.get_children():
            if not self._review_checked.get(item, False):
                self._toggle_review_check(item)

    def _review_uncheck_all(self):
        for item in self.review_tree.get_children():
            if self._review_checked.get(item, False):
                self._toggle_review_check(item)

    def _review_invert_check(self):
        for item in self.review_tree.get_children():
            self._toggle_review_check(item)

    def _update_review_check_count(self):
        n = sum(1 for v in self._review_checked.values() if v)
        if hasattr(self, "review_check_count_var"):
            self.review_check_count_var.set("已勾 " + str(n) + " 条")

    def _select_review_index(self, index: int):
        children = list(self.review_tree.get_children())
        if not children:
            return False
        index = max(0, min(len(children) - 1, int(index)))
        target = children[index]
        self.review_tree.selection_set(target)
        self.review_tree.focus(target)
        self.review_tree.see(target)
        self._on_review_tree_select()
        return True

    def _move_review_selection(self, delta: int):
        children = list(self.review_tree.get_children())
        if not children:
            return
        selection = self.review_tree.selection()
        current = (children.index(selection[0])
                   if selection and selection[0] in children else 0)
        self._select_review_index(current + delta)

    def _review_tab_is_active(self):
        try:
            return self.notebook.select() == str(self.tab_review)
        except tk.TclError:
            return False

    def _review_accept_shortcut(self, event=None):
        if self._review_tab_is_active():
            self._accept_and_next()
            return "break"

    def _review_navigation_shortcut(self, delta: int):
        if self._review_tab_is_active():
            self._move_review_selection(delta)
            return "break"

    def _on_review_tree_select(self, event=None):
        """选中行(非勾选)时刷新右侧详情"""
        sel = self.review_tree.selection()
        if not sel:
            return
        tree_id = sel[0]
        item = next((it for it in self._review_items
                     if it["tree_id"] == tree_id), None)
        if not item:
            return
        try:
            with open(item["cache_file"], encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            messagebox.showerror("读取失败", str(e))
            return
        self._render_detail(data, item)

    def _get_selected_review(self):
        """单条操作用:返回当前 tree 选中行对应的 item"""
        sel = self.review_tree.selection()
        if not sel:
            messagebox.showwarning("提示", "请先在左侧选中一条")
            return None
        tree_id = sel[0]
        return next((it for it in self._review_items
                     if it["tree_id"] == tree_id), None)

    def _get_checked_reviews(self) -> list:
        """批量操作用:返回所有打勾的 item 列表"""
        return [it for it in self._review_items
                if self._review_checked.get(it["tree_id"], False)]

    def _batch_accept_checked(self):
        checked = self._get_checked_reviews()
        if not checked:
            messagebox.showwarning("提示",
                                   "请先勾选要批量收录的条目")
            return
        if not messagebox.askyesno(
                "确认批量收录",
                "将批量收录 " + str(len(checked)) + " 条记录入库,继续?"):
            return

        project_dir = self.settings.get("project_dir", "")
        ok_n = 0
        fail_n = 0
        for item in checked:
            paper_dir = (Path(project_dir) / "extract_cache"
                         / item["paper_id"])
            try:
                for cf in paper_dir.glob(item["schema_ver"] + "_*.json"):
                    lang = cf.stem.split("_")[-1]
                    with open(cf, encoding="utf-8") as f:
                        data = json.load(f)
                    insert_paper(self.db_path,
                                 item["paper_id"] + "__" + lang,
                                 data, item["schema_ver"],
                                 data.get("_model", ""), lang)
                ok_n += 1
            except Exception as e:
                self._log("✗ 收录失败 " + item["paper_id"]
                          + ": " + str(e)[:120])
                fail_n += 1

        messagebox.showinfo(
            "批量收录完成",
            "成功 " + str(ok_n) + " 条" + chr(10)
            + "失败 " + str(fail_n) + " 条")
        self._refresh_review_tab()
        self._refresh_manage_tab()

    def _batch_reject_checked(self):
        checked = self._get_checked_reviews()
        if not checked:
            messagebox.showwarning("提示",
                                   "请先勾选要批量丢弃的条目")
            return
        if not messagebox.askyesno(
                "确认批量丢弃",
                "将批量丢弃 " + str(len(checked)) + " 条记录,"
                "缓存文件会被删除,确认继续?"):
            return

        ok_n = 0
        fail_n = 0
        for item in checked:
            try:
                with open(item["cache_file"], encoding="utf-8") as f:
                    data = json.load(f)
                reject_paper(self.db_path, item["paper_id"],
                             data, "manual_batch_reject")
                Path(item["cache_file"]).unlink(missing_ok=True)
                ok_n += 1
            except Exception as e:
                self._log("✗ 丢弃失败 " + item["paper_id"]
                          + ": " + str(e)[:120])
                fail_n += 1

        messagebox.showinfo(
            "批量丢弃完成",
            "成功 " + str(ok_n) + " 条" + chr(10)
            + "失败 " + str(fail_n) + " 条")
        self._refresh_review_tab()

    def _render_detail(self, data: dict, item: dict):
        """右侧详情:顶部显示文献身份字段区块,然后是 schema 字段"""
        from schema_loader import IDENTITY_FIELDS as _IDENTITY_FIELDS_ZH
        from schema_loader import IDENTITY_KEYS as _IDENTITY_KEYS

        d = self.review_detail
        self._review_current_doi = ""
        d._identity_copy_values = {}
        d._identity_copy_all_text = ""
        d._identity_doi = ""
        d.configure(state=tk.NORMAL)
        d.delete("1.0", tk.END)
        d.insert(tk.END, "📄 " + item["paper_id"] + chr(10), "title")
        d.insert(tk.END,
                 "版本: " + item["schema_ver"]
                 + "  语言: " + item["lang"]
                 + "  模型: " + str(data.get("_model", "N/A"))
                 + "  tokens: " + str(data.get("_tokens_used", "N/A"))
                 + chr(10) + chr(10))

        # ── 文献身份信息区块(硬编码字段)──
        d.insert(tk.END, "═══ 文献身份信息 ═══" + chr(10), "section")
        d.insert(tk.END,
                 "单击字段值复制；按住鼠标拖动可任意选择文字。" +
                 chr(10), "evidence")
        label_key = "label_" + item["lang"]
        identity_lines = ["文献ID: " + item["paper_id"]]
        for f in _IDENTITY_FIELDS_ZH:
            key = f["key"]
            label = f.get(label_key, key)
            val = data.get(key, "N/A")
            if isinstance(val, list):
                parts = []
                for it in val:
                    if isinstance(it, dict):
                        parts.append(str(it.get("value", "")))
                    else:
                        parts.append(str(it))
                val = "; ".join(p for p in parts if p) or "N/A"
            else:
                val = str(val) if val not in (None, "") else "N/A"
            d.insert(
                tk.END, "  ▸ " + label + ": ",
                "identity_label")
            tag = "na" if val in ("N/A", "None", "") else "identity_block"
            if key == "DOI" and tag != "na":
                self._review_current_doi = val
                d._identity_doi = val
                self._insert_copyable_identity_value(
                    d, key, label, val,
                    ("identity_block", "doi_link"))
            else:
                self._insert_copyable_identity_value(
                    d, key, label, val, tag)
            identity_lines.append(label + ": " + val)
        d._identity_copy_all_text = chr(10).join(identity_lines)
        d.insert(tk.END, chr(10))

        # ── Schema 字段(已剔除身份字段,避免重复)──
        try:
            schema = load_schema(item["schema_ver"])
            fields = schema.get("fields", [])
        except Exception:
            fields = []

        d.insert(tk.END, "═══ 抽取字段 ═══" + chr(10), "section")
        for field in fields:
            key = field["key"]
            # 身份字段已在上方显示,这里跳过
            if key in _IDENTITY_KEYS:
                continue
            label = field.get(label_key, key)
            value = data.get(key, "N/A")
            d.insert(tk.END, "▸ " + label + " (" + key + ")"
                     + chr(10), "field")
            if isinstance(value, list):
                for item_val in value:
                    if isinstance(item_val, dict):
                        v = item_val.get("value", "")
                        src = item_val.get("source", "")
                        conf = item_val.get("confidence", "")
                        ev = item_val.get("evidence", "")
                        flag = "★[AI自创] " if src == "ai_generated" else ""
                        tag = "ai_tag" if src == "ai_generated" else None
                        d.insert(tk.END,
                                 "   • " + flag + str(v)
                                 + " [" + str(conf) + "]"
                                 + chr(10), tag)
                        if ev:
                            d.insert(tk.END,
                                     "     证据: " + str(ev) + chr(10),
                                     "evidence")
                    else:
                        d.insert(tk.END,
                                 "   • " + str(item_val) + chr(10))
            else:
                v = str(value) if value is not None else "N/A"
                tag = "na" if v in ("N/A", "", "None") else None
                d.insert(tk.END, "   " + v + chr(10), tag)
            d.insert(tk.END, chr(10))
        # 保持 NORMAL 以允许鼠标选中和 Ctrl+C；键盘编辑由绑定拦截。
        d.configure(state=tk.NORMAL)

    def _copy_to_clipboard(self, text, description="内容"):
        value = str(text or "")
        if not value:
            return False
        self.clipboard_clear()
        self.clipboard_append(value)
        self.update_idletasks()
        self.status_var.set("已复制" + description)
        self.after(
            3000,
            lambda: self.status_var.set("就绪")
            if self.status_var.get() == "已复制" + description else None)
        return True

    def _copy_selected_text(self, widget):
        try:
            selected = widget.get(tk.SEL_FIRST, tk.SEL_LAST)
        except tk.TclError:
            return False
        return self._copy_to_clipboard(selected, "选中文字")

    def _readonly_copy_shortcut(self, event):
        self._copy_selected_text(event.widget)
        return "break"

    @staticmethod
    def _readonly_select_all(event):
        event.widget.tag_add(tk.SEL, "1.0", "end-1c")
        event.widget.tag_raise(tk.SEL)
        event.widget.mark_set(tk.INSERT, "1.0")
        event.widget.see("1.0")
        return "break"

    @staticmethod
    def _raise_text_selection(event):
        """Ensure selection paint stays above coloured content tags."""
        try:
            event.widget.tag_raise(tk.SEL)
        except tk.TclError:
            pass

    def _install_readonly_copy_support(self, widget):
        # Preserve the visible selection when focus moves to a copy command.
        widget.configure(exportselection=False)
        widget.bind("<Control-c>", self._readonly_copy_shortcut)
        widget.bind("<Control-C>", self._readonly_copy_shortcut)
        widget.bind("<Control-a>", self._readonly_select_all)
        widget.bind("<Control-A>", self._readonly_select_all)
        widget.bind("<ButtonPress-1>", self._raise_text_selection, add="+")
        widget.tag_raise(tk.SEL)
        widget.bind(
            "<Button-3>",
            lambda event: self._show_text_copy_menu(event, widget))
        if not hasattr(widget, "_identity_copy_values"):
            widget._identity_copy_values = {}
        widget._identity_copy_all_text = ""
        widget._identity_doi = ""

    def _insert_copyable_identity_value(self, widget, key, label,
                                        value, base_tag):
        safe_key = re.sub(r"[^A-Za-z0-9_]", "_", str(key))
        copy_tag = "copy_identity_" + safe_key
        widget._identity_copy_values[copy_tag] = (str(label), str(value))
        if isinstance(base_tag, (tuple, list)):
            tags = tuple(base_tag) + (copy_tag,)
        else:
            tags = (base_tag, copy_tag) if base_tag else (copy_tag,)
        widget.insert(tk.END, str(value) + chr(10), tags)
        widget.tag_bind(
            copy_tag, "<ButtonRelease-1>",
            lambda event, tag=copy_tag:
            self._copy_identity_tag_if_not_selecting(event, tag))
        widget.tag_bind(
            copy_tag, "<Enter>",
            lambda event: event.widget.configure(cursor="hand2"))
        widget.tag_bind(
            copy_tag, "<Leave>",
            lambda event: event.widget.configure(cursor="xterm"))

    def _copy_identity_tag_if_not_selecting(self, event, tag):
        try:
            if event.widget.tag_ranges(tk.SEL):
                return None
        except tk.TclError:
            return None
        field = event.widget._identity_copy_values.get(tag)
        if field:
            label, value = field
            self._copy_to_clipboard(value, "字段：" + label)
        return None

    def _identity_tag_at_pointer(self, widget, x, y):
        try:
            tags = widget.tag_names("@" + str(x) + "," + str(y))
        except tk.TclError:
            return None
        return next((tag for tag in tags
                     if str(tag).startswith("copy_identity_")), None)

    def _show_text_copy_menu(self, event, widget):
        copy_tag = self._identity_tag_at_pointer(
            widget, event.x, event.y)
        field = widget._identity_copy_values.get(copy_tag) if copy_tag else None
        try:
            has_selection = bool(widget.tag_ranges(tk.SEL))
        except tk.TclError:
            has_selection = False

        previous_menu = getattr(widget, "_copy_context_menu", None)
        if previous_menu is not None:
            try:
                previous_menu.destroy()
            except tk.TclError:
                pass
        menu = tk.Menu(widget, tearoff=False)
        widget._copy_context_menu = menu
        menu.add_command(
            label="复制选中文字",
            state=tk.NORMAL if has_selection else tk.DISABLED,
            command=lambda: self._copy_selected_text(widget))
        menu.add_command(
            label=("复制当前字段：" + field[0]
                   if field else "复制当前字段"),
            state=tk.NORMAL if field else tk.DISABLED,
            command=(lambda selected=field:
                     self._copy_to_clipboard(
                         selected[1], "字段：" + selected[0])
                     if selected else None))
        menu.add_separator()
        identity_text = str(getattr(
            widget, "_identity_copy_all_text", "") or "")
        menu.add_command(
            label="复制全部身份信息",
            state=tk.NORMAL if identity_text else tk.DISABLED,
            command=lambda: self._copy_to_clipboard(
                identity_text, "全部身份信息"))
        doi = str(getattr(widget, "_identity_doi", "") or "")
        menu.add_command(
            label="复制 DOI",
            state=(tk.NORMAL if doi and doi not in ("N/A", "None")
                   else tk.DISABLED),
            command=lambda: self._copy_to_clipboard(doi, " DOI"))
        menu.add_separator()
        menu.add_command(
            label="全选",
            command=lambda: (
                widget.tag_add(tk.SEL, "1.0", "end-1c"),
                widget.mark_set(tk.INSERT, "1.0")))
        self.theme_manager.apply_widget(menu)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    @staticmethod
    def _readonly_text_key(event):
        """让 Text 保持可选中复制，但拦截输入、删除和粘贴。"""
        key = str(event.keysym).lower()
        control = bool(event.state & 0x4)
        if control and key in ("c", "a"):
            return None
        if key in (
                "left", "right", "up", "down",
                "home", "end", "prior", "next"):
            return None
        return "break"

    def _copy_review_doi(self):
        doi = str(getattr(self, "_review_current_doi", "") or "").strip()
        if not doi or doi in ("N/A", "None"):
            messagebox.showinfo("提示", "当前条目没有可复制的 DOI")
            return "break"
        self.clipboard_clear()
        self.clipboard_append(doi)
        self.status_var.set("已复制 DOI：" + doi)
        self.after(
            3000,
            lambda: self.status_var.set("就绪")
            if self.status_var.get().startswith("已复制 DOI：") else None)
        return "break"

    def _accept_and_next(self):
        children = list(self.review_tree.get_children())
        selection = self.review_tree.selection()
        next_index = (children.index(selection[0])
                      if selection and selection[0] in children else 0)
        self._accept_paper(show_message=False, next_index=next_index)

    def _accept_paper(self, show_message: bool = True,
                      next_index: int = 0):
        item = self._get_selected_review()
        if not item:
            return False
        project_dir = self.settings.get("project_dir", "")
        paper_dir = Path(project_dir) / "extract_cache" / item["paper_id"]
        count = 0
        try:
            for cache_file in paper_dir.glob(item["schema_ver"] + "_*.json"):
                lang = cache_file.stem.split("_")[-1]
                with open(cache_file, encoding="utf-8") as f:
                    data = json.load(f)
                insert_paper(self.db_path,
                             item["paper_id"] + "__" + lang,
                             data, item["schema_ver"],
                             data.get("_model", ""), lang)
                count += 1
        except Exception as exc:
            messagebox.showerror(
                "收录失败", type(exc).__name__ + ": " + str(exc)[:200])
            return False
        if show_message:
            messagebox.showinfo(
                "成功",
                "已入库 " + str(count) + " 条(" + item["paper_id"] + ")")
        else:
            self.review_status_label._theme_role = "success"
            self.theme_manager.apply_widget(self.review_status_label)
            self.review_status_var.set(
                "✓ 已收录 " + item["paper_id"] + "，已定位下一篇"
            )
            self.after(2500, lambda: self.review_status_var.set(""))
        self._refresh_review_tab()
        if not show_message:
            self._select_review_index(next_item_index(
                next_index, len(self.review_tree.get_children())))
        self._refresh_manage_tab()
        return True

    def _reject_paper(self):
        item = self._get_selected_review()
        if not item:
            return
        if not messagebox.askyesno("确认",
                                   "丢弃 " + item["paper_id"] + "?"):
            return
        with open(item["cache_file"], encoding="utf-8") as f:
            data = json.load(f)
        reject_paper(self.db_path, item["paper_id"], data, "manual_reject")
        Path(item["cache_file"]).unlink()
        messagebox.showinfo("已丢弃", item["paper_id"])
        self._refresh_review_tab()

    def _rerun_paper(self):
        item = self._get_selected_review()
        if not item:
            return
        if not messagebox.askyesno(
                "确认",
                "立即重跑 " + item["paper_id"] + " ?" + chr(10) +
                "将删除当前 AI 缓存并重新调用 AI（本窗口完成，无需回抽取页）"):
            return

        # 删旧 AI 缓存
        try:
            Path(item["cache_file"]).unlink()
        except Exception:
            pass

        # 显示"正在重跑"状态
        self.review_status_var.set("⏳ 正在重跑 " + item["paper_id"] + " ...")
        self.review_status_label._theme_role = "accent"
        self.theme_manager.apply_widget(self.review_status_label)
        self.status_var.set("正在重跑: " + item["paper_id"])

        def _worker():
            ok = self._do_rerun_one(
                item["paper_id"], item["schema_ver"], item["lang"])

            def _on_done():
                self._refresh_review_tab()
                self._refresh_manage_tab()
                if ok:
                    self.review_status_var.set("✓ 重跑完成: " + item["paper_id"])
                    self.review_status_label._theme_role = "success"
                    self.theme_manager.apply_widget(self.review_status_label)
                    self.status_var.set("重跑完成")
                    self.after(2000, lambda: self.review_status_var.set(""))
                else:
                    self.review_status_var.set("✗ 重跑失败，请查看抽取页日志")
                    self.review_status_label._theme_role = "danger"
                    self.theme_manager.apply_widget(self.review_status_label)
                    self.status_var.set("重跑失败")
                    self.after(5000, lambda: self.review_status_var.set(""))

            self.after(0, _on_done)

        threading.Thread(target=_worker, daemon=True).start()

    def _open_cache_dir(self):
        item = self._get_selected_review()
        if not item:
            return
        os.startfile(Path(item["cache_file"]).parent)

    # ════════════════════════════════════════════════════
    # 数据管理页
    # ════════════════════════════════════════════════════

    def _build_manage_tab(self):
        frame = self.tab_manage

        top = ttk.Frame(frame)
        top.pack(fill=tk.X, padx=8, pady=4)

        ttk.Label(top, text="当前数据库:",
                  font=("微软雅黑", 9, "bold")
                  ).pack(side=tk.LEFT, padx=(4, 0))
        self.db_name_var = tk.StringVar()
        self.db_combo = ttk.Combobox(top, textvariable=self.db_name_var,
                                     width=20, state="readonly")
        self.db_combo.pack(side=tk.LEFT, padx=4)
        self.db_combo.bind("<<ComboboxSelected>>", self._on_db_change)
        ttk.Button(top, text="➕ 新建数据库",
                   command=lambda: self._create_new_db()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(top, text="🔄 刷新",
                   command=lambda: self._refresh_manage_tab()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Separator(top, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=6)
        ttk.Button(top, text="📤 导出Excel(中)",
                   command=lambda: self._export_excel("zh")
                   ).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="📤 导出Excel(英)",
                   command=lambda: self._export_excel("en")
                   ).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="📝 导出Markdown(中)",
                   command=lambda: self._export_markdown("zh")
                   ).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="📝 导出Markdown(英)",
                   command=lambda: self._export_markdown("en")
                   ).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="📦 备份zip",
                   command=lambda: self._export_zip()).pack(side=tk.LEFT, padx=2)
        ttk.Button(top, text="📥 导入zip",
                   command=lambda: self._import_zip()).pack(side=tk.LEFT, padx=2)



        self.stats_label = tk.Label(frame, text="", fg="#1F4E79",
                                    font=("微软雅黑", 10, "bold"))
        self.stats_label._theme_role = "accent"
        self.stats_label.pack(pady=2)

        filter_frame = ttk.Frame(frame)
        filter_frame.pack(fill=tk.X, padx=8, pady=2)
        ttk.Label(filter_frame, text="筛选:").pack(side=tk.LEFT)
        self.filter_var = tk.StringVar()
        self.filter_var.trace_add(
            "write", lambda *a: self._schedule_filter())
        ttk.Entry(filter_frame, textvariable=self.filter_var,
                  width=30).pack(side=tk.LEFT, padx=4)
        ttk.Label(filter_frame, text="字段:").pack(side=tk.LEFT, padx=(8, 2))
        self.filter_field_var = tk.StringVar(value="全部字段")
        self.filter_field_combo = ttk.Combobox(
            filter_frame, textvariable=self.filter_field_var,
            values=[
                "全部字段", "文献ID", "DOI", "标题", "作者", "期刊",
                "年份", "关键词", "抽取内容", "语言", "版本", "模型",
            ],
            width=10, state="readonly")
        self.filter_field_combo.pack(side=tk.LEFT)
        self.filter_field_combo.bind(
            "<<ComboboxSelected>>", lambda event: self._schedule_filter())
        ttk.Button(filter_frame, text="清除",
                   command=lambda: self.filter_var.set("")
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Separator(filter_frame, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=8)
        self.ask_var = tk.StringVar()
        ask_entry = ttk.Entry(filter_frame, textvariable=self.ask_var,
                              width=36)
        ask_entry.pack(side=tk.LEFT, padx=4)
        ask_entry.bind("<Return>", lambda e: self._ask_ai_find_papers())
        ttk.Label(filter_frame, text="返回:").pack(side=tk.LEFT, padx=(12, 2))
        self.max_hits_var = tk.StringVar(value="20")
        max_hits_combo = ttk.Combobox(
            filter_frame, textvariable=self.max_hits_var,
            values=["10", "20", "50", "100", "全部"],
            width=6)   # ← 注意:不加 state="readonly",允许手动输入任意数字
        max_hits_combo.pack(side=tk.LEFT)
        ttk.Label(filter_frame, text="篇",
                  foreground="#888").pack(side=tk.LEFT)
        self._ai_button(filter_frame, text="🔎 问 AI 找论文",
                       command=lambda: self._ask_ai_find_papers()
                       ).pack(side=tk.LEFT, padx=2)


        sel_frame = ttk.Frame(frame)
        sel_frame.pack(fill=tk.X, padx=8, pady=2)
        ttk.Button(sel_frame, text="全选",
                   command=self._select_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(sel_frame, text="全不选",
                   command=self._deselect_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(sel_frame, text="反选",
                   command=self._invert_select).pack(side=tk.LEFT, padx=2)
        self.sel_count_label = tk.Label(sel_frame, text="已选 0 条",
                                        foreground="#0070C0",
                                        font=("微软雅黑", 9))
        self.sel_count_label.pack(side=tk.LEFT, padx=8)

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
        self.manage_tree_frame = tree_frame

        cols = ("seq", "paper_id", "schema_ver", "lang", "model", "added_at")
        self.manage_tree = ttk.Treeview(tree_frame, columns=cols,
                                        show="headings", height=12,
                                        selectmode=tk.EXTENDED)
        self.manage_tree.heading("seq",        text="#")
        self.manage_tree.heading("paper_id",   text="文献ID")
        self.manage_tree.heading("schema_ver", text="提示词版本")
        self.manage_tree.heading("lang",       text="语言")
        self.manage_tree.heading("model",      text="模型")
        self.manage_tree.heading("added_at",   text="入库时间")
        self.manage_tree.column("seq",        width=50, anchor=tk.CENTER,
                                stretch=False)
        self.manage_tree.column("paper_id",   width=280)
        self.manage_tree.column("schema_ver", width=130)
        self.manage_tree.column("lang",       width=55, anchor=tk.CENTER)
        self.manage_tree.column("model",      width=130)
        self.manage_tree.column("added_at",   width=170)
        # 同一篇 __zh/__en 共享序号时,用底色区分奇偶序号,视觉上更清晰
        self.manage_tree.tag_configure("seq_odd",  background="#FFFFFF")
        self.manage_tree.tag_configure("seq_even", background="#F5F9FF")
        self.manage_tree.bind("<<TreeviewSelect>>", self._update_sel_count)
        self.manage_tree.bind("<Double-1>",
                              self._on_manage_tree_double_click)

        scrollbar = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL,
                                  command=self.manage_tree.yview)
        self.manage_tree.configure(yscrollcommand=scrollbar.set)
        self.manage_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.LEFT, fill=tk.Y)

        page_bar = ttk.Frame(frame)
        self.manage_page_bar = page_bar
        page_bar.pack(side=tk.BOTTOM, fill=tk.X, padx=8, pady=(0, 2),
                      before=tree_frame)
        ttk.Label(page_bar, text="每页:").pack(side=tk.LEFT, padx=(2, 2))
        self.manage_page_size_var = tk.StringVar(
            value=str(self._manage_page_size))
        page_size_combo = ttk.Combobox(
            page_bar, textvariable=self.manage_page_size_var,
            values=[str(value) for value in PAGE_SIZE_OPTIONS],
            width=6, state="readonly")
        page_size_combo.pack(side=tk.LEFT)
        page_size_combo.bind(
            "<<ComboboxSelected>>", self._on_manage_page_size_change)
        self.manage_first_page_button = ttk.Button(
            page_bar, text="首页", command=lambda: self._go_manage_page(1))
        self.manage_first_page_button.pack(side=tk.LEFT, padx=(12, 2))
        self.manage_prev_page_button = ttk.Button(
            page_bar, text="上一页", command=lambda: self._go_manage_page(-1, True))
        self.manage_prev_page_button.pack(side=tk.LEFT, padx=2)
        self.manage_next_page_button = ttk.Button(
            page_bar, text="下一页", command=lambda: self._go_manage_page(1, True))
        self.manage_next_page_button.pack(side=tk.LEFT, padx=2)
        self.manage_last_page_button = ttk.Button(
            page_bar, text="末页", command=self._go_manage_last_page)
        self.manage_last_page_button.pack(side=tk.LEFT, padx=2)
        self.manage_page_info_var = tk.StringVar(value="第 1 / 1 页")
        ttk.Label(page_bar, textvariable=self.manage_page_info_var).pack(
            side=tk.LEFT, padx=10)
        page_hint = tk.Label(
            page_bar, text="选择状态可跨页保留", anchor=tk.E)
        page_hint._theme_role = "muted"
        page_hint.pack(side=tk.RIGHT, padx=4)

        bottom = ttk.Frame(frame)
        self.manage_action_bar = bottom
        # Reserve the action bar before the expandable table gets its space.
        # This keeps the buttons visible on shorter screens and larger scaling.
        bottom.pack(side=tk.BOTTOM, fill=tk.X, padx=8, pady=(2, 6),
                    before=tree_frame)
        ttk.Button(bottom, text="❌ 删除选中",
                   command=lambda: self._delete_selected_db()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(bottom, text="👁 查看详情",
                   command=lambda: self._view_db_detail()
                   ).pack(side=tk.LEFT, padx=4)
        self._ai_button(bottom, text="↺ 重抽选中",
                        command=lambda: self._rerun_selected_db()
                        ).pack(side=tk.LEFT, padx=4)

        self._ai_button(bottom, text="💬 单篇问 AI",
                        command=lambda: self._open_single_chat()
                        ).pack(side=tk.LEFT, padx=4)

        self._ai_button(bottom, text="💬 打开AI对话",
                        command=lambda: self._open_chat_window()
                        ).pack(side=tk.LEFT, padx=12)
        ttk.Button(bottom, text="📚 引用工具",
                   command=lambda: self._open_citation_builder()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(bottom, text="🔧 重命名ID",
                   command=lambda: self._rename_paper_id_dialog()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(bottom, text="🧹 扫描重复",
                   command=lambda: self._scan_database_duplicates()
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(bottom, text="🔍 查找替换",
                   command=lambda: self._find_replace_dialog()
                  ).pack(side=tk.LEFT, padx=4)


        self.after(100, self._refresh_db_list)


    # ── 数据库切换 ─────────────────────────────────

    def _refresh_db_list(self):
        project_dir = self.settings.get("project_dir", "")
        dbs = list_databases(project_dir)
        if not dbs:
            create_database(project_dir, "main")
            dbs = ["main"]
        self.db_combo["values"] = dbs
        saved = self.settings.get("current_db", "main")
        if saved in dbs:
            self.db_name_var.set(saved)
        else:
            self.db_name_var.set(dbs[0])
        self._apply_db_change()
        self._refresh_manage_tab()

    def _apply_db_change(self):
        project_dir = self.settings.get("project_dir", "")
        db_name = self.db_name_var.get()
        self.db_path = get_db_path_by_name(project_dir, db_name)
        init_db(self.db_path)
        self.settings["current_db"] = db_name
        save_settings(self.settings)

    def _on_db_change(self, event=None):
        self._manage_selected_ids.clear()
        self._manage_page = 1
        self._apply_db_change()
        self._refresh_manage_tab()
        self._chat_messages = []

    def _create_new_db(self):
        name = simpledialog.askstring(
            "新建数据库",
            "输入数据库名称(字母数字下划线):")
        if not name:
            return
        project_dir = self.settings.get("project_dir", "")
        create_database(project_dir, name)
        self._refresh_db_list()
        messagebox.showinfo("成功", "已创建数据库:" + name)

    # ── 选择与筛选 ─────────────────────────────────

    def _select_all(self):
        self._capture_manage_page_selection()
        self._manage_selected_ids.update(
            str(row[1]) for row in self._manage_filtered_rows)
        self._render_manage_page()

    def _deselect_all(self):
        self._manage_selected_ids.clear()
        self._render_manage_page()

    def _invert_select(self):
        self._capture_manage_page_selection()
        filtered_ids = {str(row[1]) for row in self._manage_filtered_rows}
        self._manage_selected_ids.symmetric_difference_update(filtered_ids)
        self._render_manage_page()

    def _capture_manage_page_selection(self):
        """Merge the visible page's Treeview selection into cross-page state."""
        if self._manage_selection_sync or not hasattr(self, "manage_tree"):
            return
        page_ids = set()
        selected_ids = set()
        selected_items = set(self.manage_tree.selection())
        for item in self.manage_tree.get_children():
            values = self.manage_tree.item(item, "values")
            if len(values) < 2:
                continue
            paper_id = str(values[1])
            page_ids.add(paper_id)
            if item in selected_items:
                selected_ids.add(paper_id)
        self._manage_selected_ids.difference_update(page_ids)
        self._manage_selected_ids.update(selected_ids)

    def _selected_manage_ids(self) -> list:
        """Return selected record IDs in stable database display order."""
        self._capture_manage_page_selection()
        return [
            str(row[1]) for row in self._all_papers_cache
            if str(row[1]) in self._manage_selected_ids
        ]

    def _update_sel_count(self, event=None):
        self._capture_manage_page_selection()
        n = len(self._manage_selected_ids)
        if hasattr(self, "sel_count_label"):
            suffix = "（跨页）" if n > len(self.manage_tree.selection()) else ""
            self.sel_count_label.configure(
                text="已选 " + str(n) + " 条" + suffix)

    def _render_manage_page(self):
        if not hasattr(self, "manage_tree"):
            return
        page_rows, safe_page, total_pages = paginate_rows(
            self._manage_filtered_rows, self._manage_page,
            self._manage_page_size)
        self._manage_page = safe_page
        self._manage_selection_sync = True
        try:
            for item in self.manage_tree.get_children():
                self.manage_tree.delete(item)
            selected_items = []
            for row in page_rows:
                seq = row[0]
                tag = "seq_even" if (seq % 2 == 0) else "seq_odd"
                item = self.manage_tree.insert(
                    "", tk.END, values=row, tags=(tag,))
                if str(row[1]) in self._manage_selected_ids:
                    selected_items.append(item)
            if selected_items:
                self.manage_tree.selection_set(selected_items)
        finally:
            self._manage_selection_sync = False

        total_rows = len(self._manage_filtered_rows)
        start = ((safe_page - 1) * self._manage_page_size + 1
                 if total_rows else 0)
        end = min(safe_page * self._manage_page_size, total_rows)
        if hasattr(self, "manage_page_info_var"):
            self.manage_page_info_var.set(
                "第 " + str(safe_page) + " / " + str(total_pages) +
                " 页 · " + str(start) + "-" + str(end) +
                " / " + str(total_rows) + " 条")
            first_state = tk.DISABLED if safe_page <= 1 else tk.NORMAL
            last_state = (tk.DISABLED if safe_page >= total_pages
                          else tk.NORMAL)
            self.manage_first_page_button.configure(state=first_state)
            self.manage_prev_page_button.configure(state=first_state)
            self.manage_next_page_button.configure(state=last_state)
            self.manage_last_page_button.configure(state=last_state)
        self._update_sel_count()

    def _go_manage_page(self, target: int, relative: bool = False):
        self._capture_manage_page_selection()
        self._manage_page = (
            self._manage_page + target if relative else target)
        self._render_manage_page()

    def _go_manage_last_page(self):
        _, _, total_pages = paginate_rows(
            self._manage_filtered_rows, 1, self._manage_page_size)
        self._go_manage_page(total_pages)

    def _on_manage_page_size_change(self, event=None):
        self._capture_manage_page_selection()
        self._manage_page_size = normalize_page_size(
            self.manage_page_size_var.get())
        self.manage_page_size_var.set(str(self._manage_page_size))
        self.settings["manage_page_size"] = self._manage_page_size
        save_settings(self.settings)
        self._manage_page = 1
        self._render_manage_page()

    def _select_manage_base_ids(self, base_ids, replace=True) -> set:
        """Select matching bases across pages and reveal the first match."""
        targets = {str(base) for base in base_ids}
        matching_ids = [
            str(row[1]) for row in self._manage_filtered_rows
            if str(row[1]).rsplit("__", 1)[0] in targets
        ]
        if replace:
            self._manage_selected_ids.clear()
        self._manage_selected_ids.update(matching_ids)
        if matching_ids:
            first_id = matching_ids[0]
            first_index = next(
                index for index, row in enumerate(self._manage_filtered_rows)
                if str(row[1]) == first_id)
            self._manage_page = first_index // self._manage_page_size + 1
        self._render_manage_page()
        return {
            paper_id.rsplit("__", 1)[0] for paper_id in matching_ids
        }

    def _on_manage_tree_double_click(self, event):
        """双击文献 ID 列，直接打开单篇重命名窗口。"""
        if self.manage_tree.identify("region", event.x, event.y) != "cell":
            return
        if self.manage_tree.identify_column(event.x) != "#2":
            return
        item = self.manage_tree.identify_row(event.y)
        if not item:
            return
        values = self.manage_tree.item(item, "values")
        self._manage_selected_ids = {str(values[1])}
        self._render_manage_page()
        self._rename_paper_id_dialog()

    def _apply_filter(self, reset_page=True):
        if not hasattr(self, "manage_tree"):
            return
        keyword = _normalize_search_text(self.filter_var.get())
        field = self.filter_field_var.get()
        tokens = [token for token in keyword.split(" ") if token]

        matching_bases = set()
        for base_id, documents in self._manage_base_search_index.items():
            document = documents.get(
                field, documents.get("全部字段", ""))
            if not tokens or all(token in document for token in tokens):
                matching_bases.add(base_id)

        self._manage_filtered_rows = []
        for row in self._all_papers_cache:
            base = str(row[1]).rsplit("__", 1)[0]
            if base in matching_bases:
                self._manage_filtered_rows.append(row)
        if reset_page:
            self._manage_page = 1
        self._render_manage_page()

    def _schedule_filter(self, delay_ms=250):
        """输入防抖：连续键入时只执行最后一次筛选。"""
        if not hasattr(self, "manage_tree"):
            return
        if self._filter_after_id is not None:
            try:
                self.after_cancel(self._filter_after_id)
            except Exception:
                pass
        self._filter_after_id = self.after(
            delay_ms, self._run_scheduled_filter)

    def _run_scheduled_filter(self):
        self._filter_after_id = None
        self._apply_filter()

    def _refresh_manage_tab(self):
        self._capture_manage_page_selection()

        papers = get_all_papers(self.db_path)

        # ── 按 base_id 去重分配序号:__zh 和 __en 共享同一序号 ──
        # 排序规则:先按 base_id 字母序,保证序号稳定,新增文献插入到合适位置
        base_to_seq = {}
        next_seq = 1
        # 先排序:按 base_id 升序,同一 base 下 zh 在前 en 在后
        sorted_papers = sorted(
            papers,
            key=lambda p: (
                p.get("_paper_id", "").rsplit("__", 1)[0],
                p.get("_lang", ""),
            )
        )

        # 分配序号
        for p in sorted_papers:
            base = p.get("_paper_id", "").rsplit("__", 1)[0]
            if base not in base_to_seq:
                base_to_seq[base] = next_seq
                next_seq += 1

        # 把序号映射保存到 self,供 AI 对话上下文复用
        self._paper_seq_map = base_to_seq

        # ── 写入 Tree ──
        self._all_papers_cache = []
        self._manage_papers_by_id = {}
        self._manage_search_index = {}
        self._manage_base_search_index = {}
        cached_search_documents = build_cached_search_documents(
            self.db_path, sorted_papers)
        for p in sorted_papers:
            base = p.get("_paper_id", "").rsplit("__", 1)[0]
            seq = base_to_seq.get(base, 0)
            row = (
                seq,
                p.get("_paper_id", ""),
                p.get("_schema_ver", ""),
                p.get("_lang", ""),
                p.get("_model_used", ""),
                p.get("_added_at", "")[:19],
            )
            self._all_papers_cache.append(row)
            full_id = p.get("_paper_id", "")
            self._manage_papers_by_id[full_id] = p
            self._manage_search_index[full_id] = (
                cached_search_documents.get(full_id)
                or _build_paper_search_document(p))
            base_documents = self._manage_base_search_index.setdefault(
                base, {})
            for field_name, document in (
                    self._manage_search_index[full_id].items()):
                if document:
                    base_documents[field_name] = (
                        base_documents.get(field_name, "") + " " +
                        document).strip()

        valid_ids = {str(row[1]) for row in self._all_papers_cache}
        self._manage_selected_ids.intersection_update(valid_ids)
        self._apply_filter(reset_page=False)

        stats = get_stats(self.db_path)
        unique_count = len(base_to_seq)
        self.stats_label.configure(
            text="数据库:" + self.db_name_var.get()
                 + "   去重后:" + str(unique_count) + " 篇"
                 + "   总记录:" + str(stats["total"]) + " 条"
                 + "   已拒绝:" + str(stats["rejected"]) + " 条")
        self._update_sel_count()

    def _get_visible_base_ids(self) -> list:
        """返回当前筛选结果中的基础文献 ID，保持界面顺序并去重。"""
        result = []
        seen = set()
        for row in self._manage_filtered_rows:
            base = str(row[1]).rsplit("__", 1)[0]
            if base and base not in seen:
                seen.add(base)
                result.append(base)
        return result

    def _scan_database_duplicates(self):
        """扫描当前文献库，预览并安全合并跨基础 ID 的重复文献。"""
        groups = find_duplicate_groups(self.db_path)
        if not groups:
            messagebox.showinfo(
                "重复扫描",
                "当前文献库「" + self.db_name_var.get() +
                "」未发现 DOI 或身份信息重复的基础文献。")
            return

        win = tk.Toplevel(self)
        win.title("文献库重复扫描 — " + self.db_name_var.get())
        win.geometry("1080x680")
        win.grab_set()

        ttk.Label(
            win,
            text=("扫描规则：DOI 相同；或标题相同且年份/第一作者一致。"
                  "中英文记录属于同一基础文献，不会互相误判。"),
            foreground="#1F4E79",
            font=("微软雅黑", 9)
        ).pack(anchor=tk.W, padx=10, pady=(10, 4))
        ttk.Label(
            win,
            text=("选择一行作为该重复组的保留 ID；合并时保留它的非空值，"
                  "仅用其他记录填补空字段。★ 为系统推荐保留项。"),
            foreground="#666",
            font=("微软雅黑", 8)
        ).pack(anchor=tk.W, padx=10, pady=(0, 6))

        cols = ("group", "keep", "base_id", "langs",
                "complete", "doi", "title")
        tree = ttk.Treeview(
            win, columns=cols, show="headings",
            selectmode=tk.BROWSE, height=20)
        headings = {
            "group": "组",
            "keep": "推荐",
            "base_id": "基础文献 ID",
            "langs": "语言",
            "complete": "完整字段",
            "doi": "DOI",
            "title": "标题",
        }
        widths = {
            "group": 45, "keep": 55, "base_id": 220,
            "langs": 70, "complete": 75, "doi": 185, "title": 350,
        }
        for col in cols:
            tree.heading(col, text=headings[col])
            tree.column(
                col, width=widths[col],
                anchor=tk.CENTER if col in (
                    "group", "keep", "langs", "complete") else tk.W)
        tree.tag_configure("recommended", background="#E8F5E9")
        tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

        item_map = {}
        first_recommended = None
        for group_no, group in enumerate(groups, 1):
            for record in group["records"]:
                recommended = (
                    record["base_id"] == group["recommended_keep"])
                item_id = tree.insert(
                    "", tk.END,
                    values=(
                        group_no,
                        "★" if recommended else "",
                        record["base_id"],
                        "/".join(record["languages"]),
                        record["completeness"],
                        record["doi"],
                        record["title"],
                    ),
                    tags=("recommended",) if recommended else ())
                item_map[item_id] = (group, record)
                if first_recommended is None and recommended:
                    first_recommended = item_id

        if first_recommended:
            tree.selection_set(first_recommended)
            tree.see(first_recommended)

        reason_var = tk.StringVar()
        tk.Label(
            win, textvariable=reason_var, foreground="#8A5A00",
            font=("微软雅黑", 9), anchor=tk.W
        ).pack(fill=tk.X, padx=10, pady=2)

        def _update_reason(event=None):
            selected = tree.selection()
            if not selected:
                reason_var.set("")
                return
            group, _ = item_map[selected[0]]
            reason_var.set("判重依据：" + group["reason"])

        tree.bind("<<TreeviewSelect>>", _update_reason)
        _update_reason()

        btn_row = ttk.Frame(win)
        btn_row.pack(fill=tk.X, padx=10, pady=8)

        def _merge_selected_group():
            selected = tree.selection()
            if not selected:
                messagebox.showwarning(
                    "提示", "请先选择要保留的一行", parent=win)
                return
            group, record = item_map[selected[0]]
            keep_id = record["base_id"]
            remove_ids = [
                base_id for base_id in group["base_ids"]
                if base_id != keep_id]
            if not messagebox.askyesno(
                    "确认合并重复文献",
                    "当前文献库：" + self.db_name_var.get() + chr(10) +
                    "保留 ID：" + keep_id + chr(10) +
                    "合并并移除：" + ", ".join(remove_ids) + chr(10) +
                    "判重依据：" + group["reason"] + chr(10) + chr(10) +
                    "目标记录的非空字段不会被覆盖；其他记录只用于填补空字段。"
                    + chr(10) + "建议重要数据先备份。确认继续？",
                    parent=win):
                return

            result = merge_duplicate_group(
                self.db_path, keep_id, remove_ids,
                self.settings.get("project_dir", ""))
            if not result.get("ok"):
                messagebox.showerror(
                    "合并失败", result.get("error", "未知错误"),
                    parent=win)
                return

            self._refresh_manage_tab()
            win.destroy()
            messagebox.showinfo(
                "重复合并完成",
                "保留：" + keep_id + chr(10) +
                "合并记录：" + str(result["merged_rows"]) + " 条" +
                chr(10) +
                "补全空字段：" + str(result["filled_fields"]) + " 个")

        ttk.Button(
            btn_row, text="✓ 保留所选 ID 并合并本组",
            command=_merge_selected_group
        ).pack(side=tk.LEFT, padx=4)
        ttk.Button(
            btn_row, text="🔧 重命名所选 ID",
            command=lambda: _rename_from_duplicate_tree()
        ).pack(side=tk.LEFT, padx=4)
        ttk.Button(
            btn_row, text="关闭", command=win.destroy
        ).pack(side=tk.RIGHT, padx=4)

        def _rename_from_duplicate_tree():
            selected = tree.selection()
            if not selected:
                return
            _, record = item_map[selected[0]]
            target_base = record["base_id"]
            # 重复项可能不在当前页；清除筛选后由跨页选择器定位。
            self.filter_var.set("")
            self._apply_filter()
            self._select_manage_base_ids([target_base])
            win.destroy()
            self._rename_paper_id_dialog()

    # ── 删除与详情 ─────────────────────────────────

    def _delete_selected_db(self):
        selected_ids = self._selected_manage_ids()
        if not selected_ids:
            messagebox.showwarning("提示", "请先选中记录")
            return
        if not messagebox.askyesno(
                "确认",
                "删除选中的 " + str(len(selected_ids)) + " 条记录?"):
            return
        for paper_id in selected_ids:
            delete_paper(self.db_path, paper_id)
        self._manage_selected_ids.difference_update(selected_ids)
        self._refresh_manage_tab()

    def _rerun_selected_db(self):
        """
        管理页批量重抽:对选中的每条记录调用 _do_rerun_one,
        重抽完会自动覆盖入库。后台线程跑,过程中不可中断。
        """
        selected_ids = self._selected_manage_ids()
        if not selected_ids:
            messagebox.showwarning("提示", "请先选中要重抽的记录")
            return
        if not self.settings.get("api_key"):
            messagebox.showerror("错误", "请先到【设置】页填写 API Key")
            return
        if not messagebox.askyesno(
                "确认",
                "重抽选中的 " + str(len(selected_ids)) + " 条记录?" + chr(10) +
                "AI 会重新跑一次,完成后自动覆盖数据库" + chr(10) +
                "(过程中无法中断,请耐心等待)"):
            return

        # 收集任务信息(必须在主线程读 tree,因为 tkinter 控件非线程安全)
        tasks = []
        for full_id in selected_ids:
            paper = self._manage_papers_by_id.get(full_id, {})
            schema_ver = paper.get("_schema_ver", "")
            lang = paper.get("_lang", "")
            base_id = full_id.rsplit("__", 1)[0]
            tasks.append((base_id, schema_ver, lang))

        self._log("========= 批量重抽 " + str(len(tasks)) + " 条 =========")

        def _worker():
            ok_n = 0
            fail_n = 0
            for base_id, ver, lang in tasks:
                if self._do_rerun_one(base_id, ver, lang):
                    ok_n += 1
                else:
                    fail_n += 1
                # 每跑完一条刷新管理页,看时间戳变化
                self.after(0, self._refresh_manage_tab)

            self._safe_log("========= 批量重抽完成: 成功 " + str(ok_n) +
                           ",失败 " + str(fail_n) + " =========")
            self.after(0, lambda: messagebox.showinfo(
                "重抽完成",
                "成功 " + str(ok_n) + " 条" + chr(10) +
                "失败 " + str(fail_n) + " 条"))

        threading.Thread(target=_worker, daemon=True).start()

    def _view_db_detail(self):
        selected_ids = self._selected_manage_ids()
        if not selected_ids:
            return
        paper_id = selected_ids[0]
        data = get_paper(self.db_path, paper_id)
        if not data:
            messagebox.showwarning("提示", "找不到该记录")
            return

        # 从 paper_id 推断 schema_ver 和 lang
        schema_ver = data.get("_schema_ver", "")
        lang = data.get("_lang", "zh")
        base_id = paper_id.rsplit("__", 1)[0]

        # 构造一个与 _render_detail 兼容的 item 字典
        item = {
            "paper_id": base_id,
            "schema_ver": schema_ver,
            "lang": lang,
            "cache_file": "",   # 详情窗口不需要缓存路径
            "tree_id": None,
        }

        win = tk.Toplevel(self)
        win.title("详情: " + paper_id)
        win.geometry("860x700")

        # ── 顶部工具栏 ──
        toolbar = ttk.Frame(win)
        toolbar.pack(fill=tk.X, padx=8, pady=4)

        # 编辑/保存状态
        edit_state = {"editing": False, "field_widgets": {}}

        def _toggle_edit():
            if not edit_state["editing"]:
                _enter_edit_mode()
            else:
                _save_edits()

        edit_btn = ttk.Button(toolbar, text="✏ 编辑字段",
                              command=_toggle_edit)
        edit_btn.pack(side=tk.LEFT, padx=4)
        ttk.Button(toolbar, text="📋 复制身份信息",
                   command=lambda: _copy_identity(win)
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(toolbar, text="📋 复制 DOI",
                   command=lambda: _copy_doi(win)
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(toolbar, text="📋 复制原始JSON",
                   command=lambda: _copy_json(win)).pack(side=tk.LEFT, padx=4)
        ttk.Label(toolbar,
                  text="schema: " + schema_ver + "  lang: " + lang
                       + "  model: " + str(data.get("_model", "N/A")),
                  foreground="#888",
                  font=("微软雅黑", 8)).pack(side=tk.LEFT, padx=12)

        # ── 主体：PanedWindow 左渲染右编辑 ──
        paned = ttk.PanedWindow(win, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        # 左侧：渲染视图
        left_frame = ttk.LabelFrame(paned, text="渲染视图")
        paned.add(left_frame, weight=3)

        render_text = scrolledtext.ScrolledText(
            left_frame, font=("微软雅黑", 9), wrap=tk.WORD,
            state=tk.NORMAL)
        render_text.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        render_text.bind("<Key>", self._readonly_text_key)
        self._install_readonly_copy_support(render_text)

        render_text.tag_config("title", foreground="#1F4E79",
                               font=("微软雅黑", 11, "bold"))
        render_text.tag_config("identity_block", background="#F0F7FF",
                               foreground="#1F4E79",
                               font=("微软雅黑", 9))
        render_text.tag_config("identity_label", foreground="#1F4E79",
                               font=("微软雅黑", 9, "bold"))
        render_text.tag_config("section", foreground="#107C10",
                               font=("微软雅黑", 10, "bold"))
        render_text.tag_config("field", foreground="#0070C0",
                               font=("微软雅黑", 9, "bold"))
        render_text.tag_config("ai_tag", background="#FFF2CC")
        render_text.tag_config("evidence", foreground="#666",
                               font=("微软雅黑", 8, "italic"))
        render_text.tag_config("na", foreground="#999")
        render_text.tag_config(
            "doi_link", foreground="#0066CC",
            underline=True, font=("微软雅黑", 9, "bold"))

        def _copy_doi(parent):
            doi = str(data.get("DOI", "") or "").strip()
            if not doi or doi in ("N/A", "None"):
                messagebox.showinfo(
                    "提示", "当前条目没有可复制的 DOI", parent=parent)
                return "break"
            parent.clipboard_clear()
            parent.clipboard_append(doi)
            self.status_var.set("已复制 DOI：" + doi)
            self.after(
                3000,
                lambda: self.status_var.set("就绪")
                if self.status_var.get().startswith("已复制 DOI：") else None)
            return "break"

        def _copy_identity(parent):
            from schema_loader import IDENTITY_FIELDS
            label_key = "label_" + item["lang"]
            lines = ["文献ID: " + item["paper_id"]]
            for field in IDENTITY_FIELDS:
                key = field["key"]
                value = data.get(key, "N/A")
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False)
                lines.append(
                    field.get(label_key, key) + ": " + str(value))
            parent.clipboard_clear()
            parent.clipboard_append(chr(10).join(lines))
            self.status_var.set("已复制文献身份信息")
            self.after(
                3000,
                lambda: self.status_var.set("就绪")
                if self.status_var.get() == "已复制文献身份信息" else None)

        # 右侧：可编辑字段面板（初始隐藏，点编辑后展开）
        right_frame = ttk.LabelFrame(paned, text="字段编辑")
        # 先不 add，等进入编辑模式再加

        edit_canvas = tk.Canvas(right_frame, highlightthickness=0)
        edit_vsb = ttk.Scrollbar(right_frame, orient=tk.VERTICAL,
                                 command=edit_canvas.yview)
        edit_canvas.configure(yscrollcommand=edit_vsb.set)
        edit_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        edit_vsb.pack(side=tk.RIGHT, fill=tk.Y)
        edit_inner = ttk.Frame(edit_canvas)
        edit_inner_id = edit_canvas.create_window(
            (0, 0), window=edit_inner, anchor=tk.NW)
        edit_inner.bind("<Configure>",
                        lambda e: edit_canvas.configure(
                            scrollregion=edit_canvas.bbox("all")))
        edit_canvas.bind("<Configure>",
                         lambda e: edit_canvas.itemconfigure(
                             edit_inner_id, width=e.width))

        save_btn_frame = ttk.Frame(right_frame)

        def _render_view(d=None):
            """把 data 渲染到左侧 ScrolledText，复用 _render_detail 逻辑"""
            if d is None:
                d = data
            from schema_loader import IDENTITY_FIELDS, IDENTITY_KEYS
            t = render_text
            t._identity_copy_values = {}
            t._identity_copy_all_text = ""
            t._identity_doi = ""
            t.configure(state=tk.NORMAL)
            t.delete("1.0", tk.END)
            t.insert(tk.END, "📄 " + item["paper_id"] + chr(10), "title")
            t.insert(tk.END,
                     "版本: " + item["schema_ver"] +
                     "  语言: " + item["lang"] +
                     "  模型: " + str(d.get("_model", "N/A")) +
                     "  tokens: " + str(d.get("_tokens_used", "N/A")) +
                     chr(10) + chr(10))

            label_key = "label_" + item["lang"]
            t.insert(tk.END, "═══ 文献身份信息 ═══" + chr(10), "section")
            t.insert(tk.END,
                     "单击字段值复制；按住鼠标拖动可任意选择文字。" +
                     chr(10), "evidence")
            identity_lines = ["文献ID: " + item["paper_id"]]
            for f in IDENTITY_FIELDS:
                key = f["key"]
                label = f.get(label_key, key)
                val = d.get(key, "N/A")
                if isinstance(val, list):
                    parts = []
                    for it in val:
                        if isinstance(it, dict):
                            parts.append(str(it.get("value", "")))
                        else:
                            parts.append(str(it))
                    val = "; ".join(p for p in parts if p) or "N/A"
                else:
                    val = str(val) if val not in (None, "") else "N/A"
                t.insert(
                    tk.END, "  ▸ " + label + ": ",
                    "identity_label")
                tag = "na" if val in ("N/A", "None", "") else "identity_block"
                if key == "DOI" and tag != "na":
                    t._identity_doi = val
                    self._insert_copyable_identity_value(
                        t, key, label, val,
                        ("identity_block", "doi_link"))
                else:
                    self._insert_copyable_identity_value(
                        t, key, label, val, tag)
                identity_lines.append(label + ": " + val)
            t._identity_copy_all_text = chr(10).join(identity_lines)
            t.insert(tk.END, chr(10))

            try:
                schema = load_schema(item["schema_ver"])
                fields = schema.get("fields", [])
            except Exception:
                fields = []

            t.insert(tk.END, "═══ 抽取字段 ═══" + chr(10), "section")

            # 优先按 schema 顺序渲染；schema 缺失或没有定义任何字段时，
            # 退化为遍历 data 里所有非元数据字段，保证至少有内容显示
            if fields:
                rendered_keys = set()
                for field in fields:
                    key = field["key"]
                    if key in IDENTITY_KEYS:
                        continue
                    rendered_keys.add(key)
                    label = field.get(label_key, key)
                    value = d.get(key, "N/A")
                    t.insert(tk.END, "▸ " + label + " (" + key + ")"
                             + chr(10), "field")
                    if isinstance(value, list):
                        for item_val in value:
                            if isinstance(item_val, dict):
                                v = item_val.get("value", "")
                                src = item_val.get("source", "")
                                conf = item_val.get("confidence", "")
                                ev = item_val.get("evidence", "")
                                flag = ("★[AI自创] "
                                        if src == "ai_generated" else "")
                                tag = ("ai_tag"
                                       if src == "ai_generated" else None)
                                t.insert(tk.END,
                                         "   • " + flag + str(v) +
                                         " [" + str(conf) + "]" + chr(10),
                                         tag)
                                if ev:
                                    t.insert(tk.END,
                                             "     证据: " + str(ev) + chr(10),
                                             "evidence")
                            else:
                                t.insert(tk.END,
                                         "   • " + str(item_val) + chr(10))
                    else:
                        v = str(value) if value is not None else "N/A"
                        tag = "na" if v in ("N/A", "", "None") else None
                        t.insert(tk.END, "   " + v + chr(10), tag)
                    t.insert(tk.END, chr(10))

                # schema 已渲染完，再补一遍 data 里 schema 没覆盖到的字段
                # （审核后用户手工加的字段、旧 schema 字段都不会丢）
                extra_keys = [k for k in d.keys()
                              if not k.startswith("_")
                              and k not in IDENTITY_KEYS
                              and k not in rendered_keys]
                if extra_keys:
                    t.insert(tk.END,
                             "─── 其他字段（schema 未定义）───" + chr(10),
                             "section")
                    for key in extra_keys:
                        value = d.get(key, "N/A")
                        t.insert(tk.END, "▸ " + key + chr(10), "field")
                        if isinstance(value, list):
                            for item_val in value:
                                if isinstance(item_val, dict):
                                    v = item_val.get("value", "")
                                    src = item_val.get("source", "")
                                    conf = item_val.get("confidence", "")
                                    ev = item_val.get("evidence", "")
                                    flag = ("★[AI自创] "
                                            if src == "ai_generated" else "")
                                    tag = ("ai_tag"
                                           if src == "ai_generated"
                                           else None)
                                    t.insert(tk.END,
                                             "   • " + flag + str(v) +
                                             " [" + str(conf) + "]"
                                             + chr(10), tag)
                                    if ev:
                                        t.insert(tk.END,
                                                 "     证据: " + str(ev)
                                                 + chr(10),
                                                 "evidence")
                                else:
                                    t.insert(tk.END,
                                             "   • " + str(item_val)
                                             + chr(10))
                        else:
                            v = str(value) if value is not None else "N/A"
                            tag = "na" if v in ("N/A", "", "None") else None
                            t.insert(tk.END, "   " + v + chr(10), tag)
                        t.insert(tk.END, chr(10))
            else:
                # schema 完全缺失：直接遍历 data 显示
                t.insert(tk.END,
                         "（schema「" + str(item["schema_ver"])
                         + "」未找到，按原始字段显示）" + chr(10),
                         "evidence")
                for key, value in d.items():
                    if key.startswith("_") or key in IDENTITY_KEYS:
                        continue
                    t.insert(tk.END, "▸ " + key + chr(10), "field")
                    if isinstance(value, list):
                        for item_val in value:
                            if isinstance(item_val, dict):
                                v = item_val.get("value", "")
                                src = item_val.get("source", "")
                                conf = item_val.get("confidence", "")
                                ev = item_val.get("evidence", "")
                                flag = ("★[AI自创] "
                                        if src == "ai_generated" else "")
                                tag = ("ai_tag"
                                       if src == "ai_generated" else None)
                                t.insert(tk.END,
                                         "   • " + flag + str(v) +
                                         " [" + str(conf) + "]" + chr(10),
                                         tag)
                                if ev:
                                    t.insert(tk.END,
                                             "     证据: " + str(ev)
                                             + chr(10),
                                             "evidence")
                            else:
                                t.insert(tk.END,
                                         "   • " + str(item_val) + chr(10))
                    else:
                        v = str(value) if value is not None else "N/A"
                        tag = "na" if v in ("N/A", "", "None") else None
                        t.insert(tk.END, "   " + v + chr(10), tag)
                    t.insert(tk.END, chr(10))

            # 保持 NORMAL 以允许选中和 Ctrl+C；键盘编辑由绑定拦截。
            t.configure(state=tk.NORMAL)

        def _enter_edit_mode():
            """展开右侧编辑面板，为每个非元数据字段生成输入框"""
            edit_state["editing"] = True
            edit_btn.configure(text="💾 保存修改")
            paned.add(right_frame, weight=2)

            # 清空旧控件
            for w in edit_inner.winfo_children():
                w.destroy()
            edit_state["field_widgets"].clear()

            row_idx = 0
            for key, val in data.items():
                if key.startswith("_"):
                    continue
                ttk.Label(edit_inner, text=key,
                          font=("微软雅黑", 8, "bold"),
                          foreground="#0070C0",
                          wraplength=160).grid(
                    row=row_idx, column=0, sticky=tk.NW, padx=6, pady=2)

                # 列表型（JSON array）用多行文本框
                if isinstance(val, list):
                    display = json.dumps(val, ensure_ascii=False, indent=2)
                    widget = tk.Text(edit_inner, height=4, width=28,
                                     font=("Consolas", 8), wrap=tk.WORD)
                    widget.insert("1.0", display)
                else:
                    display = str(val) if val is not None else ""
                    widget = tk.Text(edit_inner, height=2, width=28,
                                     font=("微软雅黑", 9), wrap=tk.WORD)
                    widget.insert("1.0", display)

                widget.grid(row=row_idx, column=1, sticky=tk.EW,
                            padx=6, pady=2)
                edit_state["field_widgets"][key] = (widget, isinstance(val, list))
                row_idx += 1

            edit_inner.columnconfigure(1, weight=1)

            # 保存按钮
            for w in save_btn_frame.winfo_children():
                w.destroy()
            save_btn_frame.pack(fill=tk.X, padx=4, pady=4)
            ttk.Button(save_btn_frame, text="💾 保存修改",
                       command=_save_edits).pack(side=tk.LEFT, padx=4)
            ttk.Button(save_btn_frame, text="✗ 取消",
                       command=_cancel_edit).pack(side=tk.LEFT, padx=4)

        def _save_edits():
            """读取所有编辑框的值，写回数据库并刷新渲染视图"""
            changed = {}
            for key, (widget, is_list) in edit_state["field_widgets"].items():
                raw = widget.get("1.0", tk.END).strip()
                if is_list:
                    try:
                        new_val = json.loads(raw)
                    except Exception:
                        messagebox.showwarning(
                            "格式错误",
                            "字段「" + key + "」的 JSON 格式有误，已跳过该字段。"
                            + chr(10) + "请检查括号和引号是否配对。")
                        continue
                else:
                    new_val = raw
                if new_val != data.get(key):
                    changed[key] = new_val

            if not changed:
                messagebox.showinfo("提示", "没有检测到修改")
                return

            # 写回 data 并入库
            for key, val in changed.items():
                data[key] = val

            schema_v = data.get("_schema_ver", "")
            model_u = data.get("_model_used", data.get("_model", ""))
            lang_v = data.get("_lang", "zh")
            insert_paper(self.db_path, paper_id, data,
                         schema_v, model_u, lang_v)

            # 同步 extract_cache
            project_dir = self.settings.get("project_dir", "")
            if project_dir:
                cache_file = (Path(project_dir) / "extract_cache" /
                              base_id /
                              (schema_v + "_" + lang_v + ".json"))
                if cache_file.exists():
                    try:
                        with open(cache_file, encoding="utf-8") as f:
                            cdata = json.load(f)
                        for key, val in changed.items():
                            cdata[key] = val
                        with open(cache_file, "w", encoding="utf-8") as f:
                            json.dump(cdata, f,
                                      ensure_ascii=False, indent=2)
                    except Exception:
                        pass

            messagebox.showinfo(
                "已保存",
                "已修改 " + str(len(changed)) + " 个字段并写回数据库")
            _cancel_edit()
            _render_view()
            self._refresh_manage_tab()

        def _cancel_edit():
            edit_state["editing"] = False
            edit_btn.configure(text="✏ 编辑字段")
            try:
                paned.forget(right_frame)
            except Exception:
                pass
            save_btn_frame.pack_forget()

        def _copy_json(parent):
            parent.clipboard_clear()
            parent.clipboard_append(
                json.dumps(data, ensure_ascii=False, indent=2))
            messagebox.showinfo("已复制", "原始 JSON 已复制到剪贴板")

        # 初始渲染
        _render_view()

    # ── 重命名 paper_id ────────────────────────────────

    def _rename_paper_id_dialog(self):
        """管理页「🔧 重命名ID」按钮:弹出重命名弹窗"""
        selected_ids = self._selected_manage_ids()
        if not selected_ids:
            messagebox.showwarning("提示", "请先选中至少一条记录")
            return

        # 收集选中的 base_id(去重,去掉 __zh/__en 后缀)
        base_ids = []
        seen = set()
        for full_id in selected_ids:
            base = full_id.rsplit("__", 1)[0]
            if base not in seen:
                seen.add(base)
                base_ids.append(base)

        project_dir = self.settings.get("project_dir", "")

        win = tk.Toplevel(self)
        win.title("重命名文献 ID")
        win.geometry("900x680")
        win.grab_set()

        # ── 说明文字 ──
        ttk.Label(win,
                  text="命名模板:使用 {字段名} 作为占位符,"
                       "例如 {First_Author}_{Year}_{Journal_Abbr}",
                  font=("微软雅黑", 9),
                  foreground="#1F4E79"
                  ).pack(anchor=tk.W, padx=12, pady=(10, 2))
        ttk.Label(win,
                  text="Journal_Abbr 中的点号和空格会自动去除"
                       "(如 Adv. Mater. → AdvMater)",
                  font=("微软雅黑", 8),
                  foreground="#888").pack(anchor=tk.W, padx=12)

        # ── 模板输入框 ──
        tpl_frame = ttk.Frame(win)
        tpl_frame.pack(fill=tk.X, padx=12, pady=6)
        ttk.Label(tpl_frame, text="模板:").pack(side=tk.LEFT)
        tpl_var = tk.StringVar(value="{First_Author}_{Year}_{Journal_Abbr}")
        tpl_entry = ttk.Entry(tpl_frame, textvariable=tpl_var, width=50)
        tpl_entry.pack(side=tk.LEFT, padx=6)
        ttk.Button(tpl_frame, text="预览",
                   command=lambda: _do_preview()).pack(side=tk.LEFT, padx=4)

        preset_frame = ttk.Frame(win)
        preset_frame.pack(fill=tk.X, padx=12, pady=(0, 4))
        ttk.Label(preset_frame, text="快捷模板:").pack(side=tk.LEFT)
        presets = [
            ("作者_年份_期刊",
             "{First_Author}_{Year}_{Journal_Abbr}"),
            ("年份_作者_题目",
             "{Year}_{First_Author}_{Title}"),
            ("DOI",
             "doi_{DOI}"),
            ("期刊_年份_作者",
             "{Journal_Abbr}_{Year}_{First_Author}"),
        ]
        for label, template in presets:
            ttk.Button(
                preset_frame, text=label,
                command=lambda value=template: (
                    tpl_var.set(value), _do_preview())
            ).pack(side=tk.LEFT, padx=2)

        direct_frame = ttk.Frame(win)
        direct_frame.pack(fill=tk.X, padx=12, pady=(0, 6))
        ttk.Label(direct_frame, text="直接改为:").pack(side=tk.LEFT)
        direct_var = tk.StringVar(
            value=base_ids[0] if len(base_ids) == 1 else "")
        ttk.Entry(
            direct_frame, textvariable=direct_var, width=48
        ).pack(side=tk.LEFT, padx=6)
        ttk.Button(
            direct_frame, text="预览直接 ID",
            command=lambda: _do_direct_preview()
        ).pack(side=tk.LEFT, padx=4)
        ttk.Label(
            direct_frame,
            text="（仅选中一篇时可用；双击主列表 ID 可快速进入）",
            foreground="#888"
        ).pack(side=tk.LEFT, padx=4)

        # ── 预览表格 ──
        ttk.Label(win, text="预览(旧ID → 新ID):",
                  font=("微软雅黑", 9, "bold")
                  ).pack(anchor=tk.W, padx=12, pady=(4, 2))

        preview_frame = ttk.Frame(win)
        preview_frame.pack(fill=tk.BOTH, expand=True, padx=12, pady=4)

        cols = ("old_id", "new_id", "status")
        preview_tree = ttk.Treeview(preview_frame, columns=cols,
                                    show="headings", height=12)
        preview_tree.heading("old_id", text="旧 ID")
        preview_tree.heading("new_id", text="新 ID(预览)")
        preview_tree.heading("status", text="状态")
        preview_tree.column("old_id", width=280)
        preview_tree.column("new_id", width=280)
        preview_tree.column("status", width=120, anchor=tk.CENTER)
        preview_tree.tag_configure("ok", background="#E8F5E9")
        preview_tree.tag_configure("conflict", background="#FFE0E0")
        preview_tree.tag_configure("same", background="#FFF9C4")
        preview_tree.tag_configure("no_data", background="#F0F0F0")

        sb = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL,
                           command=preview_tree.yview)
        preview_tree.configure(yscrollcommand=sb.set)
        preview_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb.pack(side=tk.LEFT, fill=tk.Y)

        status_label = tk.Label(win, text="",
                                font=("微软雅黑", 9),
                                foreground="#1F4E79")
        status_label.pack(anchor=tk.W, padx=12, pady=2)

        # ── 预览逻辑 ──
        preview_data = []   # [(old_base_id, new_base_id, tag), ...]

        def _sanitize(s: str) -> str:
            """把字段值里不适合作文件名的字符替换掉"""
            s = re.sub(r'[<>:"/\\|?*\s]', '_', s)
            s = re.sub(r'_+', '_', s)
            return s.strip('_')[:40]

        def _render_id(template: str, data: dict) -> str:
            """把模板里的 {字段名} 替换为实际值"""
            def replacer(m):
                key = m.group(1)
                val = data.get(key, "")
                if not val or val in ("N/A", "None"):
                    return "NA"
                val = str(val)
                # Journal_Abbr 特殊处理:去点号和空格
                if key in ("Journal_Abbr", "Journal_Full", "Journal"):
                    val = val.replace(".", "").replace(" ", "")
                return _sanitize(val)
            result = re.sub(r'\{(\w+)\}', replacer, template)
            # 最终清理:连续下划线合并,首尾去掉
            result = re.sub(r'_+', '_', result).strip('_')
            return result if result else "unnamed"

        def _resolve_conflicts(id_list: list) -> list:
            """
            对重复的新 ID 加 _b/_c/_d 后缀。
            返回 [(old, new_resolved, tag), ...]
            tag: ok / conflict / same / no_data
            """
            existing_ids = set()
            for p in get_all_papers(self.db_path):
                base = p.get("_paper_id", "").rsplit("__", 1)[0]
                existing_ids.add(base)

            used = {}
            result = []
            suffix_letters = "bcdefghijklmnopqrstuvwxyz"
            for old_id, new_id_raw, tag in id_list:
                if tag in ("same", "no_data"):
                    result.append((old_id, new_id_raw, tag))
                    continue

                candidate = new_id_raw
                idx = 0
                while candidate in existing_ids and candidate != old_id:
                    if idx >= len(suffix_letters):
                        candidate = new_id_raw + "_" + str(idx)
                    else:
                        candidate = new_id_raw + "_" + suffix_letters[idx]
                    idx += 1

                if candidate in used:
                    used[candidate] += 1
                    candidate = (candidate + "_" +
                                 suffix_letters[min(used[candidate] - 1,
                                                    len(suffix_letters) - 1)])
                else:
                    used[candidate] = 1

                final_tag = "conflict" if candidate != new_id_raw else "ok"
                result.append((old_id, candidate, final_tag))

            return result

        def _do_preview():
            nonlocal preview_data
            template = tpl_var.get().strip()
            if not template:
                messagebox.showwarning("提示", "请输入命名模板", parent=win)
                return

            for it in preview_tree.get_children():
                preview_tree.delete(it)
            preview_data.clear()

            raw_list = []
            for base_id in base_ids:
                # 优先读 zh,没有读 en,都没有读原始 ID
                d = (get_paper(self.db_path, base_id + "__zh")
                     or get_paper(self.db_path, base_id + "__en")
                     or get_paper(self.db_path, base_id))

                if not d:
                    raw_list.append((base_id, base_id, "no_data"))
                    continue

                new_id = _render_id(template, d)
                if new_id == base_id:
                    raw_list.append((base_id, base_id, "same"))
                else:
                    raw_list.append((base_id, new_id, "ok"))

            resolved = _resolve_conflicts(raw_list)
            preview_data.extend(resolved)

            ok_count = sum(1 for _, _, t in resolved if t == "ok")
            conflict_count = sum(1 for _, _, t in resolved if t == "conflict")
            same_count = sum(1 for _, _, t in resolved if t == "same")
            no_data_count = sum(1 for _, _, t in resolved if t == "no_data")

            status_map = {
                "ok": "✓ 可重命名",
                "conflict": "⚠ 已去重后缀",
                "same": "— 无变化",
                "no_data": "✗ 无数据",
            }
            for old_id, new_id, tag in resolved:
                preview_tree.insert("", tk.END,
                                    values=(old_id, new_id,
                                            status_map.get(tag, tag)),
                                    tags=(tag,))

            status_label.configure(
                text=("共 " + str(len(resolved)) + " 篇 | "
                      "可改名: " + str(ok_count) + " | "
                      "加后缀: " + str(conflict_count) + " | "
                      "无变化: " + str(same_count) + " | "
                      "无数据: " + str(no_data_count)))

        def _do_direct_preview():
            nonlocal preview_data
            if len(base_ids) != 1:
                messagebox.showwarning(
                    "提示",
                    "直接 ID 仅支持单篇。批量重命名请使用上方模板。",
                    parent=win)
                return
            new_id = _sanitize(direct_var.get().strip())
            if not new_id:
                messagebox.showwarning(
                    "提示", "请输入新的文献 ID", parent=win)
                return

            old_id = base_ids[0]
            tag = "same" if new_id == old_id else "ok"
            resolved = _resolve_conflicts([(old_id, new_id, tag)])
            preview_data.clear()
            preview_data.extend(resolved)
            for item_id in preview_tree.get_children():
                preview_tree.delete(item_id)

            status_map = {
                "ok": "✓ 可重命名",
                "conflict": "⚠ 已去重后缀",
                "same": "— 无变化",
                "no_data": "✗ 无数据",
            }
            for current_old, current_new, current_tag in resolved:
                preview_tree.insert(
                    "", tk.END,
                    values=(
                        current_old,
                        current_new,
                        status_map.get(current_tag, current_tag)),
                    tags=(current_tag,))
            status_label.configure(
                text="直接重命名预览：" + old_id + " → " +
                     resolved[0][1])

        # ── 执行按钮 ──
        btn_row = ttk.Frame(win)
        btn_row.pack(pady=8)

        def _do_apply():
            actionable = [(o, n) for o, n, t in preview_data
                          if t in ("ok", "conflict")]
            if not actionable:
                messagebox.showinfo("提示", "没有需要重命名的条目", parent=win)
                return
            if not messagebox.askyesno(
                    "确认",
                    "将重命名 " + str(len(actionable)) + " 篇文献的 ID。"
                    + chr(10) + "此操作不可撤销,确认继续?",
                    parent=win):
                return
            # 注意：_apply_rename 必须是 App 类的方法（同样要是 4 空格缩进）
            if hasattr(self, "_apply_rename"):
                self._apply_rename(actionable, project_dir, win)
            else:
                messagebox.showwarning(
                    "提示",
                    "_apply_rename 方法尚未实现或缩进错误，已跳过",
                    parent=win)

        ttk.Button(btn_row, text="▶ 执行重命名",
                   command=_do_apply).pack(side=tk.LEFT, padx=8)
        ttk.Button(btn_row, text="✗ 取消",
                   command=win.destroy).pack(side=tk.LEFT, padx=8)

        # 打开时自动预览一次
        _do_preview()

    def _find_replace_dialog(self):
        win = tk.Toplevel(self)
        win.title("查找替换 — 当前数据库:" + self.db_name_var.get())
        win.geometry("1020x720")
        win.grab_set()

        # ── 顶部说明 ──
        ttk.Label(win,
                text="一次性替换数据库中所有匹配文本，"
                    "适合清理 \"• N/A\"、\"Weak Inference/N/A\" 等污染数据",
                font=("微软雅黑", 9),
                foreground="#1F4E79"
                ).pack(anchor=tk.W, padx=12, pady=(10, 4))

        # ── 主体：左侧历史 + 右侧操作 ──
        main_paned = ttk.PanedWindow(win, orient=tk.HORIZONTAL)
        main_paned.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)

        # ── 左侧历史记录面板 ──
        hist_frame = ttk.LabelFrame(main_paned, text="历史记录（点击复用）",
                                    padding=4)
        main_paned.add(hist_frame, weight=1)

        hist_listbox = tk.Listbox(hist_frame, font=("微软雅黑", 8),
                                selectmode=tk.SINGLE, activestyle="dotbox")
        hist_sb = ttk.Scrollbar(hist_frame, orient=tk.VERTICAL,
                                command=hist_listbox.yview)
        hist_listbox.configure(yscrollcommand=hist_sb.set)
        hist_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        hist_sb.pack(side=tk.RIGHT, fill=tk.Y)

        hist_btn_row = ttk.Frame(hist_frame)
        hist_btn_row.pack(fill=tk.X, pady=2)

        # ── 右侧操作面板 ──
        op_frame = ttk.Frame(main_paned)
        main_paned.add(op_frame, weight=3)

        # ── 输入区 ──
        in_frame = ttk.LabelFrame(op_frame, text="查找与替换", padding=8)
        in_frame.pack(fill=tk.X, padx=4, pady=4)

        ttk.Label(in_frame, text="查找:").grid(
            row=0, column=0, sticky=tk.W, pady=2)
        find_var = tk.StringVar()
        ttk.Entry(in_frame, textvariable=find_var, width=55).grid(
            row=0, column=1, columnspan=3, sticky=tk.W,
            pady=2, padx=4)

        ttk.Label(in_frame, text="替换为:").grid(
            row=1, column=0, sticky=tk.W, pady=2)
        repl_var = tk.StringVar()
        ttk.Entry(in_frame, textvariable=repl_var, width=55).grid(
            row=1, column=1, columnspan=3, sticky=tk.W,
            pady=2, padx=4)
        ttk.Label(in_frame,
                text="(留空则把匹配项删除)",
                foreground="#888"
                ).grid(row=2, column=1, sticky=tk.W, padx=4)

        # ── 范围选项 ──
        scope_frame = ttk.LabelFrame(op_frame, text="范围与选项", padding=8)
        scope_frame.pack(fill=tk.X, padx=4, pady=4)

        ttk.Label(scope_frame, text="字段:").grid(
            row=0, column=0, sticky=tk.W, pady=2)
        all_papers = get_all_papers(self.db_path)
        field_set = set()
        for p in all_papers:
            for k in p.keys():
                if not k.startswith("_"):
                    field_set.add(k)
        field_options = ["全部字段"] + sorted(field_set)
        field_var = tk.StringVar(value="全部字段")
        ttk.Combobox(scope_frame, textvariable=field_var,
                    values=field_options, width=28,
                    state="readonly"
                    ).grid(row=0, column=1, sticky=tk.W, padx=4)

        ttk.Label(scope_frame, text="语言:").grid(
            row=0, column=2, sticky=tk.W, padx=(16, 4))
        lang_var = tk.StringVar(value="全部")
        ttk.Combobox(scope_frame, textvariable=lang_var,
                    values=["全部", "zh", "en"], width=8,
                    state="readonly"
                    ).grid(row=0, column=3, sticky=tk.W, padx=4)

        case_var = tk.BooleanVar(value=False)
        word_var = tk.BooleanVar(value=False)
        regex_var = tk.BooleanVar(value=False)
        delete_empty_var = tk.BooleanVar(value=True)
        sync_cache_var = tk.BooleanVar(value=True)

        ttk.Checkbutton(scope_frame, text="大小写敏感",
                        variable=case_var
                        ).grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Checkbutton(scope_frame, text="整词匹配",
                        variable=word_var
                        ).grid(row=1, column=1, sticky=tk.W, pady=4)
        ttk.Checkbutton(scope_frame, text="正则表达式",
                        variable=regex_var
                        ).grid(row=1, column=2, sticky=tk.W, pady=4)
        ttk.Checkbutton(scope_frame,
                        text="替换后值为空时改成 N/A",
                        variable=delete_empty_var
                        ).grid(row=2, column=0, columnspan=2,
                            sticky=tk.W, pady=4)
        ttk.Checkbutton(scope_frame,
                        text="同步更新 extract_cache 缓存(推荐)",
                        variable=sync_cache_var
                        ).grid(row=2, column=2, columnspan=2,
                            sticky=tk.W, pady=4)

        # ── 预览区 ──
        ttk.Label(op_frame, text="预览(替换前 → 替换后):",
                font=("微软雅黑", 9, "bold")
                ).pack(anchor=tk.W, padx=4, pady=(4, 2))

        preview_frame = ttk.Frame(op_frame)
        preview_frame.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

        cols = ("paper_id", "field", "before", "after")
        preview_tree = ttk.Treeview(preview_frame, columns=cols,
                                    show="headings", height=8)
        preview_tree.heading("paper_id", text="文献ID")
        preview_tree.heading("field", text="字段")
        preview_tree.heading("before", text="替换前")
        preview_tree.heading("after", text="替换后")
        preview_tree.column("paper_id", width=160)
        preview_tree.column("field", width=110)
        preview_tree.column("before", width=200)
        preview_tree.column("after", width=200)

        psb = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL,
                            command=preview_tree.yview)
        preview_tree.configure(yscrollcommand=psb.set)
        preview_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        psb.pack(side=tk.RIGHT, fill=tk.Y)

        status_label = tk.Label(op_frame, text="",
                                font=("微软雅黑", 9),
                                foreground="#1F4E79")
        status_label.pack(anchor=tk.W, padx=4, pady=2)

        plan = {"items": []}

        # ── 历史记录辅助函数 ──
        def _refresh_hist_list():
            hist_listbox.delete(0, tk.END)
            for rec in reversed(self._find_replace_history):
                label = (rec["find"][:20] +
                        (" → " + rec["replace"][:15] if rec["replace"] else " → [删除]") +
                        "  [" + rec["field"] + "/" + rec["lang"] + "]")
                hist_listbox.insert(tk.END, label)

        def _load_hist_item(event=None):
            """双击历史记录，把参数填回表单"""
            sel = hist_listbox.curselection()
            if not sel:
                return
            # 历史列表是逆序显示的，所以索引需要转换
            real_idx = len(self._find_replace_history) - 1 - sel[0]
            rec = self._find_replace_history[real_idx]
            find_var.set(rec["find"])
            repl_var.set(rec["replace"])
            field_var.set(rec["field"])
            lang_var.set(rec["lang"])
            case_var.set(rec["case_sensitive"])
            word_var.set(rec["whole_word"])
            regex_var.set(rec["use_regex"])
            delete_empty_var.set(rec.get("delete_empty", True))
            sync_cache_var.set(rec.get("sync_cache", True))

        def _delete_hist_item():
            sel = hist_listbox.curselection()
            if not sel:
                return
            real_idx = len(self._find_replace_history) - 1 - sel[0]
            del self._find_replace_history[real_idx]
            _refresh_hist_list()

        def _clear_hist():
            if not messagebox.askyesno("确认", "清空所有历史记录？", parent=win):
                return
            self._find_replace_history.clear()
            _refresh_hist_list()

        hist_listbox.bind("<Double-Button-1>", _load_hist_item)
        # 历史操作按钮：竖排显示
        ttk.Button(hist_btn_row, text="↑ 复用",
                   command=_load_hist_item).pack(fill=tk.X, pady=2)
        ttk.Button(hist_btn_row, text="✗ 删除",
                   command=_delete_hist_item).pack(fill=tk.X, pady=2)
        ttk.Button(hist_btn_row, text="🗑 清空",
                   command=_clear_hist).pack(fill=tk.X, pady=2)

        # ── 一键清理 N/A 污染（结构层面归一化，独立于查找替换流程）──
        ttk.Separator(hist_btn_row,
                      orient=tk.HORIZONTAL).pack(fill=tk.X, pady=6)
        self._ai_button(hist_btn_row, text="🧹 一键清理 N/A",
                        command=lambda: self._cleanup_na_pollution(win)
                        ).pack(fill=tk.X, pady=2)
        ttk.Label(hist_btn_row,
                  text="把 [\"N/A\"]、[]、null 等\n统一为字符串 \"N/A\"",
                  foreground="#888",
                  font=("微软雅黑", 8),
                  justify=tk.LEFT,
                  wraplength=140).pack(fill=tk.X, pady=2)

        # ── 预览/执行逻辑（与原版一致）──
        def _build_pattern():
            pat = find_var.get()
            if not pat:
                return None
            if not regex_var.get():
                pat = re.escape(pat)
            if word_var.get():
                pat = r"\b" + pat + r"\b"
            flags = 0 if case_var.get() else re.IGNORECASE
            try:
                return re.compile(pat, flags)
            except re.error as e:
                messagebox.showerror("正则错误", str(e), parent=win)
                return None

        def _replace_in_value(val, pattern, replacement):
            if isinstance(val, str):
                new_val, n = pattern.subn(replacement, val)
                return new_val, n
            if isinstance(val, list):
                total = 0
                new_list = []
                for it in val:
                    nv, c = _replace_in_value(it, pattern, replacement)
                    total += c
                    new_list.append(nv)
                return new_list, total
            if isinstance(val, dict):
                total = 0
                new_dict = {}
                for k, v in val.items():
                    nv, c = _replace_in_value(v, pattern, replacement)
                    total += c
                    new_dict[k] = nv
                return new_dict, total
            return val, 0

        def _format_preview(val) -> str:
            if isinstance(val, str):
                return val[:120]
            if isinstance(val, list):
                parts = []
                for it in val:
                    if isinstance(it, dict):
                        parts.append(str(it.get("value", "")))
                    else:
                        parts.append(str(it))
                return ("; ".join(parts))[:120]
            return str(val)[:120]

        def _do_preview():
            for item in preview_tree.get_children():
                preview_tree.delete(item)
            plan["items"].clear()

            pattern = _build_pattern()
            if not pattern:
                return

            replacement = repl_var.get()
            target_field = field_var.get()
            target_lang = lang_var.get()

            papers = get_all_papers(self.db_path)
            total_hits = 0
            affected_papers = set()

            for p in papers:
                lang = p.get("_lang", "")
                if target_lang != "全部" and lang != target_lang:
                    continue
                full_id = p.get("_paper_id", "")

                for key, val in p.items():
                    if key.startswith("_"):
                        continue
                    if (target_field != "全部字段"
                            and key != target_field):
                        continue
                    new_val, n = _replace_in_value(val, pattern, replacement)
                    if n == 0:
                        continue

                    if (delete_empty_var.get()
                            and isinstance(new_val, str)
                            and new_val.strip() == ""):
                        new_val = "N/A"

                    total_hits += n
                    affected_papers.add(full_id)
                    plan["items"].append({
                        "paper_id": full_id,
                        "lang": lang,
                        "field": key,
                        "new_value": new_val,
                    })
                    preview_tree.insert(
                        "", tk.END,
                        values=(full_id, key,
                                _format_preview(val),
                                _format_preview(new_val)))

            status_label.configure(
                text=("命中 " + str(total_hits) + " 处 | 涉及 "
                    + str(len(affected_papers)) + " 篇文献 | "
                    + "字段更新 " + str(len(plan["items"])) + " 次"))

        def _do_apply():
            if not plan["items"]:
                messagebox.showinfo("提示",
                                    "没有可替换的内容，请先点「预览」",
                                    parent=win)
                return
            if not messagebox.askyesno(
                    "确认执行",
                    "将对数据库执行 " + str(len(plan["items"]))
                    + " 处字段替换。" + chr(10)
                    + "建议先到「设置 → 数据备份」导出一份数据库备份，"
                    + "确认继续？",
                    parent=win):
                return

            # ── 执行前把当前参数存入历史记录 ──
            find_text = find_var.get()
            if find_text:
                rec = {
                    "find":           find_text,
                    "replace":        repl_var.get(),
                    "field":          field_var.get(),
                    "lang":           lang_var.get(),
                    "case_sensitive": case_var.get(),
                    "whole_word":     word_var.get(),
                    "use_regex":      regex_var.get(),
                    "delete_empty":   delete_empty_var.get(),
                    "sync_cache":     sync_cache_var.get(),
                }
                # 去重：完全相同的规则不重复记录
                existing = [r for r in self._find_replace_history
                            if r["find"] == rec["find"]
                            and r["replace"] == rec["replace"]
                            and r["field"] == rec["field"]
                            and r["lang"] == rec["lang"]]
                if not existing:
                    self._find_replace_history.append(rec)
                    # 最多保留 50 条
                    if len(self._find_replace_history) > 50:
                        self._find_replace_history.pop(0)
                    _refresh_hist_list()

            self._apply_find_replace(plan["items"], sync_cache_var.get(), win)

        # ── 按钮区（固定在底部，不会被挤走）──
        btn_row = ttk.Frame(op_frame)
        btn_row.pack(side=tk.BOTTOM, fill=tk.X, pady=8, padx=4)
        ttk.Button(btn_row, text="🔍 预览",
                command=_do_preview).pack(side=tk.LEFT, padx=8)
        ttk.Button(btn_row, text="▶ 执行替换",
                command=_do_apply).pack(side=tk.LEFT, padx=8)
        ttk.Button(btn_row, text="✗ 关闭",
                command=win.destroy).pack(side=tk.LEFT, padx=8)

        # 初始化历史列表
        _refresh_hist_list()

    def _apply_find_replace(self, items: list, sync_cache: bool,
                            parent_win=None):
        """
        执行替换:
        逐条读出 paper 完整 JSON,把对应字段更新后调用 insert_paper 覆盖入库;
        sync_cache=True 时同步更新 extract_cache/<paper_id>/<schema_ver>_<lang>.json
        """
        # 按 paper_id 聚合,避免一篇文献被多次写入
        grouped = {}
        for it in items:
            pid = it["paper_id"]
            grouped.setdefault(pid, []).append(it)

        ok_n = 0
        fail_n = 0
        cache_updated = 0
        project_dir = self.settings.get("project_dir", "")

        for full_id, updates in grouped.items():
            try:
                data = get_paper(self.db_path, full_id)
                if not data:
                    fail_n += 1
                    continue

                # 应用所有字段更新
                for u in updates:
                    data[u["field"]] = u["new_value"]

                schema_ver = data.get("_schema_ver", "")
                model_used = data.get("_model_used",
                                      data.get("_model", ""))
                lang = data.get("_lang", "")

                insert_paper(self.db_path, full_id, data,
                             schema_ver, model_used, lang)
                ok_n += 1

                # 同步缓存
                if sync_cache and project_dir:
                    base_id = full_id.rsplit("__", 1)[0]
                    cache_file = (Path(project_dir) / "extract_cache"
                                  / base_id
                                  / (schema_ver + "_" + lang + ".json"))
                    if cache_file.exists():
                        try:
                            with open(cache_file, encoding="utf-8") as f:
                                cdata = json.load(f)
                            for u in updates:
                                cdata[u["field"]] = u["new_value"]
                            with open(cache_file, "w",
                                      encoding="utf-8") as f:
                                json.dump(cdata, f,
                                          ensure_ascii=False, indent=2)
                            cache_updated += 1
                        except Exception:
                            pass

            except Exception as e:
                self._log("✗ 替换失败 " + full_id + ": "
                          + str(e)[:120])
                fail_n += 1

        self._refresh_manage_tab()
        if parent_win:
            try:
                parent_win.destroy()
            except Exception:
                pass

        msg = ("成功更新 " + str(ok_n) + " 篇" + chr(10)
               + "失败 " + str(fail_n) + " 篇")
        if sync_cache:
            msg += chr(10) + "同步缓存文件 " + str(cache_updated) + " 个"
        messagebox.showinfo("替换完成", msg)

    def _apply_rename(self, pairs: list, project_dir: str, parent_win=None):
        """
        执行批量重命名:逐条调用 database.rename_paper_id,
        完成后刷新管理页并弹结果摘要。
        """
        from database import rename_paper_id

        success, failed = [], []
        messages = []

        for old_id, new_id in pairs:
            result = rename_paper_id(self.db_path, old_id, new_id, project_dir)
            if result["ok"]:
                success.append((old_id, new_id))
                if result.get("cache_msg"):
                    messages.append(new_id + ": " + result["cache_msg"])
            else:
                failed.append((old_id, result.get("error", "未知错误")))

        self._refresh_manage_tab()
        if parent_win:
            try:
                parent_win.destroy()
            except Exception:
                pass

        summary = ("✓ 成功: " + str(len(success)) + " 篇\n"
                   "✗ 失败: " + str(len(failed)) + " 篇")
        if failed:
            summary += "\n\n失败详情:\n"
            summary += "\n".join("  " + o + " → " + e for o, e in failed[:10])
        if messages:
            summary += "\n\n提示:\n" + "\n".join(messages[:5])

        messagebox.showinfo("重命名完成", summary)

    # ── N/A问题统一处理 ─────────────────────────────
    @staticmethod
    def _normalize_empty_values(data: dict) -> tuple:
        """
        归一化单条记录里的所有空值变体为统一字符串 "N/A"。
        返回 (cleaned_data, changed_keys)：
          cleaned_data: 处理后的 dict
          changed_keys: 被改动的字段列表
        元数据字段（_ 开头）原样保留。
        """
        if not isinstance(data, dict):
            return data, []

        def _is_na_value(v):
            if v is None:
                return True
            if isinstance(v, str):
                return v.strip() in (
                    "", "N/A", "n/a", "None", "null", "NA", "na")
            if isinstance(v, dict):
                return _is_na_value(v.get("value", ""))
            return False

        cleaned = {}
        changed_keys = []
        for key, val in data.items():
            if key.startswith("_"):
                cleaned[key] = val
                continue

            new_val = val
            if _is_na_value(val):
                new_val = "N/A"
            elif isinstance(val, list):
                if len(val) == 0:
                    new_val = "N/A"
                elif all(_is_na_value(it) for it in val):
                    new_val = "N/A"
                else:
                    # 有真值的 list：保留但过滤掉零星 N/A 元素
                    filtered = [it for it in val if not _is_na_value(it)]
                    new_val = filtered if filtered else "N/A"

            if new_val != val:
                changed_keys.append(key)
            cleaned[key] = new_val

        return cleaned, changed_keys
    
    def _cleanup_na_pollution(self, parent_win=None):
        """
        一键清理数据库里的 N/A 污染：
        把所有 ["N/A"]、[]、null、{"value":"N/A"} 等空值变体
        统一归一化为字符串 "N/A"。
        同步更新 extract_cache 里的对应缓存文件。
        """
        if not messagebox.askyesno(
                "确认清理",
                "将扫描当前数据库「" + self.db_name_var.get()
                + "」的所有记录，把以下空值变体统一为字符串 \"N/A\"："
                + chr(10) + chr(10)
                + "  • [\"N/A\"] / [\"\"] / [] / [None] 等空数组" + chr(10)
                + "  • null / \"None\" / \"\" 等空标量" + chr(10)
                + "  • [{\"value\":\"N/A\",...}] 等只含 N/A 的对象数组"
                + chr(10) + chr(10)
                + "建议先到「设置 → 数据备份 → 导出数据库」备份。"
                + chr(10) + "确认继续？",
                parent=parent_win):
            return

        papers = get_all_papers(self.db_path)
        project_dir = self.settings.get("project_dir", "")

        ok_n = 0
        skip_n = 0
        cache_n = 0
        total_field_changes = 0

        for p in papers:
            full_id = p.get("_paper_id", "")
            if not full_id:
                continue

            cleaned, changed_keys = self._normalize_empty_values(p)
            if not changed_keys:
                skip_n += 1
                continue

            try:
                schema_ver = cleaned.get("_schema_ver", "")
                model_used = cleaned.get(
                    "_model_used", cleaned.get("_model", ""))
                lang = cleaned.get("_lang", "")
                insert_paper(self.db_path, full_id, cleaned,
                             schema_ver, model_used, lang)
                ok_n += 1
                total_field_changes += len(changed_keys)

                # 同步更新 extract_cache
                if project_dir:
                    base_id = full_id.rsplit("__", 1)[0]
                    cache_file = (Path(project_dir) / "extract_cache"
                                  / base_id
                                  / (schema_ver + "_" + lang + ".json"))
                    if cache_file.exists():
                        try:
                            with open(cache_file,
                                      encoding="utf-8") as f:
                                cdata = json.load(f)
                            cdata2, _ = self._normalize_empty_values(cdata)
                            with open(cache_file, "w",
                                      encoding="utf-8") as f:
                                json.dump(cdata2, f,
                                          ensure_ascii=False, indent=2)
                            cache_n += 1
                        except Exception:
                            pass

                self._log("✓ 已归一化 " + full_id + "（"
                          + str(len(changed_keys)) + " 个字段）")
            except Exception as e:
                self._log("✗ 归一化失败 " + full_id + ": "
                          + str(e)[:120])

        self._refresh_manage_tab()
        msg = ("扫描 " + str(len(papers)) + " 条记录" + chr(10)
               + "已归一化: " + str(ok_n) + " 条" + chr(10)
               + "字段更新总计: " + str(total_field_changes) + " 处" + chr(10)
               + "无需处理: " + str(skip_n) + " 条" + chr(10)
               + "同步缓存文件: " + str(cache_n) + " 个")
        messagebox.showinfo("N/A 清理完成", msg,
                            parent=parent_win or self)

    # ── 导出 ───────────────────────────────────────

    def _export_excel(self, lang: str):
        papers = get_all_papers(self.db_path)
        papers = [p for p in papers if p.get("_lang") == lang]
        if not papers:
            messagebox.showwarning(
                "提示", "没有 " + lang + " 语言的入库数据")
            return

        schema_ver = papers[0].get("_schema_ver", "")
        try:
            schema = load_schema(schema_ver)
        except Exception as e:
            messagebox.showerror("错误", "加载Schema失败:" + str(e))
            return

        out_path = filedialog.asksaveasfilename(
            initialdir=str(Path(self.settings.get("project_dir", "")) /
                           "exports"),
            initialfile="export_" + self.db_name_var.get() + "_" + lang + "_" +
                        datetime.now().strftime("%Y%m%d") + ".xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")])
        if not out_path:
            return
        try:
            export_excel(papers, schema, out_path, lang)
            messagebox.showinfo(
                "成功",
                "已导出 " + str(len(papers)) + " 条到:" + out_path)
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def _export_markdown(self, lang: str = "zh"):
        """导出当前数据库为 Markdown 文件"""
        from exporter import export_markdown
        project_dir = self.settings.get("project_dir", "")
        schema_ver = self.schema_var.get()
        if not schema_ver:
            messagebox.showwarning("提示", "请先在抽取页选择提示词版本")
            return
        try:
            schema = load_schema(schema_ver)
        except Exception as e:
            messagebox.showerror("错误", "加载 Schema 失败：" + str(e))
            return

        papers = get_all_papers(self.db_path)
        lang_papers = [p for p in papers if p.get("_lang", "zh") == lang]
        if not lang_papers:
            messagebox.showinfo("提示", "当前数据库没有语言为 " + lang + " 的记录")
            return

        default_name = (self.db_name_var.get() + "_" + schema_ver +
                        "_" + lang + ".md")
        out_path = filedialog.asksaveasfilename(
            title="保存 Markdown 文件",
            defaultextension=".md",
            initialfile=default_name,
            filetypes=[("Markdown", "*.md"), ("所有文件", "*.*")])
        if not out_path:
            return

        try:
            export_markdown(
                lang_papers, schema, out_path, lang,
                seq_map=getattr(self, "_paper_seq_map", None))
            messagebox.showinfo(
                "导出成功",
                "已导出 " + str(len(lang_papers)) + " 篇文献" + chr(10) +
                "路径：" + out_path)
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _export_zip(self):
        out_path = filedialog.asksaveasfilename(
            initialfile="backup_" + datetime.now().strftime("%Y%m%d_%H%M") + ".zip",
            defaultextension=".zip",
            filetypes=[("Zip", "*.zip")])
        if not out_path:
            return
        try:
            result = export_zip(
                self.settings.get("project_dir", ""), out_path)
            messagebox.showinfo(
                "备份完成",
                "已备份 " + str(len(result.get("databases", []))) +
                " 个文献库、" + str(result.get("file_count", 0)) +
                " 个文件。" + NL + out_path)
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def _import_zip(self):
        zip_path = filedialog.askopenfilename(filetypes=[("Zip", "*.zip")])
        if not zip_path:
            return
        try:
            info = inspect_backup(zip_path)
        except Exception as e:
            messagebox.showerror("无法读取备份", str(e))
            return

        win = tk.Toplevel(self)
        win.title("导入文献库备份")
        win.geometry("590x440")
        win.transient(self)
        win.grab_set()
        body = ttk.Frame(win, padding=14)
        body.pack(fill=tk.BOTH, expand=True)

        version_desc = (
            "新版备份 v" + str(info["format_version"])
            if info.get("has_manifest") else "旧版兼容备份")
        ttk.Label(
            body,
            text=version_desc + "；包含文献库：" +
                 "、".join(info.get("databases", [])),
            wraplength=550, foreground="#1F4E79",
        ).pack(fill=tk.X, pady=(0, 12))

        source_var = tk.StringVar(
            value=(info.get("databases") or ["main"])[0])
        source_row = ttk.Frame(body)
        source_row.pack(fill=tk.X, pady=4)
        ttk.Label(source_row, text="源文献库：", width=14).pack(side=tk.LEFT)
        ttk.Combobox(
            source_row, textvariable=source_var,
            values=info.get("databases", []),
            state="readonly", width=30).pack(side=tk.LEFT)

        mode_var = tk.StringVar(value="merge")
        mode_box = ttk.LabelFrame(body, text="导入方式", padding=8)
        mode_box.pack(fill=tk.X, pady=10)
        ttk.Radiobutton(
            mode_box, text="安全合并到当前库（只补空字段，不覆盖已有值）",
            variable=mode_var, value="merge").pack(anchor=tk.W, pady=3)
        ttk.Radiobutton(
            mode_box, text="导入为新文献库（备份含多库时全部保留）",
            variable=mode_var, value="new_library").pack(anchor=tk.W, pady=3)
        ttk.Radiobutton(
            mode_box, text="替换当前库（替换前自动生成 .preimport 备份）",
            variable=mode_var, value="replace").pack(anchor=tk.W, pady=3)

        new_name_var = tk.StringVar(
            value="imported_" + datetime.now().strftime("%Y%m%d"))
        name_row = ttk.Frame(body)
        name_row.pack(fill=tk.X, pady=4)
        ttk.Label(name_row, text="新库名称前缀：", width=14).pack(side=tk.LEFT)
        ttk.Entry(
            name_row, textvariable=new_name_var,
            width=32).pack(side=tk.LEFT)

        confirmed = {"value": False}

        def confirm_import():
            if mode_var.get() == "replace":
                if not messagebox.askyesno(
                        "确认替换",
                        "当前库会先自动备份，然后由所选源库替换。继续吗？",
                        parent=win):
                    return
            confirmed["value"] = True
            win.destroy()

        actions = ttk.Frame(body)
        actions.pack(side=tk.BOTTOM, fill=tk.X, pady=(14, 0))
        ttk.Button(
            actions, text="取消", command=win.destroy
        ).pack(side=tk.RIGHT, padx=4)
        ttk.Button(
            actions, text="开始导入", command=confirm_import
        ).pack(side=tk.RIGHT, padx=4)
        self.wait_window(win)
        if not confirmed["value"]:
            return

        try:
            mode = mode_var.get()
            result = import_zip(
                zip_path,
                self.settings.get("project_dir", ""),
                mode=mode,
                target_db_name=self.db_name_var.get(),
                source_db_name=(
                    None if mode == "new_library" else source_var.get()),
                new_db_name=(
                    new_name_var.get() if mode == "new_library" else None),
            )
            self._refresh_db_list()
            lines = [
                "导入完成",
                "文献库：" +
                "、".join(result.get("imported_databases", [])),
            ]
            if result.get("verified_files"):
                lines.append(
                    "完整性校验：" +
                    str(result["verified_files"]) + " 个文件通过")
            if mode == "merge":
                lines.append(
                    "新增 " + str(result.get("inserted", 0)) +
                    " 条；合并/跳过 " + str(result.get("merged", 0)) + " 条")
            if result.get("backups"):
                lines.append("替换前备份：" + result["backups"][0])
            messagebox.showinfo("成功", NL.join(lines))
        except Exception as e:
            messagebox.showerror("错误", str(e))

    # ── 引用导出 ───────────────────────────────────

    def _citation_library_records(
            self, language_mode="字段最完整（双语互补）") -> list:
        grouped = {}
        base_order = []
        for paper in get_all_papers(self.db_path):
            full_id = paper.get("_paper_id", "")
            base = full_id.rsplit("__", 1)[0]
            if base not in grouped:
                base_order.append(base)
            grouped.setdefault(base, []).append(paper)
        records = []
        for base in base_order:
            record = select_citation_record(
                grouped[base], language_mode)
            record["_citation_base_id"] = base
            records.append(record)
        return records

    def _citation_source_records(
            self, language_mode="字段最完整（双语互补）") -> tuple:
        selected_ids = self._selected_manage_ids()
        if selected_ids:
            selected_bases = {
                str(paper_id).rsplit("__", 1)[0]
                for paper_id in selected_ids
            }
            base_ids = []
            for row in self._all_papers_cache:
                base = str(row[1]).rsplit("__", 1)[0]
                if base in selected_bases and base not in base_ids:
                    base_ids.append(base)
            scope = "已选文献"
        else:
            base_ids = self._get_visible_base_ids()
            scope = (
                "当前筛选结果"
                if self.filter_var.get().strip() else "当前文献库全部文献"
            )

        library_records = self._citation_library_records(language_mode)
        by_base = {
            record.get("_citation_base_id"): record
            for record in library_records}
        records = [by_base[base] for base in base_ids if base in by_base]
        return records, scope

    def _open_citation_builder(self):
        ordered, scope = self._citation_source_records()
        if not ordered:
            messagebox.showwarning(
                "引用格式生成", "当前选择或筛选范围内没有可用文献。")
            return

        project_dir = self.settings.get("project_dir", "")
        schemes_payload = load_citation_schemes(project_dir)
        if schemes_payload.get("warnings"):
            messagebox.showwarning(
                "引用方案恢复",
                NL.join(schemes_payload["warnings"]),
                parent=self)
        win = tk.Toplevel(self)
        win.title("引用格式生成器 — " + self.db_name_var.get())
        win.geometry("1260x840")
        win.minsize(1040, 650)
        win.transient(self)

        office_bridge = OfficeBridge()
        session_store = CitationSessionStore(project_dir)

        def help_button(parent, title, text):
            button = tk.Button(
                parent, text="?", width=2, padx=1, pady=0,
                relief=tk.FLAT, cursor="hand2",
                command=lambda: messagebox.showinfo(
                    title, text, parent=parent.winfo_toplevel()))
            button._theme_role = "accent"
            return button

        header = ttk.Frame(win, padding=(10, 8))
        header.pack(fill=tk.X)
        count_var = tk.StringVar()
        ttk.Label(
            header,
            text="来源：" + scope,
            foreground="#1F4E79",
            font=("微软雅黑", 10, "bold"),
        ).pack(side=tk.LEFT)
        ttk.Label(header, textvariable=count_var).pack(side=tk.LEFT, padx=12)
        citation_language_var = tk.StringVar(
            value="字段最完整（双语互补）")
        citation_language_combo = ttk.Combobox(
            header, textvariable=citation_language_var,
            values=[
                "字段最完整（双语互补）",
                "中文记录优先",
                "英文记录优先",
            ],
            state="readonly", width=21)
        citation_language_combo.pack(side=tk.LEFT, padx=8)
        ttk.Label(
            header,
            text="未选文献时自动使用当前筛选结果",
            foreground="#777777",
        ).pack(side=tk.RIGHT)

        pane = ttk.Panedwindow(win, orient=tk.HORIZONTAL)
        pane.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 8))
        left = ttk.Frame(pane)
        right = ttk.Frame(pane)
        pane.add(left, weight=5)
        pane.add(right, weight=7)

        order_box = ttk.LabelFrame(left, text="引用顺序", padding=6)
        order_box.pack(fill=tk.BOTH, expand=True)
        columns = ("seq", "title", "author", "year")
        order_tree = ttk.Treeview(
            order_box, columns=columns, show="headings",
            selectmode=tk.BROWSE)
        order_tree.heading("seq", text="#")
        order_tree.heading("title", text="论文名称")
        order_tree.heading("author", text="第一作者")
        order_tree.heading("year", text="年份")
        order_tree.column("seq", width=42, anchor=tk.CENTER, stretch=False)
        order_tree.column("title", width=310)
        order_tree.column("author", width=105)
        order_tree.column("year", width=55, anchor=tk.CENTER)
        order_scroll = ttk.Scrollbar(
            order_box, orient=tk.VERTICAL, command=order_tree.yview)
        order_tree.configure(yscrollcommand=order_scroll.set)
        order_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        order_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        move_bar = ttk.Frame(left)
        move_bar.pack(fill=tk.X, pady=6)
        sort_var = tk.StringVar(value="排序方式")
        sort_combo = ttk.Combobox(
            move_bar, textvariable=sort_var,
            values=["第一作者", "年份（升序）", "年份（降序）",
                    "论文名称", "期刊名称"],
            state="readonly", width=13)
        sort_combo.pack(side=tk.RIGHT, padx=3)

        scheme_box = ttk.LabelFrame(right, text="引用方案", padding=8)
        scheme_box.pack(fill=tk.X)
        scheme_name_var = tk.StringVar()
        scheme_select_var = tk.StringVar()
        number_var = tk.StringVar(value="[{n}]")
        separator_var = tk.StringVar(value=NL)
        inline_style_var = tk.StringVar(value="数字编号")
        inline_template_var = tk.StringVar(value=DEFAULT_INLINE_TEMPLATE)
        inline_item_template_var = tk.StringVar(
            value=DEFAULT_INLINE_ITEM_TEMPLATE)
        inline_separator_var = tk.StringVar(value=",")
        compress_ranges_var = tk.BooleanVar(value=True)
        content_mode_var = tk.StringVar(value="bibliography")
        privacy_mode_var = tk.StringVar(value="privacy")
        target_app_var = tk.StringVar(value="auto")
        target_status_var = tk.StringVar(value="尚未检测 Word/WPS")
        current_scheme_id = {"value": DEFAULT_SCHEME["id"]}

        first_row = ttk.Frame(scheme_box)
        first_row.pack(fill=tk.X, pady=2)
        ttk.Label(first_row, text="读取方案：").pack(side=tk.LEFT)
        scheme_combo = ttk.Combobox(
            first_row, textvariable=scheme_select_var,
            state="readonly", width=24)
        scheme_combo.pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(first_row, text="方案名称：").pack(side=tk.LEFT)
        ttk.Entry(
            first_row, textvariable=scheme_name_var,
            width=24).pack(side=tk.LEFT, fill=tk.X, expand=True)

        second_row = ttk.Frame(scheme_box)
        second_row.pack(fill=tk.X, pady=4)
        ttk.Label(second_row, text="序号类型：").pack(side=tk.LEFT)
        number_combo = ttk.Combobox(
            second_row, textvariable=number_var,
            values=["[{n}]", "{n}.", "({n})", "{circled}", ""],
            width=13)
        number_combo.pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(
            second_row,
            text="可自定义；{n}=序号，{circled}=圆圈序号",
            foreground="#777777",
        ).pack(side=tk.LEFT)

        inline_row = ttk.Frame(scheme_box)
        inline_row.pack(fill=tk.X, pady=3)
        ttk.Label(inline_row, text="正文引用：").pack(side=tk.LEFT)
        inline_style_combo = ttk.Combobox(
            inline_row, textvariable=inline_style_var,
            values=["数字编号", "作者-年份"],
            state="readonly", width=12)
        inline_style_combo.pack(side=tk.LEFT, padx=(0, 3))
        help_button(
            inline_row, "正文引用类型",
            "numeric 使用文档级编号，例如 [1–3]；author_year 使用作者和年份，"
            "例如 (Wang, 2024)。此设置会随引用方案保存。"
        ).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(inline_row, text="外层模板：").pack(side=tk.LEFT)
        ttk.Entry(
            inline_row, textvariable=inline_template_var,
            width=18).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Label(inline_row, text="分隔：").pack(side=tk.LEFT)
        ttk.Entry(
            inline_row, textvariable=inline_separator_var,
            width=5).pack(side=tk.LEFT, padx=(0, 6))
        inline_item_row = ttk.Frame(scheme_box)
        inline_item_row.pack(fill=tk.X, pady=(0, 3))
        ttk.Label(
            inline_item_row,
            text="作者年份单项模板：").pack(side=tk.LEFT)
        ttk.Entry(
            inline_item_row, textvariable=inline_item_template_var,
            width=34).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Label(
            inline_item_row,
            text="可用 {author}、{year}、{title}",
            foreground="#777777").pack(side=tk.LEFT)
        inline_options_row = ttk.Frame(scheme_box)
        inline_options_row.pack(fill=tk.X, pady=(0, 3))
        ttk.Checkbutton(
            inline_options_row, text="数字编号连续时压缩为范围（如 1,2,3 → 1–3）",
            variable=compress_ranges_var).pack(side=tk.LEFT)

        ttk.Label(
            scheme_box,
            text="格式模板（[[...]] 为可选块：块内字段为空时整块隐藏）：",
        ).pack(anchor=tk.W, pady=(6, 2))
        template_text = tk.Text(
            scheme_box, height=3, wrap=tk.WORD,
            font=("Consolas", 10), undo=True)
        template_text.pack(fill=tk.X)

        token_bar = ttk.Frame(scheme_box)
        token_bar.pack(fill=tk.X, pady=(4, 2))
        ttk.Label(token_bar, text="插入字段：").pack(side=tk.LEFT)
        token_var = tk.StringVar(value=AVAILABLE_FIELDS[0])
        token_combo = ttk.Combobox(
            token_bar, textvariable=token_var,
            values=AVAILABLE_FIELDS, state="readonly", width=23)
        token_combo.pack(side=tk.LEFT, padx=3)

        insert_box = ttk.LabelFrame(right, text="插入 Word / WPS", padding=6)
        insert_box.pack(fill=tk.X, pady=(8, 0))

        content_row = ttk.Frame(insert_box)
        content_row.pack(fill=tk.X, pady=2)
        ttk.Label(content_row, text="内容模式：").pack(side=tk.LEFT)
        ttk.Radiobutton(
            content_row, text="正文引用",
            variable=content_mode_var, value="inline"
        ).pack(side=tk.LEFT, padx=(0, 2))
        help_button(
            content_row, "正文引用模式",
            "在当前光标处插入 [1]、[1–3] 或 (作者, 年份) 等正文标记。"
            "数字编号会按当前文档分别记录，重复引用沿用原编号。"
        ).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Radiobutton(
            content_row, text="完整参考文献",
            variable=content_mode_var, value="bibliography"
        ).pack(side=tk.LEFT, padx=(0, 2))
        help_button(
            content_row, "完整参考文献模式",
            "按当前引用方案把作者、标题、期刊、年份、DOI 等完整条目插入"
            "光标位置；多篇按照左侧顺序一次性插入。"
        ).pack(side=tk.LEFT)

        privacy_row = ttk.Frame(insert_box)
        privacy_row.pack(fill=tk.X, pady=2)
        ttk.Label(privacy_row, text="记录模式：").pack(side=tk.LEFT)
        ttk.Radiobutton(
            privacy_row, text="隐私模式",
            variable=privacy_mode_var, value="privacy"
        ).pack(side=tk.LEFT, padx=(0, 2))
        help_button(
            privacy_row, "隐私模式",
            "默认推荐。paper_id、编号和插入历史只保存在 Lazybones 项目中，"
            "不向 Word/WPS 文档写入隐藏标记，适合投稿和对外共享。"
        ).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Radiobutton(
            privacy_row, text="智能编辑模式",
            variable=privacy_mode_var, value="smart"
        ).pack(side=tk.LEFT, padx=(0, 2))
        help_button(
            privacy_row, "智能编辑模式",
            "在文档内容控件中只写入无意义的随机 UUID，用于后续识别引用。"
            "该标记可被技术检查发现，投稿前应生成“投稿干净版”。"
        ).pack(side=tk.LEFT)

        target_row = ttk.Frame(insert_box)
        target_row.pack(fill=tk.X, pady=(3, 1))
        ttk.Label(target_row, text="目标：").pack(side=tk.LEFT)
        ttk.Combobox(
            target_row, textvariable=target_app_var,
            values=["auto", "word", "wps"],
            state="readonly", width=10).pack(side=tk.LEFT, padx=(0, 3))
        help_button(
            target_row, "目标程序",
            "auto 自动连接已打开的 Word，未找到时再连接 WPS；如果两者同时"
            "打开，请明确选择 word 或 wps，避免插入到错误文档。"
        ).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Label(
            target_row, textvariable=target_status_var,
            foreground="#1F4E79").pack(side=tk.LEFT, fill=tk.X, expand=True)

        preview_box = ttk.LabelFrame(right, text="实时预览", padding=6)
        preview_box.pack(fill=tk.BOTH, expand=True, pady=(8, 0))
        preview_text = scrolledtext.ScrolledText(
            preview_box, wrap=tk.WORD,
            font=("Times New Roman", 10), padx=8, pady=8)
        preview_text.pack(fill=tk.BOTH, expand=True)

        action_bar = ttk.Frame(win, padding=(10, 0, 10, 10))
        action_bar.pack(side=tk.BOTTOM, fill=tk.X, before=pane)

        def display_value(record, *keys):
            for key in keys:
                value = record.get(key)
                if value not in (None, "", "N/A", [], {}):
                    if isinstance(value, list):
                        return "; ".join(
                            str(item.get("value", ""))
                            if isinstance(item, dict) else str(item)
                            for item in value)
                    return str(value)
            return ""

        def refresh_order_tree(selected_index=None):
            order_tree.delete(*order_tree.get_children())
            for index, record in enumerate(ordered):
                order_tree.insert(
                    "", tk.END, iid=str(index),
                    values=(
                        index + 1,
                        display_value(record, "Title"),
                        display_value(
                            record, "First_Author", "Authors_Full", "Authors"),
                        display_value(record, "Year"),
                    ))
            count_var.set("共 " + str(len(ordered)) + " 篇")
            if ordered and selected_index is not None:
                selected_index = max(
                    0, min(selected_index, len(ordered) - 1))
                order_tree.selection_set(str(selected_index))
                order_tree.see(str(selected_index))
            update_preview()

        def change_citation_language(event=None):
            refreshed, _ = self._citation_source_records(
                citation_language_var.get())
            refreshed_by_base = {
                record.get("_citation_base_id"): record
                for record in refreshed
            }
            current_order = [
                record.get("_citation_base_id") for record in ordered]
            ordered[:] = [
                refreshed_by_base[base] for base in current_order
                if base in refreshed_by_base
            ]
            refresh_order_tree(0)

        citation_language_combo.bind(
            "<<ComboboxSelected>>", change_citation_language)

        def selected_index():
            selected = order_tree.selection()
            return int(selected[0]) if selected else None

        def move_selected(offset=None, absolute=None):
            index = selected_index()
            if index is None:
                return
            if absolute == "top":
                target = 0
            elif absolute == "bottom":
                target = len(ordered) - 1
            else:
                target = max(
                    0, min(len(ordered) - 1, index + int(offset or 0)))
            if target == index:
                return
            record = ordered.pop(index)
            ordered.insert(target, record)
            refresh_order_tree(target)

        def remove_selected():
            index = selected_index()
            if index is None:
                return
            ordered.pop(index)
            refresh_order_tree(min(index, len(ordered) - 1))

        ttk.Button(
            move_bar, text="置顶",
            command=lambda: move_selected(absolute="top")
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            move_bar, text="上移",
            command=lambda: move_selected(offset=-1)
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            move_bar, text="下移",
            command=lambda: move_selected(offset=1)
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            move_bar, text="置底",
            command=lambda: move_selected(absolute="bottom")
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            move_bar, text="移除",
            command=remove_selected
        ).pack(side=tk.LEFT, padx=8)

        def sort_records(event=None):
            field = sort_var.get()
            key_map = {
                "第一作者": ("First_Author", "Authors_Full", "Authors"),
                "论文名称": ("Title",),
                "期刊名称": ("Journal_Full", "Journal_Abbr", "Journal"),
                "年份（升序）": ("Year",),
                "年份（降序）": ("Year",),
            }
            keys = key_map.get(field)
            if not keys:
                return
            ordered.sort(
                key=lambda record: _normalize_search_text(
                    display_value(record, *keys)),
                reverse=(field == "年份（降序）"))
            refresh_order_tree(0)

        sort_combo.bind("<<ComboboxSelected>>", sort_records)

        drag_state = {"index": None}

        def drag_start(event):
            row = order_tree.identify_row(event.y)
            drag_state["index"] = int(row) if row else None

        def drag_end(event):
            source_index = drag_state.get("index")
            target_row = order_tree.identify_row(event.y)
            drag_state["index"] = None
            if source_index is None or not target_row:
                return
            target_index = int(target_row)
            if source_index == target_index:
                return
            record = ordered.pop(source_index)
            ordered.insert(target_index, record)
            refresh_order_tree(target_index)

        order_tree.bind("<ButtonPress-1>", drag_start, add="+")
        order_tree.bind("<ButtonRelease-1>", drag_end, add="+")

        def current_scheme() -> dict:
            return {
                "name": scheme_name_var.get().strip() or "临时方案",
                "number_format": number_var.get(),
                "template": template_text.get("1.0", tk.END).strip(),
                "separator": separator_var.get() or NL,
                "inline_style": (
                    "author_year"
                    if inline_style_var.get() == "作者-年份"
                    else "numeric"),
                "inline_template": inline_template_var.get(),
                "inline_item_template": inline_item_template_var.get(),
                "inline_separator": inline_separator_var.get(),
                "compress_ranges": compress_ranges_var.get(),
            }

        def rendered_lines():
            return format_citation_list(ordered, current_scheme())

        def rendered_text():
            if content_mode_var.get() == "inline":
                return format_inline_citation(ordered, current_scheme())
            return (separator_var.get() or NL).join(rendered_lines())

        def update_preview(*args):
            if not preview_text.winfo_exists():
                return
            preview_text.configure(state=tk.NORMAL)
            preview_text.delete("1.0", tk.END)
            if content_mode_var.get() == "inline":
                preview = format_inline_citation(
                    ordered, current_scheme())
            else:
                preview = (separator_var.get() or NL).join(rendered_lines())
            if preview:
                preview_text.insert("1.0", preview)
            else:
                preview_text.insert("1.0", "引用列表为空。")
            preview_text.configure(state=tk.DISABLED)

        def template_modified(event=None):
            if template_text.edit_modified():
                template_text.edit_modified(False)
                update_preview()

        template_text.bind("<<Modified>>", template_modified)
        number_var.trace_add("write", update_preview)
        for variable in (
                inline_style_var, inline_template_var,
                inline_item_template_var, inline_separator_var,
                compress_ranges_var, content_mode_var):
            variable.trace_add("write", update_preview)

        def schemes_by_name():
            return {
                scheme.get("name", scheme.get("id", "")): scheme
                for scheme in schemes_payload.get("schemes", [])
            }

        def refresh_scheme_combo(select_id=None):
            choices = schemes_payload.get("schemes", [])
            scheme_combo["values"] = [
                scheme.get("name", scheme.get("id", ""))
                for scheme in choices]
            selected = next(
                (scheme for scheme in choices
                 if scheme.get("id") == select_id),
                choices[0] if choices else DEFAULT_SCHEME)
            scheme_select_var.set(selected.get("name", ""))
            load_scheme(selected)

        def load_scheme(scheme):
            scheme = normalize_citation_scheme(scheme)
            current_scheme_id["value"] = scheme.get(
                "id", DEFAULT_SCHEME["id"])
            scheme_name_var.set(scheme.get("name", ""))
            number_var.set(scheme.get("number_format", "[{n}]"))
            separator_var.set(scheme.get("separator", NL))
            template_text.delete("1.0", tk.END)
            template_text.insert(
                "1.0", scheme.get("template", DEFAULT_TEMPLATE))
            inline_style_var.set(
                "作者-年份"
                if scheme.get("inline_style") == "author_year"
                else "数字编号")
            inline_template_var.set(
                scheme.get("inline_template", DEFAULT_INLINE_TEMPLATE))
            inline_item_template_var.set(scheme.get(
                "inline_item_template", DEFAULT_INLINE_ITEM_TEMPLATE))
            inline_separator_var.set(scheme.get("inline_separator", ","))
            compress_ranges_var.set(bool(
                scheme.get("compress_ranges", True)))
            template_text.edit_modified(False)
            update_preview()

        def on_scheme_selected(event=None):
            scheme = schemes_by_name().get(scheme_select_var.get())
            if scheme:
                load_scheme(scheme)

        scheme_combo.bind("<<ComboboxSelected>>", on_scheme_selected)

        def insert_token():
            template_text.insert(
                tk.INSERT, "{" + token_var.get() + "}")
            template_text.edit_modified(True)

        ttk.Button(
            token_bar, text="插入",
            command=insert_token).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            token_bar, text="插入可选块",
            command=lambda: (
                template_text.insert(
                    tk.INSERT, "[[{" + token_var.get() + "}]]"),
                template_text.edit_modified(True),
            )).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            token_bar, text="恢复完整默认格式",
            command=lambda: (
                template_text.delete("1.0", tk.END),
                template_text.insert("1.0", DEFAULT_TEMPLATE),
                template_text.edit_modified(True),
            )).pack(side=tk.RIGHT, padx=2)

        scheme_actions = ttk.Frame(scheme_box)
        scheme_actions.pack(fill=tk.X, pady=(5, 0))

        def save_scheme(as_new):
            name = scheme_name_var.get().strip()
            template = template_text.get("1.0", tk.END).strip()
            if not name or not template:
                messagebox.showwarning(
                    "引用方案", "方案名称和格式模板不能为空。", parent=win)
                return
            if template.count("[[") != template.count("]]"):
                messagebox.showwarning(
                    "引用方案",
                    "可选块标记不成对，请检查 [[ 和 ]]。",
                    parent=win)
                return
            if template.count("{") != template.count("}"):
                messagebox.showwarning(
                    "引用方案",
                    "字段标记不成对，请检查 { 和 }。",
                    parent=win)
                return
            scheme_id = None
            if not as_new:
                active = next(
                    (scheme for scheme in schemes_payload.get("schemes", [])
                     if scheme.get("id") == current_scheme_id["value"]),
                    None)
                if active and active.get("builtin"):
                    messagebox.showinfo(
                        "内置方案",
                        "内置方案不能覆盖，请使用“另存新方案”。",
                        parent=win)
                    return
                scheme_id = current_scheme_id["value"]
            try:
                saved = upsert_citation_scheme(
                    project_dir, schemes_payload, name,
                    number_var.get(), template,
                    separator_var.get() or NL, scheme_id=scheme_id,
                    inline_style=(
                        "author_year"
                        if inline_style_var.get() == "作者-年份"
                        else "numeric"),
                    inline_template=inline_template_var.get(),
                    inline_item_template=inline_item_template_var.get(),
                    inline_separator=inline_separator_var.get(),
                    compress_ranges=compress_ranges_var.get())
            except ValueError as error:
                messagebox.showwarning(
                    "引用方案", str(error), parent=win)
                return
            refresh_scheme_combo(saved["id"])
            messagebox.showinfo(
                "引用方案", "方案已保存，下次启动可继续使用。", parent=win)

        def remove_scheme():
            scheme_id = current_scheme_id["value"]
            if not delete_citation_scheme(
                    project_dir, schemes_payload, scheme_id):
                messagebox.showinfo(
                    "引用方案", "内置默认方案不能删除。", parent=win)
                return
            refresh_scheme_combo(DEFAULT_SCHEME["id"])

        def set_default_scheme():
            schemes_payload["default_scheme_id"] = current_scheme_id["value"]
            save_citation_schemes(project_dir, schemes_payload)
            messagebox.showinfo(
                "引用方案", "已设为下次打开时的默认方案。", parent=win)

        ttk.Button(
            scheme_actions, text="另存新方案",
            command=lambda: save_scheme(True)
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            scheme_actions, text="覆盖当前方案",
            command=lambda: save_scheme(False)
        ).pack(side=tk.LEFT, padx=2)
        ttk.Button(
            scheme_actions, text="删除当前方案",
            command=remove_scheme
        ).pack(side=tk.LEFT, padx=8)
        ttk.Button(
            scheme_actions, text="设为启动默认",
            command=set_default_scheme
        ).pack(side=tk.LEFT, padx=2)

        def detect_target(parent=win, quiet=False):
            try:
                target = office_bridge.get_target(target_app_var.get())
                target_status_var.set(
                    "● " + target.app_name + " — " + target.document_name)
                return target
            except OfficeBridgeError as error:
                target_status_var.set("○ 未连接 Word/WPS")
                if not quiet:
                    messagebox.showwarning(
                        "未连接文档", str(error), parent=parent)
                return None

        ttk.Button(
            target_row, text="检测",
            command=lambda: detect_target(win)
        ).pack(side=tk.RIGHT, padx=2)

        def prepare_insertion(records, target):
            paper_ids = [
                str(record.get("_citation_base_id", ""))
                for record in records
                if record.get("_citation_base_id")]
            session, number_map, token_map = session_store.prepare(
                target.public(), paper_ids)
            scheme = current_scheme()
            if content_mode_var.get() == "inline":
                text = format_inline_citation(
                    records, scheme, number_map=number_map)
            else:
                lines = [
                    format_citation(
                        record, scheme,
                        number_map.get(
                            str(record.get("_citation_base_id", "")), index))
                    for index, record in enumerate(records, 1)
                ]
                text = (scheme.get("separator") or NL).join(lines)
            return paper_ids, session, token_map, text

        def insert_records(records, parent=win, clear_callback=None):
            if not records:
                messagebox.showwarning(
                    "插入引用", "请先选择至少一篇文献。", parent=parent)
                return
            target = detect_target(parent)
            if target is None:
                return
            try:
                paper_ids, session, token_map, text = prepare_insertion(
                    records, target)
                smart_token = ""
                if privacy_mode_var.get() == "smart":
                    smart_token = session_store.create_smart_group(
                        target.public(), session, paper_ids)
                result = office_bridge.insert(
                    text, app_id=target.app_id,
                    smart_token=smart_token)
                session_store.record_insertion(
                    result["target"], session, paper_ids,
                    current_scheme_id["value"], content_mode_var.get(),
                    privacy_mode_var.get(), text)
                if (privacy_mode_var.get() == "smart"
                        and not result.get("smart_applied")):
                    messagebox.showwarning(
                        "已按隐私模式插入",
                        "当前 Word/WPS 版本不支持智能内容控件，引用已经作为普通"
                        "文字插入，文档中没有写入隐藏标记。",
                        parent=parent)
                self.status_var.set(
                    "已向 " + result["target"]["app_name"] + " 插入 " +
                    str(len(records)) + " 篇引用")
                target_status_var.set(
                    "● " + result["target"]["app_name"] + " — " +
                    result["target"]["document_name"])
                if clear_callback is not None:
                    clear_callback()
            except (OfficeBridgeError, OSError, ValueError) as error:
                messagebox.showerror(
                    "插入失败", str(error), parent=parent)

        def create_clean_submission_copy(parent=win):
            target = detect_target(parent)
            if target is None:
                return
            source_name = Path(
                target.document_name or "manuscript.docx").stem
            output = filedialog.asksaveasfilename(
                parent=parent,
                initialdir=str(Path(project_dir) / "exports"),
                initialfile=source_name + "_submission_clean.docx",
                defaultextension=".docx",
                filetypes=[("Word 文档", "*.docx")])
            if not output:
                return
            try:
                result = office_bridge.create_clean_copy(
                    output, target.app_id)
            except OfficeBridgeError as error:
                messagebox.showerror(
                    "投稿干净版", str(error), parent=parent)
                return
            if result["clean"]:
                messagebox.showinfo(
                    "投稿干净版已生成",
                    "文件：" + result["output_path"] + NL +
                    "已移除 Lazybones 标记：" + str(result["removed"]) +
                    " 个" + NL + "内部复检：未发现 Lazybones 隐藏标记。" +
                    NL2 + "建议投稿前再运行 Word 的“检查文档”。",
                    parent=parent)
            else:
                messagebox.showwarning(
                    "仍检测到隐藏标记",
                    "干净版已生成，但以下内部文件仍有 Lazybones 标记：" +
                    NL.join(result["marker_files"][:10]),
                    parent=parent)

        floating_window = {"value": None}

        def open_floating_inserter():
            existing = floating_window.get("value")
            try:
                if existing is not None and existing.winfo_exists():
                    existing.deiconify()
                    existing.lift()
                    return
            except tk.TclError:
                pass

            floating = tk.Toplevel(self)
            floating_window["value"] = floating
            floating.title("Lazybones 悬浮引用插入器")
            floating.geometry("620x640")
            floating.minsize(520, 500)
            floating.attributes("-topmost", True)

            float_header = ttk.Frame(floating, padding=8)
            float_header.pack(fill=tk.X)
            pin_var = tk.BooleanVar(value=True)
            ttk.Checkbutton(
                float_header, text="始终置顶", variable=pin_var,
                command=lambda: floating.attributes(
                    "-topmost", pin_var.get())
            ).pack(side=tk.LEFT)
            ttk.Button(
                float_header, text="展开完整页面",
                command=lambda: (win.deiconify(), win.lift())
            ).pack(side=tk.RIGHT)

            float_scheme_row = ttk.Frame(floating, padding=(8, 0))
            float_scheme_row.pack(fill=tk.X, pady=3)
            ttk.Label(float_scheme_row, text="引用方案：").pack(side=tk.LEFT)
            float_scheme_combo = ttk.Combobox(
                float_scheme_row, textvariable=scheme_select_var,
                values=list(schemes_by_name()),
                state="readonly", width=24)
            float_scheme_combo.pack(side=tk.LEFT, padx=3)
            float_scheme_combo.bind(
                "<<ComboboxSelected>>", on_scheme_selected)

            float_mode_box = ttk.LabelFrame(
                floating, text="插入方式", padding=6)
            float_mode_box.pack(fill=tk.X, padx=8, pady=4)
            row_one = ttk.Frame(float_mode_box)
            row_one.pack(fill=tk.X, pady=2)
            ttk.Radiobutton(
                row_one, text="正文引用", variable=content_mode_var,
                value="inline").pack(side=tk.LEFT)
            help_button(
                row_one, "正文引用模式",
                "插入 [1]、[1–3] 或 (作者, 年份)，重复文献沿用当前文档编号。"
            ).pack(side=tk.LEFT, padx=(2, 10))
            ttk.Radiobutton(
                row_one, text="完整参考文献", variable=content_mode_var,
                value="bibliography").pack(side=tk.LEFT)
            help_button(
                row_one, "完整参考文献模式",
                "按引用方案插入完整条目，多篇按照当前列表顺序排列。"
            ).pack(side=tk.LEFT, padx=2)
            row_two = ttk.Frame(float_mode_box)
            row_two.pack(fill=tk.X, pady=2)
            ttk.Radiobutton(
                row_two, text="隐私模式", variable=privacy_mode_var,
                value="privacy").pack(side=tk.LEFT)
            help_button(
                row_two, "隐私模式",
                "只在 Lazybones 项目保存引用记录，文档中没有隐藏 paper_id。"
            ).pack(side=tk.LEFT, padx=(2, 10))
            ttk.Radiobutton(
                row_two, text="智能编辑", variable=privacy_mode_var,
                value="smart").pack(side=tk.LEFT)
            help_button(
                row_two, "智能编辑模式",
                "文档只写随机 UUID，便于后续识别；投稿前应生成干净版。"
            ).pack(side=tk.LEFT, padx=2)

            float_target_row = ttk.Frame(floating, padding=(8, 0))
            float_target_row.pack(fill=tk.X, pady=3)
            ttk.Combobox(
                float_target_row, textvariable=target_app_var,
                values=["auto", "word", "wps"], state="readonly",
                width=9).pack(side=tk.LEFT)
            ttk.Label(
                float_target_row, textvariable=target_status_var,
                foreground="#1F4E79").pack(
                    side=tk.LEFT, fill=tk.X, expand=True, padx=6)
            ttk.Button(
                float_target_row, text="检测",
                command=lambda: detect_target(floating)
            ).pack(side=tk.RIGHT)

            search_var = tk.StringVar()
            search_entry = ttk.Entry(
                floating, textvariable=search_var)
            search_entry.pack(fill=tk.X, padx=8, pady=4)

            list_frame = ttk.Frame(floating)
            list_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=3)
            float_tree = ttk.Treeview(
                list_frame, columns=("title", "author", "year"),
                show="headings", selectmode=tk.EXTENDED)
            float_tree.heading("title", text="论文名称")
            float_tree.heading("author", text="第一作者")
            float_tree.heading("year", text="年份")
            float_tree.column("title", width=310)
            float_tree.column("author", width=120)
            float_tree.column("year", width=58, anchor=tk.CENTER)
            float_scroll = ttk.Scrollbar(
                list_frame, orient=tk.VERTICAL, command=float_tree.yview)
            float_tree.configure(yscrollcommand=float_scroll.set)
            float_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            float_scroll.pack(side=tk.RIGHT, fill=tk.Y)

            available_records = self._citation_library_records(
                citation_language_var.get())
            visible_records = []
            selected_float_ids = {
                str(record.get("_citation_base_id", ""))
                for record in ordered}
            selection_guard = {"active": False}

            def refresh_float_list(*args):
                query = _normalize_search_text(search_var.get())
                visible_records.clear()
                float_tree.delete(*float_tree.get_children())
                selected_rows = []
                for record in available_records:
                    haystack = _normalize_search_text(" ".join((
                        display_value(record, "Title"),
                        display_value(
                            record, "First_Author", "Authors_Full"),
                        display_value(record, "Year"),
                        display_value(record, "DOI"),
                        display_value(record, "Journal_Full"),
                    )))
                    if query and query not in haystack:
                        continue
                    index = len(visible_records)
                    visible_records.append(record)
                    float_tree.insert(
                        "", tk.END, iid=str(index), values=(
                            display_value(record, "Title"),
                            display_value(
                                record, "First_Author", "Authors_Full"),
                            display_value(record, "Year")))
                    if str(record.get(
                            "_citation_base_id", "")) in selected_float_ids:
                        selected_rows.append(str(index))
                if selected_rows:
                    selection_guard["active"] = True
                    float_tree.selection_set(selected_rows)
                    selection_guard["active"] = False

            def sync_order_from_float():
                lookup = {
                    str(record.get("_citation_base_id", "")): record
                    for record in available_records}
                current_ids = [
                    str(record.get("_citation_base_id", ""))
                    for record in ordered]
                next_ids = [
                    base_id for base_id in current_ids
                    if base_id in selected_float_ids]
                next_ids.extend(
                    base_id for base_id in lookup
                    if (base_id in selected_float_ids
                        and base_id not in next_ids))
                ordered[:] = [lookup[base_id] for base_id in next_ids]
                refresh_order_tree(0 if ordered else None)

            def on_float_selection(event=None):
                if selection_guard["active"]:
                    return
                visible_ids = {
                    str(record.get("_citation_base_id", ""))
                    for record in visible_records}
                selected_now = {
                    str(visible_records[int(item)].get(
                        "_citation_base_id", ""))
                    for item in float_tree.selection()
                    if int(item) < len(visible_records)}
                selected_float_ids.difference_update(visible_ids)
                selected_float_ids.update(selected_now)
                sync_order_from_float()

            def selected_float_records():
                return list(ordered)

            def select_all_visible():
                float_tree.selection_set(float_tree.get_children())
                on_float_selection()

            def clear_all_selection():
                selected_float_ids.clear()
                selection_guard["active"] = True
                float_tree.selection_remove(float_tree.selection())
                selection_guard["active"] = False
                sync_order_from_float()

            search_var.trace_add("write", refresh_float_list)
            float_tree.bind(
                "<<TreeviewSelect>>", on_float_selection, add="+")
            refresh_float_list()

            float_actions = ttk.Frame(floating, padding=8)
            float_actions.pack(
                side=tk.BOTTOM, fill=tk.X, before=list_frame)
            ttk.Button(
                float_actions, text="全选",
                command=select_all_visible
            ).pack(side=tk.LEFT, padx=2)
            ttk.Button(
                float_actions, text="全不选",
                command=clear_all_selection
            ).pack(side=tk.LEFT, padx=2)
            ttk.Button(
                float_actions, text="生成投稿干净版",
                command=lambda: create_clean_submission_copy(floating)
            ).pack(side=tk.RIGHT, padx=2)
            ttk.Button(
                float_actions, text="插入",
                style="Accent.TButton",
                command=lambda: insert_records(
                    selected_float_records(), floating)
            ).pack(side=tk.RIGHT, padx=2)
            ttk.Button(
                float_actions, text="插入并清空",
                command=lambda: insert_records(
                    selected_float_records(), floating,
                    clear_all_selection)
            ).pack(side=tk.RIGHT, padx=2)
            detect_target(floating, quiet=True)

        def copy_output():
            text = rendered_text()
            win.clipboard_clear()
            win.clipboard_append(text)
            win.update()
            self.status_var.set(
                "已复制 " + str(len(ordered)) + " 条引用到剪贴板")

        def export_text(extension):
            if not ordered:
                return
            filetypes = (
                [("Markdown", "*.md")] if extension == ".md"
                else [("文本", "*.txt")])
            path = filedialog.asksaveasfilename(
                parent=win,
                initialdir=str(Path(project_dir) / "exports"),
                initialfile=(
                    "citations_" + datetime.now().strftime("%Y%m%d") +
                    extension),
                defaultextension=extension,
                filetypes=filetypes)
            if not path:
                return
            Path(path).write_text(rendered_text(), encoding="utf-8")
            messagebox.showinfo(
                "导出成功", "已导出到：" + path, parent=win)

        def export_word():
            if not ordered:
                return
            path = filedialog.asksaveasfilename(
                parent=win,
                initialdir=str(Path(project_dir) / "exports"),
                initialfile=(
                    "citations_" + datetime.now().strftime("%Y%m%d") +
                    ".docx"),
                defaultextension=".docx",
                filetypes=[("Word 文档", "*.docx")])
            if not path:
                return
            try:
                from docx import Document
                document = Document()
                document.add_heading("References", level=1)
                for citation in rendered_lines():
                    document.add_paragraph(citation)
                document.save(path)
                messagebox.showinfo(
                    "导出成功", "已导出到：" + path, parent=win)
            except Exception as e:
                messagebox.showerror("Word 导出失败", str(e), parent=win)

        ttk.Button(
            action_bar, text="关闭",
            command=win.destroy).pack(side=tk.RIGHT, padx=3)
        ttk.Button(
            action_bar, text="导出 Word",
            command=export_word).pack(side=tk.RIGHT, padx=3)
        ttk.Button(
            action_bar, text="导出 Markdown",
            command=lambda: export_text(".md")
        ).pack(side=tk.RIGHT, padx=3)
        ttk.Button(
            action_bar, text="导出 TXT",
            command=lambda: export_text(".txt")
        ).pack(side=tk.RIGHT, padx=3)
        ttk.Button(
            action_bar, text="复制全部",
            command=copy_output).pack(side=tk.RIGHT, padx=3)
        ttk.Button(
            action_bar, text="📌 悬浮插入器",
            command=open_floating_inserter,
            style="Accent.TButton").pack(side=tk.LEFT, padx=3)
        ttk.Button(
            action_bar, text="插入当前列表",
            command=lambda: insert_records(ordered, win)
        ).pack(side=tk.LEFT, padx=3)
        ttk.Button(
            action_bar, text="生成投稿干净版",
            command=lambda: create_clean_submission_copy(win)
        ).pack(side=tk.LEFT, padx=3)
        help_button(
            action_bar, "投稿干净版",
            "另存一份文档，把 Lazybones 智能内容控件转换成普通文字并删除"
            "隐藏标记；不会覆盖正在编辑的原稿。生成后仍建议运行 Word 的"
            "“检查文档”。"
        ).pack(side=tk.LEFT, padx=2)

        default_scheme_id = schemes_payload.get(
            "default_scheme_id", DEFAULT_SCHEME["id"])
        refresh_scheme_combo(default_scheme_id)
        refresh_order_tree(0)

    # ════════════════════════════════════════════════════
    # AI 对话
    # ════════════════════════════════════════════════════
    def _ask_ai_find_papers(self):
        question = self.ask_var.get().strip()
        if not question:
            messagebox.showinfo("提示", "请先输入问题")
            return

        index = self._build_minimal_index()
        if not index:
            messagebox.showinfo("提示", "数据库为空,无法检索")
            return

        # ── 读取用户选择的返回数量 ──
        raw_hits = self.max_hits_var.get().strip()
        if raw_hits in ("全部", "all", "ALL", "*", ""):
            max_hits = 9999
            hits_desc = "全部相关文献(数据库中所有匹配的论文都要返回,不要遗漏)"
        else:
            try:
                max_hits = max(1, int(raw_hits))
                hits_desc = "最多 " + str(max_hits) + " 篇最相关的"
            except ValueError:
                messagebox.showwarning(
                    "输入有误",
                    "「返回数量」请填数字或选「全部」,已自动按 20 篇处理")
                max_hits = 20
                hits_desc = "最多 20 篇最相关的"

        prompt = (
            "下面是一份文献索引,每行格式为:[文献ID] | 字段名: 值 | ..." + NL
            + index + NL
            + "用户问题:" + question + NL
            + "请返回 " + hits_desc + " 文献编号,格式为 JSON 对象,例如:" + NL
            + '{"hits":["Wang_2023_ACSNano","Li_2024_AdvMater"],'
            + '"reason":"一句话说明命中理由"}。' + NL
            + "若没有相关文献,hits 返回空数组。只输出 JSON,不要其他文字。"
        )

        self.status_var.set("AI 正在检索…")

        def _worker():
            try:
                raw = quick_ask(
                    prompt,
                    self.settings["base_url"],
                    self.settings["api_key"],
                    self.settings["model"],
                    system=("你是严谨的文献检索助手,只根据给定索引回答,"
                            "不编造文献。" +
                            ("用户要求返回全部匹配项时,请尽你所能列出所有"
                             "相关文献,不要因为数量多而省略。"
                             if max_hits >= 9999 else "")),
                    timeout=float(self.settings.get(
                        "ai_timeout",
                        120 if max_hits >= 9999 else 60)),
                    max_retries=int(self.settings.get("ai_max_retries", 2)),
                )
                import re
                obj = None
                try:
                    obj = json.loads(raw)
                except Exception:
                    m = re.search(r'\{[\s\S]*\}', raw or "")
                    if m:
                        try:
                            obj = json.loads(m.group())
                        except Exception:
                            obj = None
                hits = obj.get("hits", []) if isinstance(obj, dict) else []
                reason = (obj.get("reason", "")
                          if isinstance(obj, dict) else "")
                self.after(0, lambda: self._show_ai_find_result(
                    hits, reason, raw))
            except Exception as e:
                self.after(0, lambda: messagebox.showerror(
                    "AI 检索失败", str(e)))
                self.after(0, lambda: self.status_var.set("就绪"))

        threading.Thread(target=_worker, daemon=True).start()

    def _render_markdown_to_text(self, text_widget, markdown_text: str):
        """
        把 Markdown 文本逐行渲染到 tkinter Text 控件里。
        支持:# ~ ###### 标题、**加粗**、*斜体*(已统一为加粗)、
              `代码`、```代码块```、> 引用、- 列表、1. 编号列表、
              | 表格 |、---、空行
        额外处理:
        - 全角星号 ＊ 自动归一化为 ASCII *
        - 行尾残留的孤立 * 或 # 自动剥离
        """
        import re

        # 全角星号 → 半角(中文模型常见错误)
        markdown_text = markdown_text.replace(chr(0xFF0A), "*")

        lines = markdown_text.split(NL)
        in_code_block = False
        in_table = False
        table_buffer = []

        def flush_table():
            if not table_buffer:
                return
            rows = []
            for raw in table_buffer:
                stripped = raw.strip().strip("|")
                cells = [c.strip() for c in stripped.split("|")]
                rows.append(cells)
            rows = [r for r in rows
                    if not all(re.fullmatch(r'[-:\s]*', c) for c in r)]
            if not rows:
                table_buffer.clear()
                return
            col_count = max(len(r) for r in rows)
            col_widths = [0] * col_count
            for r in rows:
                for i, c in enumerate(r):
                    col_widths[i] = max(col_widths[i], len(c))
            for ri, r in enumerate(rows):
                line_parts = []
                for i in range(col_count):
                    cell = r[i] if i < len(r) else ""
                    line_parts.append(cell.ljust(col_widths[i]))
                line = "  " + "  |  ".join(line_parts) + "  "
                tag = "md_table_header" if ri == 0 else "md_table_cell"
                text_widget.insert(tk.END, line + NL, tag)
            text_widget.insert(tk.END, NL)
            table_buffer.clear()

        def render_inline(line: str, base_tag: str = "ai"):
            # 单星号 *xxx* 也按加粗渲染(DeepSeek 经常用单星号)
            # 用更宽松的正则:允许中间含空格、中文标点,但不跨行
            pattern = re.compile(
                r'(\*\*[^\*\n]+?\*\*)|(\*[^\*\n]+?\*)|(`[^`\n]+?`)'
            )
            pos = 0
            for m in pattern.finditer(line):
                if m.start() > pos:
                    text_widget.insert(tk.END, line[pos:m.start()], base_tag)
                token = m.group()
                if token.startswith("**​") and token.endswith("​**"):
                    text_widget.insert(tk.END, token[2:-2], "md_bold")
                elif token.startswith("*") and token.endswith("*"):
                    text_widget.insert(tk.END, token[1:-1], "md_bold")
                elif token.startswith("`") and token.endswith("`"):
                    text_widget.insert(tk.END, token[1:-1], "md_code_inline")
                pos = m.end()
            if pos < len(line):
                # 兜底:剥掉残留的孤立 * 和 #(渲染不出来的标记符)
                tail = line[pos:]
                tail = re.sub(r'(?<![\*\w])\*(?!\*)', '', tail)
                tail = re.sub(r'(?<!\*)\*(?![\*\w])', '', tail)
                text_widget.insert(tk.END, tail, base_tag)

        for raw_line in lines:
            line = raw_line.rstrip()

            # 代码块
            if line.strip().startswith("```"):
                in_code_block = not in_code_block
                text_widget.insert(tk.END, NL)
                continue
            if in_code_block:
                text_widget.insert(tk.END, line + NL, "md_code_block")
                continue

            # 表格
            if line.strip().startswith("|") and "|" in line.strip()[1:]:
                in_table = True
                table_buffer.append(line)
                continue
            elif in_table:
                flush_table()
                in_table = False

            # 水平线
            if re.fullmatch(r'-{3,}|={3,}|\*{3,}', line.strip()):
                text_widget.insert(tk.END, "─" * 50 + NL, "md_hr")
                continue

            # 标题:支持 H1~H6(原来只支持 H1~H3,H4 以上会漏)
            h_match = re.match(r'^(#{1,6})\s+(.+)$', line)
            if h_match:
                level = len(h_match.group(1))
                content = h_match.group(2).rstrip("#").strip()
                # H1~H3 用专属样式,H4~H6 统一用 md_h3 样式(避免字号太小)
                tag = "md_h" + str(level) if level <= 3 else "md_h3"
                # 标题内部也可能有 **xxx**,先剥一下
                content = re.sub(r'\*\*([^\*]+?)\*\*', r'\1', content)
                content = re.sub(r'\*([^\*]+?)\*', r'\1', content)
                text_widget.insert(tk.END, content + NL, tag)
                continue

            # 引用
            if line.startswith(">"):
                content = line[1:].lstrip()
                text_widget.insert(tk.END, "│ ", "md_quote")
                render_inline(content, base_tag="md_quote")
                text_widget.insert(tk.END, NL)
                continue

            # 无序列表
            list_match = re.match(r'^(\s*)[-\*]\s+(.+)$', line)
            if list_match:
                indent = list_match.group(1)
                content = list_match.group(2)
                text_widget.insert(tk.END, indent + "• ", "md_list")
                render_inline(content, base_tag="md_list")
                text_widget.insert(tk.END, NL)
                continue

            # 有序列表
            num_match = re.match(r'^(\s*)(\d+)\.\s+(.+)$', line)
            if num_match:
                indent = num_match.group(1)
                num = num_match.group(2)
                content = num_match.group(3)
                text_widget.insert(tk.END,
                                   indent + num + ". ", "md_list")
                render_inline(content, base_tag="md_list")
                text_widget.insert(tk.END, NL)
                continue

            # 普通段落
            if line:
                render_inline(line, base_tag="ai")
                text_widget.insert(tk.END, NL)
            else:
                text_widget.insert(tk.END, NL)

        if in_table:
            flush_table()

    def _build_chat_context(self, db_path: str = None,
                        max_chars: int = 600000,
                        field_value_max: int = 800,
                        full_fields: bool = False,
                        prefer_lang: str = "en") -> str:
        """
        构建 Chat 上下文,带大小上限。
        full_fields=False(默认):只喂 KEY_FIELDS 白名单(快、省 token)
        full_fields=True:遍历每篇文献的所有非下划线字段(全、准、慢)
        """
        if db_path is None:
            db_path = self.db_path

        papers = get_all_papers(db_path)
        lang_papers = [p for p in papers if p.get("_lang") == prefer_lang]
        if not lang_papers:
            lang_papers = papers

        seen, unique = set(), []
        for p in lang_papers:
            base = p.get("_paper_id", "").rsplit("__", 1)[0]
            if base not in seen:
                seen.add(base)
                unique.append(p)


        # 动态从当前 schema 读取字段列表，不再硬编码领域专属字段
        try:
            current_ver = self.settings.get("schema_version", "")
            _schema = load_schema(current_ver) if current_ver else {}
            KEY_FIELDS = [f["key"] for f in _schema.get("fields", [])]
        except Exception:
            KEY_FIELDS = []

        mode_tag = "完整字段" if full_fields else "核心字段"
        header = ["以下是文献数据库（" + mode_tag + "模式），共 " +
                  str(len(unique)) + " 篇：", "=" * 60]
        out = list(header)
        included = 0
        total_chars = sum(len(s) + 1 for s in out)

        # 优先使用管理页同款的全局序号映射,保证 AI 编号与 UI 一致
        seq_map = getattr(self, "_paper_seq_map", {})

        for p in unique:
            base_id = p.get("_paper_id", "").rsplit("__", 1)[0]
            seq = seq_map.get(base_id)
            if seq is None:
                seq = unique.index(p) + 1
            block = ["[" + str(seq) + "] " + base_id]

            if full_fields:
                field_keys = [k for k in p.keys() if not k.startswith("_")]
            else:
                field_keys = KEY_FIELDS if KEY_FIELDS else [
                    k for k in p.keys() if not k.startswith("_")]

            for key in field_keys:
                val = p.get(key, "N/A")
                if val in ("N/A", "", None):
                    continue
                if isinstance(val, list):
                    tags = []
                    for it in val:
                        if isinstance(it, dict):
                            v = str(it.get("value", ""))
                            ev = it.get("evidence", "") if full_fields else ""
                            if ev:
                                tags.append(v + "(证据:" + str(ev)[:120] + ")")
                            else:
                                tags.append(v)
                        else:
                            tags.append(str(it))
                    val = "; ".join(t for t in tags if t)
                else:
                    val = str(val)
                if len(val) > field_value_max:
                    val = val[:field_value_max] + "…"
                block.append("  " + key + ": " + val)
            block.append("")

            block_chars = sum(len(s) + 1 for s in block)
            if total_chars + block_chars > max_chars:
                out.append(
                    "[已省略剩余 " + str(len(unique) - included) +
                    " 篇文献,上下文已达上限 " + str(max_chars) +
                    " 字符;如需更详细分析,请用筛选缩小数据库范围]")
                break

            out.extend(block)
            total_chars += block_chars
            included += 1

        return NL.join(out)

    def _open_chat_window(self,
                          preset_context: str = None,
                          preset_title: str = None,
                          preset_session_name: str = None,
                          lock_scope: bool = False):
        """
        AI 对话窗口。
        参数:
          preset_context      预设上下文(B 模式:单篇文献的全部字段)
          preset_title        预设窗口标题
          preset_session_name 预设 session 名称
          lock_scope          是否锁定数据范围下拉框(B 模式用)
        """
        project_dir = self.settings.get("project_dir", "")
        chat_db = init_chat_db(project_dir)

        win = tk.Toplevel(self)
        win.title(preset_title or "AI 数据分析对话")
        win.geometry("1100x700")

        left = ttk.LabelFrame(win, text="对话记录", padding=4)
        left.pack(side=tk.LEFT, fill=tk.Y, padx=8, pady=8)

        btn_row = ttk.Frame(left)
        btn_row.pack(fill=tk.X, pady=4)

        session_listbox = tk.Listbox(left, width=24, height=30,
                                     font=("微软雅黑", 9))
        session_listbox.pack(fill=tk.BOTH, expand=True)

        right = ttk.Frame(win)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=4, pady=8)

        scope_frame = ttk.LabelFrame(right, text="数据范围", padding=4)
        scope_frame.pack(fill=tk.X, pady=4)
        scope_var = tk.StringVar(value=self.db_name_var.get())
        ttk.Label(scope_frame, text="读取数据库:").pack(side=tk.LEFT)
        db_list = list_databases(project_dir)
        scope_combo = ttk.Combobox(scope_frame, textvariable=scope_var,
                                   values=db_list, width=20,
                                   state="readonly")
        scope_combo.pack(side=tk.LEFT, padx=4)
        ttk.Label(scope_frame, text="切换后下次发消息生效",
                  foreground="#888").pack(side=tk.LEFT, padx=8)

        # ✅ 修复 Bug 1:B 模式锁定数据范围
        if lock_scope:
            scope_combo.configure(state=tk.DISABLED)

        # 完整字段开关:勾上后把每篇的所有字段都喂给 AI(更准但更慢/更贵)
        full_fields_var = tk.BooleanVar(value=False)
        full_fields_chk = ttk.Checkbutton(
            scope_frame, text="包含全部字段(更准但更慢)",
            variable=full_fields_var,
            command=lambda: (state.__setitem__("context_loaded", False),
                             append_display(
                                 "sys",
                                 "字段模式已切换为「" +
                                 ("完整字段" if full_fields_var.get()
                                  else "核心字段") +
                                 "」,下次发消息将重新加载")))
        full_fields_chk.pack(side=tk.LEFT, padx=8)


        chat_display = scrolledtext.ScrolledText(
            right, font=("微软雅黑", 10), wrap=tk.WORD,
            bg="#FAFAFA", state=tk.DISABLED)
        chat_display.pack(fill=tk.BOTH, expand=True, pady=4)
        chat_display.tag_config("user", foreground="#0070C0",
                                font=("微软雅黑", 10, "bold"))
        chat_display.tag_config("ai", foreground="#1A1A1A",
                                font=("微软雅黑", 10))
        chat_display.tag_config("sys", foreground="#999",
                                font=("微软雅黑", 8, "italic"))
        chat_display.tag_config("thinking", foreground="#888",
                                font=("微软雅黑", 9, "italic"),
                                background="#F0F0F0")
        chat_display.tag_config("thinking_header", foreground="#555",
                                font=("微软雅黑", 9, "bold"))
        # Markdown 渲染用的 tag
        chat_display.tag_config("md_h1", foreground="#1F4E79",
                                font=("微软雅黑", 14, "bold"),
                                spacing1=8, spacing3=4)
        chat_display.tag_config("md_h2", foreground="#2E75B6",
                                font=("微软雅黑", 12, "bold"),
                                spacing1=6, spacing3=3)
        chat_display.tag_config("md_h3", foreground="#5B9BD5",
                                font=("微软雅黑", 11, "bold"),
                                spacing1=4, spacing3=2)
        chat_display.tag_config("md_bold",
                                font=("微软雅黑", 10, "bold"),
                                foreground="#1A1A1A")
        chat_display.tag_config("md_italic",
                                font=("微软雅黑", 10, "italic"),
                                foreground="#1A1A1A")
        chat_display.tag_config("md_code_inline",
                                font=("Consolas", 9),
                                background="#F4F4F4",
                                foreground="#C7254E")
        chat_display.tag_config("md_code_block",
                                font=("Consolas", 9),
                                background="#F8F8F8",
                                foreground="#333",
                                lmargin1=20, lmargin2=20,
                                spacing1=4, spacing3=4)
        chat_display.tag_config("md_quote",
                                foreground="#666",
                                font=("微软雅黑", 10, "italic"),
                                lmargin1=20, lmargin2=20,
                                background="#FAFAFA")
        chat_display.tag_config("md_list",
                                font=("微软雅黑", 10),
                                lmargin1=20, lmargin2=40)
        chat_display.tag_config("md_table_header",
                                font=("微软雅黑", 9, "bold"),
                                background="#E7E7E7",
                                foreground="#1A1A1A")
        chat_display.tag_config("md_table_cell",
                                font=("Consolas", 9),
                                background="#FAFAFA",
                                foreground="#1A1A1A")
        chat_display.tag_config("md_hr",
                                foreground="#CCC",
                                font=("微软雅黑", 8))


        input_frame = ttk.Frame(right)
        input_frame.pack(fill=tk.X, pady=4)
        chat_entry = tk.Text(input_frame, height=3,
                             font=("微软雅黑", 10), wrap=tk.WORD)
        chat_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 4))
        send_col = ttk.Frame(input_frame)
        send_col.pack(side=tk.LEFT)

        # ✅ 修复 Bug 2:state 字典里加 preset_context
        state = {"session_id": None, "messages": [],
                 "context_loaded": False,
                 "preset_context": preset_context}
        sessions_cache = []

        def append_display(role, text, end=None):
            """
            role: user / ai / sys / thinking / thinking_header / ai_stream
            end:  追加完后的结尾,默认换行;流式输出(ai_stream/thinking)用空串
            """
            if end is None:
                end = "" if role in ("ai_stream", "thinking") else NL
            chat_display.configure(state=tk.NORMAL)
            if role == "user":
                chat_display.insert(tk.END, "你:" + text + end, "user")
            elif role == "ai":
                chat_display.insert(tk.END, "AI:" + text + end, "ai")
            elif role == "ai_stream":
                chat_display.insert(tk.END, text, "ai")
            elif role == "thinking":
                chat_display.insert(tk.END, text, "thinking")
            elif role == "thinking_header":
                chat_display.insert(tk.END, text + NL, "thinking_header")
            else:
                chat_display.insert(tk.END, text + end, "sys")
            chat_display.see(tk.END)
            chat_display.configure(state=tk.DISABLED)


        def refresh_sessions():
            session_listbox.delete(0, tk.END)
            sessions_cache.clear()
            for s in list_chat_sessions(chat_db):
                sessions_cache.append(s)
                session_listbox.insert(
                    tk.END, s["title"][:20] + "  " + s["updated_at"][:10])

        def new_session():
            name = simpledialog.askstring(
                "新建对话", "对话名称:",
                initialvalue="Chat_" + datetime.now().strftime("%m%d_%H%M"))
            if not name:
                return
            sid = str(uuid.uuid4())[:8]
            save_chat_session(chat_db, sid, name, scope_var.get())
            refresh_sessions()
            session_listbox.selection_clear(0, tk.END)
            session_listbox.selection_set(0)
            load_session()

        def delete_session():
            sel = session_listbox.curselection()
            if not sel:
                return
            s = sessions_cache[sel[0]]
            if not messagebox.askyesno(
                    "确认", "删除:" + s["title"] + "?"):
                return
            delete_chat_session(chat_db, s["session_id"])
            state["session_id"] = None
            state["messages"] = []
            chat_display.configure(state=tk.NORMAL)
            chat_display.delete("1.0", tk.END)
            chat_display.configure(state=tk.DISABLED)
            refresh_sessions()

        def load_session():
            sel = session_listbox.curselection()
            if not sel:
                return
            s = sessions_cache[sel[0]]
            state["session_id"] = s["session_id"]
            state["context_loaded"] = False
            scope_var.set(s["db_scope"])
            chat_display.configure(state=tk.NORMAL)
            chat_display.delete("1.0", tk.END)
            chat_display.configure(state=tk.DISABLED)
            history = load_chat_messages(chat_db, s["session_id"])
            state["messages"] = []
            for msg in history:
                if msg["role"] == "user":
                    append_display("user", msg["content"])
                    state["messages"].append(msg)
                elif msg["role"] == "assistant":
                    chat_display.configure(state=tk.NORMAL)
                    chat_display.insert(tk.END, "AI:", "ai")
                    self._render_markdown_to_text(chat_display, msg["content"])
                    chat_display.insert(tk.END, chr(10))
                    chat_display.configure(state=tk.DISABLED)
                    state["messages"].append(msg)
            if history:
                state["context_loaded"] = True

        ttk.Button(btn_row, text="➕ 新建",
                   command=new_session).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_row, text="🗑 删除",
                   command=delete_session).pack(side=tk.LEFT, padx=2)
        session_listbox.bind("<<ListboxSelect>>", lambda e: load_session())

        def do_send():
            if not state["session_id"]:
                messagebox.showwarning("提示", "请先新建或选择对话")
                return
            question = chat_entry.get("1.0", tk.END).strip()
            if not question:
                return
            if not self.settings.get("api_key"):
                messagebox.showerror("错误", "请先填写API Key")
                return
            chat_entry.delete("1.0", tk.END)
            append_display("user", question)
            save_chat_message(chat_db, state["session_id"],
                              "user", question)
            threading.Thread(target=run_ai, args=(question,),
                             daemon=True).start()

        def run_ai(question):
            from ai_client import get_client

            # ── 上下文加载 ──
            if not state["context_loaded"]:
                full = full_fields_var.get()
                if state.get("preset_context"):
                    # B 模式:单篇文献,preset_context 已经预先构造好,直接用
                    context = state["preset_context"]
                    db_name = "单篇文献"
                else:
                    db_name = scope_var.get()
                    target_db = get_db_path_by_name(project_dir, db_name)
                    init_db(target_db)
                    context = self._build_chat_context(
                        db_path=target_db,
                        max_chars=1000000 if full else 60000,
                        full_fields=full)

                est = int(len(context) * 0.4)
                mode_label = "完整字段" if full else "核心字段"
                win.after(0, lambda: append_display(
                    "sys",
                    "(已加载「" + db_name + "」· " + mode_label +
                    " · 约 " + str(est) + " tokens)"))
                state["messages"] = [{
                    "role": "system",
                    "content": ("你是一名专业的学术文献数据分析助手。"
                                "基于以下文献数据回答问题,"
                                "回答要紧扣给定数据,不要编造未提及的内容。"
                                + chr(10) +
                                "重要:每篇文献开头的 [N] 是该文献在用户数据库管理页的全局序号,"
                                "你引用具体文献时必须用「文献 [N]」或「#N」的格式,"
                                "这样用户可以直接在数据库管理页找到对应行。"
                                "不要用其他临时编号。"
                                + chr(10) + chr(10) +
                                context),
                }]
                state["context_loaded"] = True

            state["messages"].append({"role": "user", "content": question})

            # ── 显示"AI 正在思考..."占位 ──
            win.after(0, lambda: append_display("sys", "AI 正在思考..."))

            try:
                client = get_client(
                    self.settings["base_url"],
                    self.settings["api_key"],
                    timeout=float(self.settings.get("ai_timeout", 600)),
                )

                stream = client.chat.completions.create(
                    model=self.settings["model"],
                    messages=state["messages"],
                    temperature=0.3,
                    stream=True,
                )

                display_state = {
                    "ai_prefix_shown": False,
                    "thinking_header_shown": False,
                    "thinking_content": "",
                    "answer_content": "",
                    "answer_start_index": None,
                }

                def safe_append(role, text, end=""):
                    win.after(0, lambda: append_display(role, text, end))

                for chunk in stream:
                    if not chunk.choices:
                        continue
                    delta = chunk.choices[0].delta

                    # ── 思考过程(reasoning_content)──
                    reasoning = getattr(delta, "reasoning_content", None)
                    if reasoning:
                        if not display_state["thinking_header_shown"]:
                            safe_append("thinking_header", "💭 思考过程:")
                            display_state["thinking_header_shown"] = True
                        display_state["thinking_content"] += reasoning
                        safe_append("thinking", reasoning)

                    # ── 正文回答(流式显示)──
                    content = getattr(delta, "content", None)
                    if content:
                        if not display_state["ai_prefix_shown"]:
                            if display_state["thinking_header_shown"]:
                                safe_append("sys", "─── 回答 ───")

                            def _mark_start():
                                chat_display.configure(state=tk.NORMAL)
                                chat_display.insert(tk.END, "AI:", "ai")
                                display_state["answer_start_index"] = \
                                    chat_display.index(tk.END + "-1c")
                                chat_display.configure(state=tk.DISABLED)
                            win.after(0, _mark_start)
                            display_state["ai_prefix_shown"] = True
                        display_state["answer_content"] += content
                        safe_append("ai_stream", content)

                # ── 流结束:用 Markdown 重新渲染 ──
                full_answer = display_state["answer_content"]

                def _rerender_markdown():
                    if not display_state["answer_start_index"]:
                        return
                    chat_display.configure(state=tk.NORMAL)
                    chat_display.delete(
                        display_state["answer_start_index"], tk.END)
                    self._render_markdown_to_text(chat_display, full_answer)
                    chat_display.insert(tk.END, chr(10))
                    chat_display.see(tk.END)
                    chat_display.configure(state=tk.DISABLED)
                win.after(0, _rerender_markdown)

                state["messages"].append({"role": "assistant",
                                          "content": full_answer})
                save_chat_message(chat_db, state["session_id"],
                                  "assistant", full_answer)
                win.after(0, refresh_sessions)

            except Exception as e:
                err_msg = "调用失败:" + type(e).__name__ + ": " + str(e)
                win.after(0, lambda m=err_msg: append_display("sys", m))


        def on_scope_change(event=None):
            state["context_loaded"] = False
            append_display(
                "sys",
                "数据库已切换为「" + scope_var.get() +
                "」,下次发消息将重新加载")

        ttk.Button(send_col, text="发送",
                   command=do_send).pack(pady=2)
        ttk.Button(send_col, text="清空输入",
                   command=lambda: chat_entry.delete("1.0", tk.END)
                   ).pack(pady=2)
        chat_entry.bind("<Control-Return>", lambda e: do_send())
        scope_combo.bind("<<ComboboxSelected>>", on_scope_change)

        refresh_sessions()

        # ✅ 修复 Bug 3:B 模式自动新建 session 并加载
        if preset_session_name:
            sid = str(uuid.uuid4())[:8]
            save_chat_session(chat_db, sid, preset_session_name,
                              scope_var.get())
            refresh_sessions()
            session_listbox.selection_clear(0, tk.END)
            session_listbox.selection_set(0)
            load_session()

    # ════════════════════════════════════════════════════
    # B. 单篇问 AI
    # ════════════════════════════════════════════════════

    def _build_single_paper_context(self, db_path: str,
                                    base_paper_id: str) -> tuple:
        """
        构造单篇文献上下文。
        """
        data = (get_paper(db_path, base_paper_id + "__zh")
                or get_paper(db_path, base_paper_id + "__en")
                or get_paper(db_path, base_paper_id))
        if not data:
            return None, base_paper_id

        lines = ["以下是单篇文献的全部抽取结果:",
                 "文献ID: " + base_paper_id,
                 "=" * 40]
        for k, v in data.items():
            if k.startswith("_"):
                continue
            if v in ("N/A", "", None):
                continue
            if isinstance(v, list):
                rendered = []
                for it in v:
                    if isinstance(it, dict):
                        val = str(it.get("value", ""))
                        src = it.get("source", "")
                        conf = it.get("confidence", "")
                        ev = it.get("evidence", "")
                        suffix = []
                        if src:
                            suffix.append("src=" + str(src))
                        if conf:
                            suffix.append("conf=" + str(conf))
                        tag = " [" + ", ".join(suffix) + "]" if suffix else ""
                        line = "- " + val + tag
                        if ev:
                            line += "  证据:" + str(ev)
                        rendered.append(line)
                    else:
                        rendered.append("- " + str(it))
                lines.append(k + ":")
                lines.extend("  " + r for r in rendered)
            else:
                lines.append(k + ": " + str(v))
        return NL.join(lines), base_paper_id

    def _open_single_chat(self):
        """单篇深聊:要求先在管理页选中一行"""
        selected_ids = self._selected_manage_ids()
        if not selected_ids:
            messagebox.showwarning("提示", "请先在表格中选中一篇文献")
            return
        full_id = selected_ids[0]
        base_paper_id = full_id.rsplit("__", 1)[0]

        context, display_id = self._build_single_paper_context(
            self.db_path, base_paper_id)
        if context is None:
            messagebox.showerror(
                "未找到", f"数据库中找不到 {base_paper_id} 的抽取结果")
            return

        self._open_chat_window(
            preset_context=context,
            preset_title=f"单篇 · {display_id}",
            preset_session_name=f"paper:{display_id}_"
                                + datetime.now().strftime("%m%d_%H%M"),
            lock_scope=True,
        )

    # ════════════════════════════════════════════════════
    # E. 字段问答(检索式)
    # ════════════════════════════════════════════════════

    def _build_minimal_index(self, db_path: str = None) -> str:
        """
        构造检索索引:每篇一行,包含 ID + 所有非空字段值。
        动态读取,不依赖任何领域专属字段名。
        """
        if db_path is None:
            db_path = self.db_path
        papers = get_all_papers(db_path)

        seen, unique = set(), []
        for p in papers:
            base = p.get("_paper_id", "").rsplit("__", 1)[0]
            if base not in seen:
                seen.add(base)
                unique.append(p)

        def _flat(v):
            if isinstance(v, list):
                parts = []
                for it in v:
                    if isinstance(it, dict):
                        parts.append(str(it.get("value", "")))
                    else:
                        parts.append(str(it))
                return "; ".join(p for p in parts if p)
            return "" if v in ("N/A", "", None) else str(v)

        # 优先展示的身份字段(如果存在)
        PRIORITY_KEYS = ["Title", "First_Author", "Year",
                         "Journal_Full", "Journal_Abbr", "Keywords",
                         "Abstract_Short"]

        seq_map = getattr(self, "_paper_seq_map", {})
        lines = []
        for p in unique:
            base_id = p.get("_paper_id", "").rsplit("__", 1)[0]
            seq = seq_map.get(base_id)
            if seq is None:
                seq = len(lines) + 1
            # 双重标识:序号 + base_id,AI 不论用哪个引用都能对上
            parts = ["[#" + str(seq) + " | " + base_id + "]"]

            # 先输出优先字段
            for key in PRIORITY_KEYS:
                val = _flat(p.get(key))
                if val:
                    parts.append(key + ": " + val[:100])

            # 再输出其余非空字段(排除元数据和已输出的)
            skip = set(PRIORITY_KEYS) | {
                k for k in p.keys() if k.startswith("_")}
            for key, raw in p.items():
                if key in skip:
                    continue
                val = _flat(raw)
                if val:
                    parts.append(key + ": " + val[:120])

            lines.append(" | ".join(parts))

        return NL.join(lines)

    def _show_ai_find_result(self, hits: list, reason: str, raw: str):
        """AI 检索结果高亮:支持 base_id 和 #序号 两种引用方式"""
        if not hits:
            self.status_var.set("AI 未找到相关文献")
            messagebox.showinfo(
                "未命中",
                "AI 未找到匹配文献。原始返回:" + (raw or "")[:500])
            return

        # 解析 hits:可能是 "Wang_2023_ACSNano" 或 "#42" 两种格式
        target_bases = set()
        target_seqs = set()
        for h in hits:
            h = str(h).strip()
            if h.startswith("#"):
                try:
                    target_seqs.add(int(h.lstrip("#")))
                except ValueError:
                    pass
            else:
                target_bases.add(h)

        # 反查序号→base_id
        seq_map = getattr(self, "_paper_seq_map", {})
        seq_to_base = {v: k for k, v in seq_map.items()}
        for s in target_seqs:
            if s in seq_to_base:
                target_bases.add(seq_to_base[s])

        # 命中项可能分布在多页；跨页选中并自动跳到第一个命中页。
        matched_bases = self._select_manage_base_ids(target_bases)
        miss = [h for h in hits
                if h not in matched_bases
                and not (h.startswith("#")
                         and seq_to_base.get(
                             int(h.lstrip("#")) if h.lstrip("#").isdigit()
                             else -1) in matched_bases)]

        msg = ("AI 命中 " + str(len(matched_bases))
               + " / " + str(len(hits)) + " 篇")
        if reason:
            msg += chr(10) + "理由:" + reason
        if miss:
            msg += chr(10) + "未在当前视图找到:" + chr(10) + "- " + (chr(10) + "- ").join(miss)
        self.status_var.set("AI 检索完成 · 命中 " + str(len(matched_bases)) + " 篇")
        messagebox.showinfo("AI 检索结果", msg)

    # ════════════════════════════════════════════════════
    # 设置页
    # ════════════════════════════════════════════════════

    def _make_scrollable(self, parent):
        """
        在 parent 里建一个垂直滚动的容器,返回内部可放控件的 frame。
        使用方式:scrollable = self._make_scrollable(self.tab_xxx)
                  ttk.Label(scrollable, text="...").pack(...)
        """
        # 外层容器,装 canvas 和 scrollbar
        outer = ttk.Frame(parent)
        outer.pack(fill=tk.BOTH, expand=True)

        canvas = tk.Canvas(outer, highlightthickness=0,
                           background="#F0F0F0")
        vsb = ttk.Scrollbar(outer, orient=tk.VERTICAL,
                            command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)

        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        # 内部真正放控件的 frame
        inner = ttk.Frame(canvas)
        inner_id = canvas.create_window((0, 0), window=inner,
                                        anchor=tk.NW)

        def _on_inner_config(event):
            # 内部 frame 大小变化时,更新 canvas 的滚动区域
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _on_canvas_config(event):
            # canvas 宽度变化时,同步内部 frame 宽度,避免出现横向空白
            canvas.itemconfigure(inner_id, width=event.width)

        inner.bind("<Configure>", _on_inner_config)
        canvas.bind("<Configure>", _on_canvas_config)

        # 鼠标滚轮支持(Windows/macOS 通用)
        def _on_mousewheel(event):
            # Windows: event.delta 一般是 ±120 的倍数
            # macOS:   event.delta 是 ±1 的倍数
            delta = -1 * (event.delta // 120) if event.delta else 0
            if delta == 0 and event.delta:
                delta = -1 if event.delta > 0 else 1
            canvas.yview_scroll(delta, "units")

        # 鼠标进入时绑定滚轮,离开时解绑(避免多个 tab 滚轮冲突)
        def _bind_wheel(event):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)

        def _unbind_wheel(event):
            canvas.unbind_all("<MouseWheel>")

        canvas.bind("<Enter>", _bind_wheel)
        canvas.bind("<Leave>", _unbind_wheel)
        inner.bind("<Enter>", _bind_wheel)
        inner.bind("<Leave>", _unbind_wheel)

        return inner


    def _build_settings_tab(self):
        frame = self._make_scrollable(self.tab_settings)

        appearance = normalize_appearance(self.settings)
        appearance_frame = ttk.LabelFrame(
            frame, text="外观与皮肤", padding=10)
        appearance_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(appearance_frame, text="主题:").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        self.theme_var = tk.StringVar(value=THEME_LABELS[
            appearance["theme"]])
        theme_combo = ttk.Combobox(
            appearance_frame, textvariable=self.theme_var,
            values=list(THEME_LABELS.values()), width=20, state="readonly")
        theme_combo.grid(row=0, column=1, sticky=tk.W, padx=4)
        theme_combo.bind("<<ComboboxSelected>>", self._preview_theme)

        ttk.Label(appearance_frame, text="字体大小:").grid(
            row=0, column=2, sticky=tk.W, padx=(18, 0), pady=4)
        self.font_scale_var = tk.StringVar(
            value=str(appearance["font_scale"]) + "%")
        font_combo = ttk.Combobox(
            appearance_frame, textvariable=self.font_scale_var,
            values=["90%", "100%", "110%", "125%"],
            width=10, state="readonly")
        font_combo.grid(row=0, column=3, sticky=tk.W, padx=4)
        font_combo.bind("<<ComboboxSelected>>", self._preview_theme)

        ttk.Label(appearance_frame, text="界面密度:").grid(
            row=0, column=4, sticky=tk.W, padx=(18, 0), pady=4)
        self.ui_density_var = tk.StringVar(value=DENSITY_LABELS[
            appearance["ui_density"]])
        density_combo = ttk.Combobox(
            appearance_frame, textvariable=self.ui_density_var,
            values=list(DENSITY_LABELS.values()), width=10,
            state="readonly")
        density_combo.grid(row=0, column=5, sticky=tk.W, padx=4)
        density_combo.bind("<<ComboboxSelected>>", self._preview_theme)

        ttk.Button(appearance_frame, text="恢复经典外观",
                   command=self._reset_theme_preview).grid(
                       row=0, column=6, sticky=tk.W, padx=(18, 4))
        self.appearance_status_var = tk.StringVar()
        appearance_status = tk.Label(
            appearance_frame, textvariable=self.appearance_status_var,
            anchor=tk.W)
        appearance_status._theme_role = "muted"
        appearance_status.grid(
            row=1, column=0, columnspan=7, sticky=tk.W, pady=(6, 2))
        appearance_hint = tk.Label(
            appearance_frame,
            text="选择后立即预览；点击页面下方“保存设置”后永久生效。"
                 "主题只改变显示，不修改文献和数据库。",
            anchor=tk.W)
        appearance_hint._theme_role = "muted"
        appearance_hint.grid(
            row=2, column=0, columnspan=7, sticky=tk.W, pady=(0, 2))
        self._update_appearance_status()

        api_frame = ttk.LabelFrame(frame, text="API 配置", padding=10)
        api_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(api_frame, text="服务商:").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        self.provider_var = tk.StringVar(
            value=self.settings.get("provider", "DeepSeek"))
        provider_combo = ttk.Combobox(
            api_frame, textvariable=self.provider_var,
            values=list(PROVIDERS.keys()),
            width=30, state="readonly")
        provider_combo.grid(row=0, column=1, sticky=tk.W, pady=4, padx=4)
        provider_combo.bind("<<ComboboxSelected>>", self._on_provider_change)

        ttk.Label(api_frame, text="API Key:").grid(
            row=1, column=0, sticky=tk.W, pady=4)
        self.api_key_var = tk.StringVar(
            value=self.settings.get("api_key", ""))
        self.api_key_entry = ttk.Entry(
            api_frame, textvariable=self.api_key_var,
            width=60, show="*")
        self.api_key_entry.grid(row=1, column=1, sticky=tk.W, pady=4, padx=4)
        self.show_key_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(api_frame, text="显示",
                        variable=self.show_key_var,
                        command=self._toggle_key_visibility
                        ).grid(row=1, column=2, padx=4)

        ttk.Label(api_frame, text="Base URL:").grid(
            row=2, column=0, sticky=tk.W, pady=4)
        self.base_url_var = tk.StringVar(
            value=self.settings.get("base_url", ""))
        ttk.Entry(api_frame, textvariable=self.base_url_var,
                  width=60).grid(row=2, column=1, sticky=tk.W, pady=4, padx=4)

        ttk.Label(api_frame, text="模型:").grid(
            row=3, column=0, sticky=tk.W, pady=4)
        self.model_var = tk.StringVar(
            value=self.settings.get("model", ""))
        self.model_combo = ttk.Combobox(
            api_frame, textvariable=self.model_var, width=30)
        self.model_combo.grid(row=3, column=1, sticky=tk.W, pady=4, padx=4)
        self._update_model_list()

        ttk.Button(api_frame, text="🔌 测试连接",
                   command=self._test_api
                   ).grid(row=4, column=1, sticky=tk.W, pady=8, padx=4)
        self.test_result_var = tk.StringVar()
        tk.Label(api_frame, textvariable=self.test_result_var,
                 fg="#1F4E79", wraplength=600,
                 justify=tk.LEFT
                 ).grid(row=5, column=0, columnspan=3,
                        sticky=tk.W, padx=4)

        ext_frame = ttk.LabelFrame(frame, text="抽取参数", padding=10)
        ext_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(ext_frame, text="最大字符数:").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        self.max_chars_var = tk.IntVar(
            value=self.settings.get("max_chars", 80000))
        ttk.Entry(ext_frame, textvariable=self.max_chars_var,
                  width=10).grid(row=0, column=1, sticky=tk.W, pady=4, padx=4)
        ttk.Label(ext_frame, text="(每篇喂给AI的最大字符)",
                  foreground="#666").grid(row=0, column=2, sticky=tk.W)
        ttk.Label(ext_frame, text="AI 超时(秒):").grid(
            row=1, column=0, sticky=tk.W, pady=4)
        self.ai_timeout_var = tk.IntVar(
            value=self.settings.get("ai_timeout", 600))
        ttk.Entry(ext_frame, textvariable=self.ai_timeout_var,
                  width=10).grid(row=1, column=1, sticky=tk.W,
                                 pady=4, padx=4)
        ttk.Label(ext_frame,
                  text="(R1 类思考模型建议 ≥ 600)",
                  foreground="#666").grid(row=1, column=2, sticky=tk.W)

        ttk.Label(ext_frame, text="AI 重试次数:").grid(
            row=2, column=0, sticky=tk.W, pady=4)
        self.ai_retries_var = tk.IntVar(
            value=self.settings.get("ai_max_retries", 3))
        ttk.Entry(ext_frame, textvariable=self.ai_retries_var,
                  width=10).grid(row=2, column=1, sticky=tk.W,
                                 pady=4, padx=4)
        ttk.Label(ext_frame,
                  text="(限流/超时时自动重试的次数)",
                  foreground="#666").grid(row=2, column=2, sticky=tk.W)

        ttk.Label(ext_frame, text="最大输入 tokens:").grid(
            row=3, column=0, sticky=tk.W, pady=4)
        self.max_input_tokens_var = tk.IntVar(
            value=self.settings.get("max_input_tokens", 100000))
        ttk.Entry(ext_frame, textvariable=self.max_input_tokens_var,
                  width=10).grid(row=3, column=1, sticky=tk.W,
                                 pady=4, padx=4)
        ttk.Label(ext_frame,
                  text="(估算超过会截断,DeepSeek 建议 ≤ 60000)",
                  foreground="#666").grid(row=3, column=2, sticky=tk.W)

        perf_frame = ttk.LabelFrame(
            frame, text="电脑性能与 GPU 加速", padding=10)
        perf_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(perf_frame, text="性能模式:").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        mode_id = self.settings.get("performance_mode", "balanced")
        self.performance_mode_var = tk.StringVar(
            value=MODE_LABELS.get(mode_id, MODE_LABELS["balanced"]))
        mode_combo = ttk.Combobox(
            perf_frame, textvariable=self.performance_mode_var,
            values=list(MODE_LABELS.values()), width=18, state="readonly")
        mode_combo.grid(row=0, column=1, sticky=tk.W, padx=4)
        mode_combo.bind("<<ComboboxSelected>>", self._on_performance_change)

        self.hardware_status_var = tk.StringVar(
            value=self.performance_manager.hardware_summary())
        ttk.Label(perf_frame, textvariable=self.hardware_status_var,
                  foreground="#1F4E79", wraplength=850).grid(
                      row=1, column=0, columnspan=6, sticky=tk.W, pady=3)
        self.performance_summary_var = tk.StringVar(
            value=self.performance_manager.budget_summary())
        ttk.Label(perf_frame, textvariable=self.performance_summary_var,
                  foreground="#356B45", wraplength=850).grid(
                      row=2, column=0, columnspan=6, sticky=tk.W, pady=3)

        custom_specs = [
            ("文档进程", "custom_cpu_workers", 2),
            ("AI 并发", "custom_ai_workers", 3),
            ("同时任务", "custom_active_documents", 3),
            ("内存上限 %", "memory_limit_percent", 75),
            ("OCR DPI", "ocr_dpi", 190),
        ]
        self.performance_custom_entries = []
        self.performance_custom_vars = {}
        for index, (label, key, default) in enumerate(custom_specs):
            column = (index % 3) * 2
            row = 3 + index // 3
            ttk.Label(perf_frame, text=label + ":").grid(
                row=row, column=column, sticky=tk.W, pady=4)
            variable = tk.IntVar(value=self.settings.get(key, default))
            self.performance_custom_vars[key] = variable
            entry = ttk.Entry(perf_frame, textvariable=variable, width=8)
            entry.grid(row=row, column=column + 1, sticky=tk.W, padx=4)
            entry.bind("<FocusOut>", self._on_performance_change)
            self.performance_custom_entries.append(entry)

        ttk.Label(perf_frame, text="GPU 模式:").grid(
            row=5, column=0, sticky=tk.W, pady=4)
        gpu_labels = {"auto": "自动（推荐）", "off": "关闭",
                      "force": "强制使用"}
        self.gpu_mode_var = tk.StringVar(value=gpu_labels.get(
            self.settings.get("gpu_mode", "auto"), "自动（推荐）"))
        gpu_combo = ttk.Combobox(
            perf_frame, textvariable=self.gpu_mode_var,
            values=list(gpu_labels.values()), width=18, state="readonly")
        gpu_combo.grid(row=5, column=1, sticky=tk.W, padx=4)
        gpu_combo.bind("<<ComboboxSelected>>", self._on_performance_change)
        self.gpu_ocr_var = tk.BooleanVar(
            value=self.settings.get("gpu_ocr_enabled", True))
        ttk.Checkbutton(
            perf_frame, text="扫描 PDF 优先使用 GPU OCR",
            variable=self.gpu_ocr_var,
            command=self._on_performance_change).grid(
                row=5, column=2, columnspan=2, sticky=tk.W, padx=6)
        ttk.Button(perf_frame, text="重新检测并自检 GPU",
                   command=self._run_gpu_self_test).grid(
                       row=5, column=4, columnspan=2, sticky=tk.W, padx=4)
        self.gpu_test_result_var = tk.StringVar(value="")
        ttk.Label(perf_frame, textvariable=self.gpu_test_result_var,
                  foreground="#1F4E79", wraplength=850).grid(
                      row=6, column=0, columnspan=6, sticky=tk.W, pady=3)
        self._on_performance_change()

        path_frame = ttk.LabelFrame(frame, text="项目路径", padding=10)
        path_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(path_frame, text="数据目录:").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        self.proj_dir_var = tk.StringVar(
            value=self.settings.get("project_dir", ""))
        ttk.Entry(path_frame, textvariable=self.proj_dir_var,
                  width=60).grid(row=0, column=1, sticky=tk.W, pady=4, padx=4)
        ttk.Button(path_frame, text="📁",
                   command=self._select_project_dir
                   ).grid(row=0, column=2, padx=4)

        shortcut_frame = ttk.Frame(path_frame)
        shortcut_frame.grid(row=1, column=0, columnspan=3,
                            sticky=tk.W, pady=6)
        ttk.Button(shortcut_frame, text="📂 打开数据目录",
                   command=self._open_project_dir
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(shortcut_frame, text="📂 打开提示词目录",
                   command=self._open_prompts_dir
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(shortcut_frame, text="📂 打开配置文件目录",
                   command=self._open_config_dir
                   ).pack(side=tk.LEFT, padx=4)

        update_frame = ttk.LabelFrame(frame, text="软件更新", padding=10)
        update_frame.pack(fill=tk.X, padx=8, pady=8)
        ttk.Label(
            update_frame, text="当前版本：v" + APP_VERSION,
            font=("微软雅黑", 9, "bold")
        ).grid(row=0, column=0, sticky=tk.W, pady=4)
        self.auto_update_var = tk.BooleanVar(
            value=self.settings.get("auto_update_check", True))
        ttk.Checkbutton(
            update_frame, text="启动后自动检查 GitHub Release 更新",
            variable=self.auto_update_var
        ).grid(row=0, column=1, sticky=tk.W, padx=12)
        ttk.Button(
            update_frame, text="检查更新",
            command=lambda: self._check_for_updates(manual=True)
        ).grid(row=0, column=2, sticky=tk.W, padx=4)
        ttk.Button(
            update_frame, text="打开发布页",
            command=lambda: webbrowser.open(GITHUB_RELEASES_URL)
        ).grid(row=0, column=3, sticky=tk.W, padx=4)
        self.update_status_var = tk.StringVar(
            value="便携版更新只替换程序文件，不修改数据库和设置。")
        update_status = tk.Label(
            update_frame, textvariable=self.update_status_var,
            anchor=tk.W, justify=tk.LEFT, wraplength=900)
        update_status._theme_role = "muted"
        update_status.grid(
            row=1, column=0, columnspan=4, sticky=tk.W, pady=(4, 0))

        danger_frame = ttk.LabelFrame(frame, text="危险操作", padding=10)
        danger_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(danger_frame, text="以下操作不可撤销",
                  foreground="red").pack(anchor=tk.W)
        btn_danger = ttk.Frame(danger_frame)
        btn_danger.pack(fill=tk.X, pady=4)
        ttk.Button(btn_danger, text="🗑 清空全部缓存",
                   command=self._clear_all_cache,
                   style="Danger.TButton"
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_danger, text="🗑 清空当前数据库",
                   command=self._clear_database,
                   style="Danger.TButton"
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_danger, text="🗑 完全重置",
                   command=self._full_reset,
                   style="Danger.TButton"
                   ).pack(side=tk.LEFT, padx=4)

        ttk.Button(frame, text="💾 保存设置",
                   command=self._save_settings,
                   style="Accent.TButton").pack(pady=10)

        # 数据备份区块
        backup_frame = ttk.LabelFrame(frame, text="数据备份", padding=10)
        backup_frame.pack(fill=tk.X, padx=8, pady=8)

        ttk.Label(backup_frame,
                  text="配置文件(API Key、模型设置等):").grid(
            row=0, column=0, sticky=tk.W, pady=4)
        btn_cfg = ttk.Frame(backup_frame)
        btn_cfg.grid(row=0, column=1, sticky=tk.W, padx=4)
        ttk.Button(btn_cfg, text="📤 导出配置",
                   command=self._export_config
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_cfg, text="📥 导入配置",
                   command=self._import_config
                   ).pack(side=tk.LEFT, padx=4)

        ttk.Label(backup_frame, text="数据库(已收录文献):").grid(
            row=1, column=0, sticky=tk.W, pady=4)
        btn_db = ttk.Frame(backup_frame)
        btn_db.grid(row=1, column=1, sticky=tk.W, padx=4)
        ttk.Button(btn_db, text="📤 导出数据库",
                   command=self._export_database
                   ).pack(side=tk.LEFT, padx=4)
        ttk.Button(btn_db, text="📥 导入数据库",
                   command=self._import_database
                   ).pack(side=tk.LEFT, padx=4)

    def _appearance_values_from_ui(self):
        try:
            font_scale = int(self.font_scale_var.get().rstrip("%"))
        except (ValueError, tk.TclError):
            font_scale = 100
        return normalize_appearance({
            "theme": THEME_IDS_BY_LABEL.get(
                self.theme_var.get(), "classic"),
            "font_scale": font_scale,
            "ui_density": DENSITY_IDS_BY_LABEL.get(
                self.ui_density_var.get(), "comfortable"),
        })

    def _preview_theme(self, event=None):
        preview = dict(self.settings)
        preview.update(self._appearance_values_from_ui())
        self.theme_manager.apply(preview)
        self._update_appearance_status()

    def _reset_theme_preview(self):
        self.theme_var.set(THEME_LABELS["classic"])
        self.font_scale_var.set("100%")
        self.ui_density_var.set(DENSITY_LABELS["comfortable"])
        self._preview_theme()

    def _update_appearance_status(self):
        if not hasattr(self, "appearance_status_var"):
            return
        selected = THEME_IDS_BY_LABEL.get(
            self.theme_var.get(), "classic")
        if selected == "system":
            resolved_label = THEME_LABELS.get(
                self.theme_manager.resolved_theme, "浅色")
            text = "跟随 Windows；当前实际使用：" + resolved_label
        else:
            text = "当前预览：" + THEME_LABELS.get(selected, "经典")
        if selected == "speed":
            text += " · 使用轻量原生控件，减少界面绘制开销"
        text += (" · 字体 " + self.font_scale_var.get() + " · " +
                 self.ui_density_var.get())
        self.appearance_status_var.set(text)

    # ── GitHub Release 自动更新 ────────────────────────

    def _auto_check_for_updates(self):
        """Quiet daily startup check for packaged portable builds."""
        if not getattr(sys, "frozen", False):
            return
        if not self.settings.get("auto_update_check", True):
            return
        try:
            last_check = float(self.settings.get("last_update_check", 0))
        except (TypeError, ValueError):
            last_check = 0
        if time.time() - last_check < 24 * 60 * 60:
            return
        self._check_for_updates(manual=False)

    def _check_for_updates(self, manual=False):
        if getattr(self, "_update_checking", False):
            if manual:
                messagebox.showinfo("软件更新", "正在检查，请稍候。")
            return
        self._update_checking = True
        if hasattr(self, "update_status_var"):
            self.update_status_var.set("正在连接 GitHub 检查更新……")

        def worker():
            try:
                info = check_latest_release(APP_VERSION)
                error = None
            except Exception as exc:
                info = None
                error = exc
            self.after(
                0, lambda: self._finish_update_check(info, error, manual))

        threading.Thread(target=worker, daemon=True).start()

    def _finish_update_check(self, info, error, manual):
        self._update_checking = False
        self.settings["last_update_check"] = int(time.time())
        save_settings(self.settings)
        if error is not None:
            text = str(error)
            if hasattr(self, "update_status_var"):
                self.update_status_var.set("检查失败：" + text)
            if manual:
                messagebox.showerror("检查更新失败", text)
            return
        if info is None:
            text = "当前已是最新版本，或仓库尚未发布正式 Release。"
            if hasattr(self, "update_status_var"):
                self.update_status_var.set(text)
            if manual:
                messagebox.showinfo("软件更新", text)
            return

        if hasattr(self, "update_status_var"):
            self.update_status_var.set(
                "发现新版本 v" + info.version + "：" + info.title)
        notes = info.notes.strip()
        if len(notes) > 1200:
            notes = notes[:1200] + "……"
        prompt = (
            "发现 Lazybones v" + info.version + chr(10) + chr(10) +
            (notes + chr(10) + chr(10) if notes else "") +
            "是否下载并安装？" + chr(10) +
            "更新前会备份现有程序文件，数据库和设置不会被修改。"
        )
        if not getattr(sys, "frozen", False):
            if messagebox.askyesno(
                    "发现新版本", prompt + chr(10) + chr(10) +
                    "当前是源码运行模式，只能打开发布页手动下载。"):
                webbrowser.open(info.release_url)
            return
        if messagebox.askyesno("发现新版本", prompt):
            self._download_and_install_update(info)

    def _download_and_install_update(self, info):
        if hasattr(self, "update_status_var"):
            self.update_status_var.set(
                "正在下载 " + info.asset_name + "……")

        def progress(downloaded, total):
            if total > 0:
                percent = min(100, round(downloaded * 100 / total))
                text = "正在下载更新：" + str(percent) + "%"
            else:
                text = "正在下载更新：" + str(downloaded // 1024 // 1024) + " MB"
            self.after(0, lambda value=text: self.update_status_var.set(value))

        def worker():
            try:
                package = download_update(info, updates_dir(), progress)
                error = None
            except Exception as exc:
                package = None
                error = exc
            self.after(
                0, lambda: self._finish_update_download(
                    info, package, error))

        threading.Thread(target=worker, daemon=True).start()

    def _finish_update_download(self, info, package, error):
        if error is not None:
            text = str(error)
            self.update_status_var.set("下载失败：" + text)
            messagebox.showerror("更新失败", text)
            return
        try:
            launch_portable_updater(package, info)
        except (OSError, UpdateError) as exc:
            self.update_status_var.set("无法启动更新器：" + str(exc))
            messagebox.showerror("更新失败", str(exc))
            return
        self.update_status_var.set("更新器已启动，程序即将关闭并自动重启。")
        self.after(250, self.destroy)

    def _performance_values_from_ui(self):
        gpu_ids = {"自动（推荐）": "auto", "关闭": "off",
                   "强制使用": "force"}
        values = {
            "performance_mode": MODE_IDS_BY_LABEL.get(
                self.performance_mode_var.get(), "balanced"),
            "gpu_mode": gpu_ids.get(self.gpu_mode_var.get(), "auto"),
            "gpu_ocr_enabled": bool(self.gpu_ocr_var.get()),
        }
        defaults = {
            "custom_cpu_workers": 2,
            "custom_ai_workers": 3,
            "custom_active_documents": 3,
            "memory_limit_percent": 75,
            "ocr_dpi": 190,
        }
        for key, variable in self.performance_custom_vars.items():
            try:
                values[key] = variable.get()
            except tk.TclError:
                values[key] = defaults[key]
        return values

    def _on_performance_change(self, event=None):
        if not hasattr(self, "performance_custom_vars"):
            return
        values = self._performance_values_from_ui()
        preview = dict(self.settings)
        preview.update(values)
        budget = self.performance_manager.refresh(preview)
        custom = values["performance_mode"] == "custom"
        state = tk.NORMAL if custom else tk.DISABLED
        for entry in self.performance_custom_entries:
            entry.configure(state=state)
        self.performance_summary_var.set(
            self.performance_manager.budget_summary())

    def _run_gpu_self_test(self):
        self.gpu_test_result_var.set("正在重新检测 GPU 并运行 OCR 自检…")
        values = self._performance_values_from_ui()
        preview = dict(self.settings)
        preview.update(values)
        self.performance_manager.refresh(preview, redetect=True)
        self.hardware_status_var.set(
            self.performance_manager.hardware_summary())
        prefer_gpu = (
            self.performance_manager.resources.gpu.directml_available)

        def worker():
            from gpu_acceleration import clear_engine_cache, self_test
            clear_engine_cache()
            result = self_test(prefer_gpu=prefer_gpu)

            def finish():
                prefix = "✓ " if result["ok"] else "✗ "
                detail = result["message"]
                if result.get("recognized"):
                    detail += "；识别结果：" + result["recognized"][:100]
                self.gpu_test_result_var.set(prefix + detail)
                self._on_performance_change()

            self.after(0, finish)

        threading.Thread(target=worker, daemon=True).start()

    # ── API / 模型 ────────────────────────────────

    def _on_provider_change(self, event):
        provider = self.provider_var.get()
        if provider in PROVIDERS:
            cfg = PROVIDERS[provider]
            if cfg["base_url"]:
                self.base_url_var.set(cfg["base_url"])
            self._update_model_list()

    def _update_model_list(self):
        provider = self.provider_var.get()
        if provider in PROVIDERS:
            self.model_combo["values"] = PROVIDERS[provider]["models"]

    def _toggle_key_visibility(self):
        self.api_key_entry.configure(
            show="" if self.show_key_var.get() else "*")

    def _test_api(self):
        self.test_result_var.set("测试中...")
        self.update_idletasks()

        def _do_test():
            success, msg = test_connection(
                self.base_url_var.get(),
                self.api_key_var.get(),
                self.model_var.get())
            prefix = "✓ " if success else "✗ "
            self.after(0, lambda: self.test_result_var.set(prefix + msg))

        threading.Thread(target=_do_test, daemon=True).start()

    # ── 路径相关 ──────────────────────────────────

    def _select_project_dir(self):
        d = filedialog.askdirectory(initialdir=self.proj_dir_var.get())
        if d:
            self.proj_dir_var.set(d)

    def _open_project_dir(self):
        p = Path(self.proj_dir_var.get())
        if p.exists():
            os.startfile(str(p))
        else:
            messagebox.showwarning("目录不存在",
                                   "数据目录不存在:" + str(p))

    def _open_prompts_dir(self):
        from schema_loader import get_prompts_dir
        p = get_prompts_dir()
        if p.exists():
            os.startfile(str(p))
        else:
            messagebox.showwarning("目录不存在",
                                   "提示词目录不存在:" + str(p))

    def _open_config_dir(self):
        from schema_loader import get_config_path
        p = get_config_path().parent
        if p.exists():
            os.startfile(str(p))
        else:
            messagebox.showwarning("目录不存在",
                                   "配置目录不存在:" + str(p))

    # ── 配置 / 数据库 备份与导入 ───────────────────

    def _export_config(self):
        from schema_loader import get_config_path
        src = get_config_path()
        if not src.exists():
            messagebox.showwarning("导出失败",
                                   "配置文件不存在,请先保存设置。")
            return
        dst = filedialog.asksaveasfilename(
            title="导出配置文件",
            defaultextension=".json",
            filetypes=[("JSON文件", "*.json")],
            initialfile="Lazybones_config.json")
        if dst:
            shutil.copy2(str(src), dst)
            messagebox.showinfo("导出成功", "配置已导出到:" + dst)

    def _import_config(self):
        src = filedialog.askopenfilename(
            title="选择配置文件",
            filetypes=[("JSON文件", "*.json")])
        if not src:
            return
        if not messagebox.askyesno(
                "确认导入",
                "导入后将覆盖当前所有设置(包括API Key),确认继续?"):
            return
        try:
            from schema_loader import get_config_path
            with open(src, encoding="utf-8") as f:
                json.load(f)
            dst = get_config_path()
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, str(dst))
            self.settings = load_settings()
            self._apply_settings_to_ui()
            messagebox.showinfo("导入成功", "配置已导入,设置页已更新。")
        except Exception as e:
            messagebox.showerror("导入失败",
                                 "文件格式错误或读取失败:" + str(e))

    def _export_database(self):
        src = Path(self.db_path)
        if not src.exists():
            messagebox.showwarning("导出失败", "数据库文件不存在。")
            return
        dst = filedialog.asksaveasfilename(
            title="导出数据库",
            defaultextension=".db",
            filetypes=[("SQLite数据库", "*.db")],
            initialfile=src.name)
        if dst:
            shutil.copy2(str(src), dst)
            messagebox.showinfo("导出成功", "数据库已导出到:" + dst)

    def _import_database(self):
        """
        ✅ 修复:替换前 gc.collect() 让可能未关闭的 SQLite 句柄释放,
        避免 Windows 上 shutil.copy2 因句柄占用失败。
        """
        src = filedialog.askopenfilename(
            title="选择数据库文件",
            filetypes=[("SQLite数据库", "*.db")])
        if not src:
            return
        if not messagebox.askyesno(
                "确认导入",
                "导入后将替换当前数据库,原有数据将被覆盖。"
                "建议先导出备份,确认继续?"):
            return
        try:
            conn = sqlite3.connect(src)
            conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
            conn.close()

            dst = Path(self.db_path)
            dst.parent.mkdir(parents=True, exist_ok=True)

            gc.collect()

            shutil.copy2(src, str(dst))
            self._init_db()
            self._refresh_manage_tab()
            self._refresh_review_tab()
            messagebox.showinfo("导入成功", "数据库已导入并重新加载。")
        except Exception as e:
            messagebox.showerror("导入失败",
                                 "文件无效或替换失败:" + str(e))

    # ── 设置保存 / UI 同步 ────────────────────────

    def _save_settings(self):
        appearance_values = self._appearance_values_from_ui()
        performance_values = self._performance_values_from_ui()
        self.settings.update({
            "provider": self.provider_var.get(),
            "api_key": self.api_key_var.get(),
            "base_url": self.base_url_var.get(),
            "model": self.model_var.get(),
            "max_chars": self.max_chars_var.get(),
            "project_dir": self.proj_dir_var.get(),
            "lang_zh": (self.lang_zh_var.get()
                        if hasattr(self, "lang_zh_var") else True),
            "lang_en": (self.lang_en_var.get()
                        if hasattr(self, "lang_en_var") else True),
            "schema_version": (self.schema_var.get()
                               if hasattr(self, "schema_var") else ""),
            "auto_confirm_clean_pairs": (
                self.auto_confirm_pairs_var.get()
                if hasattr(self, "auto_confirm_pairs_var") else True),
            "auto_open_review": (
                self.auto_open_review_var.get()
                if hasattr(self, "auto_open_review_var") else True),
            "ai_timeout": (self.ai_timeout_var.get()
                           if hasattr(self, "ai_timeout_var") else 600),
            "ai_max_retries": (self.ai_retries_var.get()
                               if hasattr(self, "ai_retries_var") else 3),
            "max_input_tokens": (self.max_input_tokens_var.get()
                                 if hasattr(self, "max_input_tokens_var")
                                 else 100000),
            "auto_update_check": (
                self.auto_update_var.get()
                if hasattr(self, "auto_update_var") else True),
        })
        self.settings.update(appearance_values)
        self.settings.update(performance_values)
        budget = self.performance_manager.refresh(self.settings)
        # Preserve the old key for older releases/config exports.
        self.settings["concurrent"] = budget.ai_workers
        save_settings(self.settings)
        self.theme_manager.apply(self.settings)
        self._ensure_project_dirs()
        self._init_db()
        messagebox.showinfo("已保存", "设置已保存")

    def _apply_settings_to_ui(self):
        """
        把 self.settings 里的值同步到设置页所有控件上。
        被 _import_config 调用,导入配置后刷新 UI。
        """
        self.provider_var.set(self.settings.get("provider", "DeepSeek"))
        self.api_key_var.set(self.settings.get("api_key", ""))
        self.base_url_var.set(self.settings.get("base_url", ""))
        self.model_var.set(self.settings.get("model", ""))
        self.max_chars_var.set(self.settings.get("max_chars", 80000))
        self.proj_dir_var.set(self.settings.get("project_dir", ""))
        if hasattr(self, "theme_var"):
            appearance = normalize_appearance(self.settings)
            self.theme_var.set(THEME_LABELS[appearance["theme"]])
            self.font_scale_var.set(str(appearance["font_scale"]) + "%")
            self.ui_density_var.set(
                DENSITY_LABELS[appearance["ui_density"]])
            self.theme_manager.apply(self.settings)
            self._update_appearance_status()
        if hasattr(self, "lang_zh_var"):
            self.lang_zh_var.set(self.settings.get("lang_zh", True))
        if hasattr(self, "lang_en_var"):
            self.lang_en_var.set(self.settings.get("lang_en", True))
        if hasattr(self, "auto_confirm_pairs_var"):
            self.auto_confirm_pairs_var.set(
                self.settings.get("auto_confirm_clean_pairs", True))
        if hasattr(self, "auto_open_review_var"):
            self.auto_open_review_var.set(
                self.settings.get("auto_open_review", True))
        # 同步新增的 3 个参数(用 hasattr 守护,兼容旧版 settings.json)
        if hasattr(self, "ai_timeout_var"):
            self.ai_timeout_var.set(self.settings.get("ai_timeout", 600))
        if hasattr(self, "ai_retries_var"):
            self.ai_retries_var.set(
                self.settings.get("ai_max_retries", 3))
        if hasattr(self, "max_input_tokens_var"):
            self.max_input_tokens_var.set(
                self.settings.get("max_input_tokens", 100000))
        if hasattr(self, "auto_update_var"):
            self.auto_update_var.set(
                self.settings.get("auto_update_check", True))
        if hasattr(self, "performance_mode_var"):
            mode = self.settings.get("performance_mode", "balanced")
            self.performance_mode_var.set(
                MODE_LABELS.get(mode, MODE_LABELS["balanced"]))
            gpu_labels = {"auto": "自动（推荐）", "off": "关闭",
                          "force": "强制使用"}
            self.gpu_mode_var.set(gpu_labels.get(
                self.settings.get("gpu_mode", "auto"), "自动（推荐）"))
            self.gpu_ocr_var.set(
                self.settings.get("gpu_ocr_enabled", True))
            defaults = {
                "custom_cpu_workers": 2,
                "custom_ai_workers": 3,
                "custom_active_documents": 3,
                "memory_limit_percent": 75,
                "ocr_dpi": 190,
            }
            for key, variable in self.performance_custom_vars.items():
                variable.set(self.settings.get(key, defaults[key]))
            self.performance_manager.refresh(self.settings)
            self.hardware_status_var.set(
                self.performance_manager.hardware_summary())
            self._on_performance_change()
        self._update_model_list()

    # ── 危险操作 ───────────────────────────────────

    def _clear_all_cache(self):
        if not messagebox.askyesno(
                "确认",
                "删除所有 text_cache 和 extract_cache?"
                "原始PDF和数据库不受影响"):
            return
        project_dir = self.settings.get("project_dir", "")
        for folder in ["text_cache", "extract_cache"]:
            p = Path(project_dir) / folder
            if p.exists():
                shutil.rmtree(p)
                p.mkdir()
        messagebox.showinfo("完成", "缓存已清空")
        self._refresh_review_tab()

    def _clear_database(self):
        """
        ✅ 修复:DELETE 之后执行 VACUUM,让数据库文件真正瘦身。
        """
        if not messagebox.askyesno(
                "确认",
                "清空当前数据库「" + self.db_name_var.get() + "」?"):
            return
        conn = sqlite3.connect(self.db_path)
        try:
            conn.execute("DELETE FROM papers")
            try:
                conn.execute("DELETE FROM rejected")
            except Exception:
                pass
            conn.commit()
            conn.execute("VACUUM")
        finally:
            conn.close()
        messagebox.showinfo("完成", "数据库已清空")
        self._refresh_manage_tab()

    def _full_reset(self):
        """✅ 修复:清空数据库后执行 VACUUM"""
        confirm = simpledialog.askstring(
            "完全重置",
            "删除所有缓存和数据库(PDF保留)输入 RESET 确认:")
        if confirm != "RESET":
            messagebox.showinfo("已取消", "未执行任何操作")
            return
        project_dir = self.settings.get("project_dir", "")
        for folder in ["text_cache", "extract_cache"]:
            p = Path(project_dir) / folder
            if p.exists():
                shutil.rmtree(p)
                p.mkdir()
        conn = sqlite3.connect(self.db_path)
        try:
            conn.execute("DELETE FROM papers")
            try:
                conn.execute("DELETE FROM rejected")
            except Exception:
                pass
            conn.commit()
            conn.execute("VACUUM")
        finally:
            conn.close()
        self._chat_messages = []
        self._find_replace_history = []   # 查找替换历史记录
        messagebox.showinfo("完成", "已完全重置")
        self._refresh_manage_tab()
        self._refresh_review_tab()

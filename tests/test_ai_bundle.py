import hashlib
import io
import json
import multiprocessing as mp
from pathlib import Path
import queue
import shutil
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from ai_bundle import BundleOptions, bundle_worker, compact_text, export_bundle, split_text


def make_pdf(path, count=2):
    import fitz
    with fitz.open() as doc:
        for index in range(count):
            page = doc.new_page()
            page.insert_text((50, 60), f"PAGE_MARKER_{index + 1} Research results: concentration 0.025 mol/L.")
            page.insert_text((50, 100), "References and numerical measurements must remain in the exported text.")
        doc.save(path)


def make_docx(path):
    from docx import Document
    from docx.oxml import OxmlElement
    from PIL import Image
    doc = Document()
    doc.add_paragraph("START_MARKER — 作者全名；温度 298 K；浓度 0.05 mol/L。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Sample", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "A|B", "0.025 ± 0.001"
    doc.add_paragraph("AFTER_TABLE_MARKER")
    paragraph = doc.add_paragraph("Equation: ")
    math = OxmlElement("m:oMath")
    fraction = OxmlElement("m:f")
    for name, value in (("num", "x"), ("den", "y")):
        part = OxmlElement("m:" + name)
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = value
        run.append(text)
        part.append(run)
        fraction.append(part)
    math.append(fraction)
    paragraph._p.append(math)
    image = io.BytesIO()
    Image.new("RGB", (20, 20), "blue").save(image, format="PNG")
    image_bytes = image.getvalue()
    image.seek(0)
    doc.add_picture(image)
    doc.sections[0].header.paragraphs[0].text = "HEADER_MARKER"
    doc.add_paragraph("END_MARKER References: example DOI 10.0000/example")
    doc.save(path)
    return image_bytes


class BundleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="lazybones-bundle-test-")
        self.root = Path(self.temp.name)
        self.output = self.root / "output"

    def tearDown(self):
        self.temp.cleanup()

    def test_lossless_unicode_split_obeys_conservative_budget(self):
        text = ("中文，±0.025 mol/L；α β 🙂\n" + "A" * 350 + "\n") * 400
        chunks = list(split_text(text, 2000))
        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(chunks), text)
        self.assertTrue(all(len(chunk.encode("utf-8")) <= 2000 for chunk in chunks))

    def test_compaction_keeps_scientific_characters_and_hyphens(self):
        self.assertEqual(compact_text("  x² ± 0.01  \r\nFe-\r\nO\n\n\n\nEND"), "x² ± 0.01\nFe-\nO\n\n\nEND")

    def test_pdf_all_pages_sources_and_part_payloads_preserved(self):
        source = self.root / "article.pdf"
        make_pdf(source, 8)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        result = export_bundle([source], self.output, BundleOptions(mode="text", part_tokens=2000, ocr=False))
        run = Path(result["directory"])
        record = result["documents"][0]
        markdown = (run / record["markdown"]).read_text(encoding="utf-8")
        self.assertEqual(record["pages"], 8)
        self.assertTrue(all(f"PAGE_MARKER_{page}" in markdown for page in range(1, 9)))
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)
        manifest = json.loads((run / "manifest.json").read_text(encoding="utf-8"))
        payloads = []
        for entry in manifest["parts"]:
            content = (run / entry["file"]).read_text(encoding="utf-8")
            self.assertLessEqual(len(content.encode("utf-8")), 1500)
            payloads.append(content.split("\n\n", 2)[2])
        self.assertEqual("".join(payloads), markdown)
        self.assertNotIn(str(source.parent), (run / "manifest.json").read_text(encoding="utf-8"))
        with zipfile.ZipFile(result["zip"]) as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("全文合集.md", archive.namelist())
            self.assertFalse(any(name.endswith(".db") for name in archive.namelist()))

    def test_word_tables_equations_images_and_order_are_exported(self):
        source = self.root / "paper.docx"
        image_bytes = make_docx(source)
        result = export_bundle([source], self.output, BundleOptions(ocr=False))
        run = Path(result["directory"])
        record = result["documents"][0]
        self.assertEqual(record["status"], "completed", record)
        text = (run / record["markdown"]).read_text(encoding="utf-8")
        self.assertLess(text.index("START_MARKER"), text.index("Sample"))
        self.assertLess(text.index("Sample"), text.index("AFTER_TABLE_MARKER"))
        for marker in ("END_MARKER", "HEADER_MARKER", "0.025 ± 0.001", "\\frac{x}{y}", "A\\|B"):
            self.assertIn(marker, text)
        self.assertEqual(record["tables"], 1)
        self.assertEqual(record["formulas"], 1)
        self.assertEqual(record["images"], 1)
        assets = run / "documents/D0001/assets"
        self.assertEqual((assets / "image-0001.png").read_bytes(), image_bytes)
        self.assertIn("oMath", (assets / "formula-0001.xml").read_text(encoding="utf-8"))
        self.assertIn("](documents/D0001/assets/", (run / "全文合集.md").read_text(encoding="utf-8"))
        part = next((run / "parts").glob("*.md")).read_text(encoding="utf-8")
        self.assertIn("](../documents/D0001/assets/", part)

    def test_exact_duplicates_alias_without_repeating_content(self):
        first = self.root / "one.pdf"
        second = self.root / "renamed.pdf"
        make_pdf(first)
        shutil.copyfile(first, second)
        result = export_bundle([first, second], self.output, BundleOptions(mode="text", ocr=False))
        self.assertEqual(result["summary"]["completed"], 1)
        self.assertEqual(result["summary"]["duplicates"], 1)
        self.assertEqual(result["documents"][1]["duplicate_of"], "D0001")

    def test_failed_file_does_not_hide_successful_file(self):
        broken = self.root / "broken.pdf"
        broken.write_bytes(b"not a PDF")
        good = self.root / "good.pdf"
        make_pdf(good)
        result = export_bundle([broken, good], self.output, BundleOptions(mode="text", ocr=False))
        self.assertEqual(result["summary"]["failed"], 1)
        self.assertEqual(result["summary"]["completed"], 1)
        self.assertIn("broken.pdf", (Path(result["directory"]) / "先读我_质量报告.md").read_text(encoding="utf-8"))

    def test_visual_mode_keeps_pdf_page_even_without_embedded_images(self):
        source = self.root / "vector.pdf"
        make_pdf(source, 1)
        result = export_bundle([source], self.output, BundleOptions(ocr=False, image_dpi=96))
        self.assertEqual(result["documents"][0]["images"], 1)
        png = Path(result["directory"]) / "documents/D0001/assets/page-0001.png"
        self.assertTrue(png.read_bytes().startswith(b"\x89PNG"))

    def test_cancel_marks_remaining_files_and_does_not_publish_zip(self):
        first = self.root / "one.pdf"
        second = self.root / "two.pdf"
        make_pdf(first)
        make_pdf(second)
        cancel = threading.Event()
        def progress(message):
            if "第 1/" in message:
                cancel.set()
        result = export_bundle([first, second], self.output, BundleOptions(mode="text", ocr=False),
                               cancel=cancel, progress=progress)
        self.assertTrue(result["cancelled"])
        self.assertIsNone(result["zip"])
        self.assertEqual([record["status"] for record in result["documents"]], ["cancelled", "not_processed"])

    def test_ocr_failure_is_reported_not_silently_successful_text(self):
        import fitz
        source = self.root / "scan.pdf"
        with fitz.open() as document:
            document.new_page()
            document.save(source)
        with patch("gpu_acceleration.run_ocr", side_effect=RuntimeError("test failure")):
            result = export_bundle([source], self.output, BundleOptions(mode="text"))
        self.assertTrue(any("OCR 失败" in warning for warning in result["documents"][0]["warnings"]))
        self.assertEqual(result["documents"][0]["status"], "needs_review")
        self.assertEqual(result["summary"]["needs_review"], 1)

    def test_new_run_never_overwrites_previous_result(self):
        source = self.root / "one.pdf"
        make_pdf(source, 1)
        first = export_bundle([source], self.output, BundleOptions(mode="text", ocr=False))
        second = export_bundle([source], self.output, BundleOptions(mode="text", ocr=False))
        self.assertNotEqual(first["directory"], second["directory"])
        self.assertTrue(Path(first["zip"]).exists())

    def test_spawned_worker_finishes_with_serializable_result(self):
        source = self.root / "worker.pdf"
        make_pdf(source, 1)
        context = mp.get_context("spawn")
        messages = context.Queue()
        process = context.Process(target=bundle_worker,
                                  args=([str(source)], str(self.output), {"mode": "text", "ocr": False},
                                        {}, messages, context.Event()))
        process.start()
        terminal = None
        try:
            for _ in range(20):
                kind, value = messages.get(timeout=20)
                if kind in {"done", "error"}:
                    terminal = kind, value
                    break
            self.assertIsNotNone(terminal)
            self.assertEqual(terminal[0], "done", terminal)
            self.assertEqual(terminal[1]["summary"]["completed"], 1)
        finally:
            process.join(10)
            if process.is_alive():
                process.terminate()
                process.join()
            messages.close()

    def test_long_document_is_not_truncated_at_extraction_limit(self):
        from docx import Document
        source = self.root / "long.docx"
        document = Document()
        text = "Measurement 0.025 mol/L; " * 5000 + "FINAL_SENTINEL"
        document.add_paragraph(text)
        document.save(source)
        result = export_bundle([source], self.output, BundleOptions(mode="text", part_tokens=8000, ocr=False))
        record = result["documents"][0]
        exported = (Path(result["directory"]) / record["markdown"]).read_text(encoding="utf-8")
        self.assertIn(text, exported)
        self.assertGreater(record["text_bytes"], 80000)
        self.assertGreater(record["parts"], 1)

    def test_table_grid_in_pdf_keeps_cell_values(self):
        import fitz
        source = self.root / "table.pdf"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((50, 40), "Research measurements and sample values are listed in the table below.")
            for x in (50, 200, 350):
                page.draw_line((x, 70), (x, 170))
            for y in (70, 120, 170):
                page.draw_line((50, y), (350, y))
            for x, y, text in ((60, 95, "Sample"), (210, 95, "Value"),
                                (60, 145, "Catalyst A"), (210, 145, "0.025")):
                page.insert_text((x, y), text)
            document.save(source)
        result = export_bundle([source], self.output, BundleOptions(mode="text", ocr=False))
        record = result["documents"][0]
        self.assertEqual(record["tables"], 1)
        data = json.loads((Path(result["directory"]) / "documents/D0001/tables.json").read_text(encoding="utf-8"))
        self.assertEqual(data[0]["rows"][1], ["Catalyst A", "0.025"])

    def test_cancel_during_zip_leaves_no_finished_or_partial_archive(self):
        source = self.root / "one.pdf"
        make_pdf(source, 1)
        cancel = threading.Event()
        def progress(message):
            if "打包 ZIP" in message:
                cancel.set()
        result = export_bundle([source], self.output, BundleOptions(mode="text", ocr=False), cancel=cancel, progress=progress)
        self.assertTrue(result["cancelled"])
        self.assertIsNone(result["zip"])
        self.assertFalse(list(self.output.glob("*.partial")))
        self.assertTrue((Path(result["directory"]) / "全文合集.md").exists())


if __name__ == "__main__":
    unittest.main()

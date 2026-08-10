import shutil
import sys
import tempfile
import unittest
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from extractor import (build_document_identity, deduplicate_pairs,
                       record_document_identity)


class DocumentDeduplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.project_dir = self.root / "project"
        self.project_dir.mkdir()

    def tearDown(self):
        self.temp_dir.cleanup()

    def _make_docx(self, name, title, doi="", author="Wang Lei",
                   year="2024", marker=""):
        path = self.root / name
        doc = Document()
        doc.core_properties.title = title
        doc.core_properties.author = author
        doc.add_paragraph(title)
        if doi:
            doc.add_paragraph("DOI: " + doi)
        doc.add_paragraph("Published in " + year)
        for i in range(50):
            doc.add_paragraph(
                f"Abstract section {i}. {marker} "
                "This study reports reproducible experimental details, "
                "materials characterization, results and discussion.")
        doc.save(path)
        return str(path)

    def test_same_file_renamed_is_removed_within_batch(self):
        first = self._make_docx(
            "paper.docx", "A Stable Paper Title", marker="alpha")
        second = str(self.root / "renamed.docx")
        shutil.copyfile(first, second)

        unique, duplicates, _ = deduplicate_pairs(
            {first: [], second: []}, str(self.project_dir))

        self.assertEqual(1, len(unique))
        self.assertEqual(1, len(duplicates))
        self.assertIn("SHA-256", duplicates[0]["reason"])

    def test_same_doi_removes_different_download_variant(self):
        first = self._make_docx(
            "publisher.docx", "The Same Research Article",
            doi="10.1234/example.2024.1", marker="publisher copy")
        second = self._make_docx(
            "repository.docx", "The Same Research Article",
            doi="10.1234/example.2024.1", marker="repository copy")

        unique, duplicates, _ = deduplicate_pairs(
            {first: [], second: []}, str(self.project_dir))

        self.assertEqual(1, len(unique))
        self.assertEqual(1, len(duplicates))
        self.assertIn("DOI", duplicates[0]["reason"])

    def test_same_author_and_year_does_not_merge_different_titles(self):
        first = self._make_docx(
            "first.docx", "Catalysis Study Number One",
            author="Li Ming", year="2025", marker="catalysis")
        second = self._make_docx(
            "second.docx", "Battery Study Number Two",
            author="Li Ming", year="2025", marker="battery")

        unique, duplicates, _ = deduplicate_pairs(
            {first: [], second: []}, str(self.project_dir))

        self.assertEqual(2, len(unique))
        self.assertEqual([], duplicates)

    def test_persisted_identity_blocks_a_later_import(self):
        first = self._make_docx(
            "first_run.docx", "Persistent Identity Test",
            doi="10.5678/persist.1", marker="first")
        identity = build_document_identity(first)
        record_document_identity(
            str(self.project_dir), "doi_10.5678_persist.1",
            first, identity)

        later = self._make_docx(
            "later_run.docx", "Persistent Identity Test",
            doi="10.5678/persist.1", marker="later")
        unique, duplicates, _ = deduplicate_pairs(
            {later: []}, str(self.project_dir))

        self.assertEqual({}, unique)
        self.assertEqual(1, len(duplicates))
        self.assertEqual(
            "doi_10.5678_persist.1",
            duplicates[0]["matched_paper_id"])

    def test_changed_si_is_kept_for_reprocessing(self):
        main = self._make_docx(
            "main.docx", "Paper With Updated SI",
            doi="10.9999/updated.si", marker="main")
        old_si = self._make_docx(
            "old_si.docx", "Supporting Information",
            marker="old supporting data")
        new_si = self._make_docx(
            "new_si.docx", "Supporting Information",
            marker="new supporting data")

        identity = build_document_identity(main)
        identity["si_sha256"] = [
            build_document_identity(old_si)["file_sha256"]]
        record_document_identity(
            str(self.project_dir), "doi_10.9999_updated.si",
            main, identity)

        unique, duplicates, _ = deduplicate_pairs(
            {main: [new_si]}, str(self.project_dir))

        self.assertEqual(1, len(unique))
        self.assertEqual([], duplicates)

    def test_text_cache_without_ai_result_does_not_block_retry(self):
        main = self._make_docx(
            "retry.docx", "Retry After AI Failure",
            doi="10.7777/retry.1", marker="retry")
        identity = build_document_identity(main)
        paper_dir = self.project_dir / "text_cache" / "retry"
        paper_dir.mkdir(parents=True)
        (paper_dir / "meta.json").write_text(
            json.dumps({
                "paper_id": "retry",
                "main_file": "retry.docx",
                "identity": identity,
            }),
            encoding="utf-8")

        unique, duplicates, _ = deduplicate_pairs(
            {main: []}, str(self.project_dir))

        self.assertEqual(1, len(unique))
        self.assertEqual([], duplicates)

    def test_identity_index_concurrent_writes_do_not_lose_records(self):
        def write(index):
            record_document_identity(
                str(self.project_dir),
                "paper_" + str(index),
                str(self.root / ("paper_" + str(index) + ".pdf")),
                {
                    "file_sha256": "file_" + str(index),
                    "doi": "10.9999/concurrent." + str(index),
                    "text_sha256": "text_" + str(index),
                    "title_key": "title" + str(index),
                    "year": "2026",
                    "first_author_key": "author",
                    "si_sha256": [],
                },
            )

        with ThreadPoolExecutor(max_workers=8) as pool:
            list(pool.map(write, range(40)))

        connection = sqlite3.connect(
            self.project_dir / "dedup_index.db")
        count = connection.execute(
            "SELECT COUNT(*) FROM identities").fetchone()[0]
        connection.close()
        self.assertEqual(count, 40)
        mirror = json.loads(
            (self.project_dir / "dedup_index.json").read_text(
                encoding="utf-8"))
        self.assertEqual(len(mirror["records"]), 40)

    def test_changed_candidate_does_not_hide_later_exact_duplicate(self):
        main = self._make_docx(
            "ordered_main.docx", "Ordered Candidate",
            doi="10.9999/ordered", marker="main")
        si_a = self._make_docx(
            "si_a.docx", "Supporting", marker="A")
        si_b = self._make_docx(
            "si_b.docx", "Supporting", marker="B")
        identity = build_document_identity(main)
        first = dict(identity)
        first["si_sha256"] = [
            build_document_identity(si_a)["file_sha256"]]
        second = dict(identity)
        second["si_sha256"] = [
            build_document_identity(si_b)["file_sha256"]]
        record_document_identity(
            str(self.project_dir), "a_changed", main, first)
        record_document_identity(
            str(self.project_dir), "z_exact", main, second)

        unique, duplicates, _ = deduplicate_pairs(
            {main: [si_b]}, str(self.project_dir))
        self.assertEqual({}, unique)
        self.assertEqual(duplicates[0]["matched_paper_id"], "z_exact")


if __name__ == "__main__":
    unittest.main()

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from database import (find_duplicate_groups, get_paper, init_db,
                      insert_paper, merge_duplicate_group,
                      rename_paper_id)
from extractor import (build_document_identity, deduplicate_pairs,
                       record_document_identity)


class DatabaseDeduplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.project = self.root / "project"
        self.database_dir = self.project / "database"
        self.database_dir.mkdir(parents=True)
        self.db_path = self.database_dir / "main.db"
        init_db(str(self.db_path))

    def tearDown(self):
        self.temp_dir.cleanup()

    def _insert(self, paper_id, doi, title, lang="zh", **fields):
        data = {
            "DOI": doi,
            "Title": title,
            "Year": "2024",
            "First_Author": "Wang Lei",
            "_paper_id": paper_id.rsplit("__", 1)[0],
            "_schema_ver": "v1.0_basic",
            "_lang": lang,
            "_model": "deepseek-chat",
        }
        data.update(fields)
        insert_paper(
            str(self.db_path), paper_id, data,
            "v1.0_basic", "deepseek-chat", lang)

    def _make_docx(self, name, doi, title):
        path = self.root / name
        doc = Document()
        doc.core_properties.title = title
        doc.core_properties.author = "Wang Lei"
        doc.add_paragraph(title)
        doc.add_paragraph("DOI: " + doi)
        doc.add_paragraph("Published in 2024")
        for index in range(50):
            doc.add_paragraph(
                f"Abstract {index}. This is a complete research article "
                "with methods, results, characterization and discussion.")
        doc.save(path)
        return str(path)

    def test_current_library_blocks_import_before_ai(self):
        doi = "10.1234/current.library.1"
        self._insert(
            "existing__zh", doi, "Current Library Identity")
        incoming = self._make_docx(
            "incoming.docx", doi, "Current Library Identity")

        unique, duplicates, _ = deduplicate_pairs(
            {incoming: []},
            str(self.project),
            current_db_path=str(self.db_path),
            current_db_name="main")

        self.assertEqual({}, unique)
        self.assertEqual(1, len(duplicates))
        self.assertTrue(duplicates[0]["current_library"])
        self.assertEqual("existing", duplicates[0]["matched_paper_id"])

    def test_current_library_with_changed_si_becomes_update(self):
        doi = "10.1234/current.library.si"
        base_id = "doi_10.1234_current.library.si"
        self._insert(
            base_id + "__zh", doi, "Current Library SI")
        main = self._make_docx(
            "si_update_main.docx", doi, "Current Library SI")
        old_si = self._make_docx(
            "old_si.docx", "10.1234/old.si", "Old SI")
        new_si = self._make_docx(
            "new_si.docx", "10.1234/new.si", "New SI")
        identity = build_document_identity(main)
        identity["si_sha256"] = [
            build_document_identity(old_si)["file_sha256"]]
        record_document_identity(
            str(self.project), base_id, main, identity)

        unique, duplicates, identities = deduplicate_pairs(
            {main: [new_si]},
            str(self.project),
            current_db_path=str(self.db_path),
            current_db_name="main")
        self.assertEqual(1, len(unique))
        self.assertEqual([], duplicates)
        self.assertEqual(
            identities[main]["_dedup_state"], "si_changed")
        self.assertTrue(
            identities[main]["_matched_current_library"])

    def test_duplicate_scan_treats_languages_as_one_base_paper(self):
        doi = "10.1234/duplicate.group.1"
        self._insert("first__zh", doi, "Duplicate Paper", lang="zh")
        self._insert("first__en", doi, "Duplicate Paper", lang="en")
        self._insert("second__zh", doi, "Duplicate Paper", lang="zh")

        groups = find_duplicate_groups(str(self.db_path))

        self.assertEqual(1, len(groups))
        self.assertEqual(["first", "second"], groups[0]["base_ids"])
        first_record = next(
            record for record in groups[0]["records"]
            if record["base_id"] == "first")
        self.assertEqual(["en", "zh"], first_record["languages"])

    def test_merge_keeps_nonempty_target_and_fills_missing_fields(self):
        doi = "10.1234/merge.group.1"
        self._insert(
            "keep__zh", doi, "Merge Paper", lang="zh",
            Performance="N/A", Journal_Full="Keep Journal")
        self._insert(
            "remove__zh", doi, "Merge Paper", lang="zh",
            Performance="Excellent", Journal_Full="Other Journal")
        self._insert(
            "remove__en", doi, "Merge Paper", lang="en",
            Performance="Excellent")

        result = merge_duplicate_group(
            str(self.db_path), "keep", ["remove"],
            str(self.project))

        self.assertTrue(result["ok"])
        merged_zh = get_paper(str(self.db_path), "keep__zh")
        self.assertEqual("Excellent", merged_zh["Performance"])
        self.assertEqual("Keep Journal", merged_zh["Journal_Full"])
        self.assertTrue(get_paper(str(self.db_path), "keep__en"))
        self.assertFalse(get_paper(str(self.db_path), "remove__zh"))
        self.assertFalse(get_paper(str(self.db_path), "remove__en"))

    def test_rename_synchronizes_database_caches_and_identity_index(self):
        self._insert(
            "old_id__zh", "10.1234/rename.1", "Rename Paper")

        text_dir = self.project / "text_cache" / "old_id"
        text_dir.mkdir(parents=True)
        (text_dir / "meta.json").write_text(
            json.dumps({"paper_id": "old_id"}),
            encoding="utf-8")

        extract_dir = self.project / "extract_cache" / "old_id"
        extract_dir.mkdir(parents=True)
        (extract_dir / "v1.0_basic_zh.json").write_text(
            json.dumps({"_paper_id": "old_id"}),
            encoding="utf-8")

        index_path = self.project / "dedup_index.json"
        index_path.write_text(
            json.dumps({
                "version": 1,
                "records": [{
                    "paper_id": "old_id",
                    "identity": {"doi": "10.1234/rename.1"},
                }],
            }),
            encoding="utf-8")

        result = rename_paper_id(
            str(self.db_path), "old_id", "new_id",
            str(self.project))

        self.assertTrue(result["ok"])
        self.assertTrue(get_paper(str(self.db_path), "new_id__zh"))
        self.assertFalse(get_paper(str(self.db_path), "old_id__zh"))
        self.assertEqual(
            "new_id",
            json.loads(
                (self.project / "text_cache" / "new_id" / "meta.json")
                .read_text(encoding="utf-8"))["paper_id"])
        self.assertEqual(
            "new_id",
            json.loads(
                (self.project / "extract_cache" / "new_id" /
                 "v1.0_basic_zh.json").read_text(
                    encoding="utf-8"))["_paper_id"])
        self.assertEqual(
            "new_id",
            json.loads(index_path.read_text(
                encoding="utf-8"))["records"][0]["paper_id"])


if __name__ == "__main__":
    unittest.main()
